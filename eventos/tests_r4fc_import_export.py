import io

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from invitaciones.models import EventoBoda, Grupoinvitacion, Invitado
from organizaciones.models import EmpresaSuscriptora, MembresiaEmpresa

from eventos.workspace_guests_io import IMPORT_HEADERS


class GuestsImportExportR4FCTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.empresa = EmpresaSuscriptora.objects.create(
            nombre_comercial="Empresa R4FC",
            slug="empresa-r4fc",
        )
        self.admin = User.objects.create_user(
            username="admin-r4fc",
            password="Pass-R4FC-2026!",
        )
        MembresiaEmpresa.objects.create(
            empresa=self.empresa,
            usuario=self.admin,
            rol="ADMIN_EMPRESA",
        )
        self.planner = User.objects.create_user(
            username="planner-r4fc",
            password="Pass-R4FC-2026!",
        )
        MembresiaEmpresa.objects.create(
            empresa=self.empresa,
            usuario=self.planner,
            rol="WEDDING_PLANNER",
        )
        self.evento = EventoBoda.objects.create(
            empresa=self.empresa,
            nombre_evento="Evento R4FC",
            wedding_planner=self.planner,
            capacidad_contratada=10,
        )

    def url(self, name):
        return reverse(
            name,
            kwargs={
                "empresa_slug": self.empresa.slug,
                "evento_id": self.evento.id,
            },
        )

    def workbook_upload(self, rows, filename="invitados.xlsx"):
        wb = Workbook()
        ws = wb.active
        ws.title = "Invitados"
        ws.append(IMPORT_HEADERS)
        for row in rows:
            ws.append(row)
        data = io.BytesIO()
        wb.save(data)
        return SimpleUploadedFile(
            filename,
            data.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )

    def family_rows(self):
        return [
            [
                "Familia Demo", "FAMILIAR", "Ana", "Demo", "ADULTO",
                "", "", "4770000000", "", "SEGUN_TIPO", "", "", "", "",
                "NO", 0,
            ],
            [
                "Familia Demo", "FAMILIAR", "Luis", "Demo", "NINO",
                "", "", "", "", "INFANTIL", "", "", "", "",
                "NO", 0,
            ],
        ]

    def test_template_download_is_valid_xlsx(self):
        self.client.force_login(self.admin)
        response = self.client.get(self.url("k9_evento_invitados_plantilla"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml.sheet", response["Content-Type"])
        wb = load_workbook(io.BytesIO(response.content), data_only=True)
        self.assertIn("Invitados", wb.sheetnames)
        self.assertIn("Instrucciones", wb.sheetnames)
        self.assertEqual(wb["Invitados"]["A1"].value, "Grupo / Familia")

    def test_export_contains_uuid_rsvp_and_absolute_link(self):
        group = Grupoinvitacion.objects.create(
            evento=self.evento,
            nombre_grupo="Familia Export",
            tipo="FAMILIAR",
            cantidad_maxima=1,
        )
        guest = Invitado.objects.create(
            grupo=group,
            nombre="Persona",
            apellidos="Export",
            tipo_persona="ADULTO",
            asistira=False,
        )
        self.client.force_login(self.admin)
        response = self.client.get(self.url("k9_evento_invitados_exportar"))
        self.assertEqual(response.status_code, 200)

        wb = load_workbook(io.BytesIO(response.content), data_only=True)
        ws = wb["Invitados"]
        headers = [cell.value for cell in ws[1]]
        values = [cell.value for cell in ws[2]]
        row = dict(zip(headers, values))

        self.assertEqual(row["UUID invitacion"], str(group.codigo))
        self.assertEqual(row["RSVP"], "NO")
        self.assertEqual(row["Tipo persona"], "ADULTO")
        self.assertIn(f"/invitacion/{group.codigo}/", row["Enlace invitacion"])
        self.assertIn("Resumen", wb.sheetnames)

    def test_import_family_is_atomic_and_preserves_adult_child_types(self):
        self.client.force_login(self.admin)
        upload = self.workbook_upload(self.family_rows())
        response = self.client.post(
            self.url("k9_evento_invitados_importar"),
            {"archivo_invitados": upload},
        )
        self.assertEqual(response.status_code, 302)

        group = Grupoinvitacion.objects.get(
            evento=self.evento,
            nombre_grupo="Familia Demo",
        )
        self.assertEqual(group.invitados.count(), 2)
        self.assertEqual(
            group.invitados.filter(tipo_persona="ADULTO").count(),
            1,
        )
        self.assertEqual(
            group.invitados.filter(tipo_persona="NINO").count(),
            1,
        )
        self.assertTrue(group.codigo)

    def test_import_personal_creates_new_uuid_and_extra_placeholder(self):
        rows = [[
            "Carlos Nuevo", "PERSONAL", "Carlos", "Perez", "ADULTO",
            "", "", "4771111111", "", "SEGUN_TIPO", "", "", "", "",
            "SI", 1,
        ]]
        self.client.force_login(self.admin)
        response = self.client.post(
            self.url("k9_evento_invitados_importar"),
            {"archivo_invitados": self.workbook_upload(rows)},
        )
        self.assertEqual(response.status_code, 302)

        group = Grupoinvitacion.objects.get(
            evento=self.evento,
            nombre_grupo="Carlos Nuevo",
        )
        self.assertEqual(group.invitados.count(), 2)
        self.assertEqual(
            group.invitados.filter(es_acompanante_extra=True).count(),
            1,
        )
        principal = group.invitados.get(es_acompanante_extra=False)
        self.assertEqual(principal.nombre, "Carlos")
        self.assertIsNone(principal.asistira)

    def test_existing_group_name_aborts_entire_import(self):
        Grupoinvitacion.objects.create(
            evento=self.evento,
            nombre_grupo="Familia Demo",
            tipo="FAMILIAR",
            cantidad_maxima=1,
        )
        rows = self.family_rows() + [[
            "Otro Grupo", "PERSONAL", "Otro", "Invitado", "ADULTO",
            "", "", "", "", "SEGUN_TIPO", "", "", "", "", "NO", 0,
        ]]

        self.client.force_login(self.admin)
        response = self.client.post(
            self.url("k9_evento_invitados_importar"),
            {"archivo_invitados": self.workbook_upload(rows)},
        )
        self.assertEqual(response.status_code, 302)
        self.assertFalse(
            Grupoinvitacion.objects.filter(
                evento=self.evento,
                nombre_grupo="Otro Grupo",
            ).exists()
        )

    def test_duplicate_person_in_file_aborts_entire_import(self):
        rows = self.family_rows()
        rows.append(list(rows[0]))

        self.client.force_login(self.admin)
        self.client.post(
            self.url("k9_evento_invitados_importar"),
            {"archivo_invitados": self.workbook_upload(rows)},
        )
        self.assertFalse(
            Grupoinvitacion.objects.filter(
                evento=self.evento,
                nombre_grupo="Familia Demo",
            ).exists()
        )

    def test_invalid_menu_aborts_entire_import(self):
        rows = self.family_rows()
        rows[0][9] = "MENU_RARO"

        self.client.force_login(self.admin)
        self.client.post(
            self.url("k9_evento_invitados_importar"),
            {"archivo_invitados": self.workbook_upload(rows)},
        )
        self.assertFalse(
            Grupoinvitacion.objects.filter(evento=self.evento).exists()
        )

    def test_capacity_overflow_aborts_without_partial_rows(self):
        self.evento.capacidad_contratada = 1
        self.evento.save(update_fields=["capacidad_contratada"])

        self.client.force_login(self.admin)
        self.client.post(
            self.url("k9_evento_invitados_importar"),
            {"archivo_invitados": self.workbook_upload(self.family_rows())},
        )
        self.assertFalse(
            Grupoinvitacion.objects.filter(evento=self.evento).exists()
        )

    def test_import_does_not_accept_csv(self):
        self.client.force_login(self.admin)
        bad = SimpleUploadedFile(
            "lista.csv",
            b"grupo,nombre\nFamilia,Ana",
            content_type="text/csv",
        )
        self.client.post(
            self.url("k9_evento_invitados_importar"),
            {"archivo_invitados": bad},
        )
        self.assertFalse(
            Grupoinvitacion.objects.filter(evento=self.evento).exists()
        )

    def test_assigned_planner_can_import_and_export(self):
        self.client.force_login(self.planner)
        imported = self.client.post(
            self.url("k9_evento_invitados_importar"),
            {"archivo_invitados": self.workbook_upload(self.family_rows())},
        )
        self.assertEqual(imported.status_code, 302)
        exported = self.client.get(self.url("k9_evento_invitados_exportar"))
        self.assertEqual(exported.status_code, 200)
        self.assertTrue(exported.content)

    def test_template_contains_no_rsvp_import_column(self):
        self.client.force_login(self.admin)
        response = self.client.get(self.url("k9_evento_invitados_plantilla"))
        wb = load_workbook(io.BytesIO(response.content), data_only=True)
        headers = [cell.value for cell in wb["Invitados"][1]]
        self.assertNotIn("RSVP", headers)
        self.assertNotIn("UUID invitacion", headers)
