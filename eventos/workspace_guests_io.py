import io
import re
from collections import OrderedDict
from dataclasses import dataclass, field

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from core.services.auditoria import registrar_auditoria
from invitaciones.guest_domain import asegurar_roster_grupo
from invitaciones.models import Grupoinvitacion, Invitado

from .workspace_guests_capacity import resumen_cupo_evento


IMPORT_HEADERS = [
    "Grupo / Familia",
    "Tipo invitacion",
    "Nombre",
    "Apellidos",
    "Tipo persona",
    "Telefono persona",
    "Correo persona",
    "Telefono contacto grupo",
    "Correo contacto grupo",
    "Menu",
    "Mesa",
    "Restricciones alimentarias",
    "Alergias",
    "Notas",
    "Permitir acompanantes extra",
    "Acompanantes extra permitidos",
]

HEADER_ALIASES = {
    "grupo_familia": "grupo",
    "grupo": "grupo",
    "familia": "grupo",
    "nombre_grupo": "grupo",
    "tipo_invitacion": "tipo_grupo",
    "tipo_grupo": "tipo_grupo",
    "nombre": "nombre",
    "apellidos": "apellidos",
    "apellido": "apellidos",
    "tipo_persona": "tipo_persona",
    "adulto_nino": "tipo_persona",
    "telefono_persona": "telefono",
    "telefono_invitado": "telefono",
    "correo_persona": "correo",
    "correo_invitado": "correo",
    "telefono_contacto_grupo": "telefono_grupo",
    "correo_contacto_grupo": "correo_grupo",
    "menu": "menu",
    "menu_asignado": "menu",
    "mesa": "mesa",
    "restricciones_alimentarias": "restricciones",
    "restricciones": "restricciones",
    "alergias": "alergias",
    "notas": "notas",
    "permitir_acompanantes_extra": "permitir_extras",
    "acompanantes_extra_permitidos": "extras",
    "acompanantes_permitidos": "extras",
}

TIPOS_GRUPO = {"PERSONAL", "FAMILIAR"}
TIPOS_PERSONA = {"ADULTO", "NINO"}
TIPOS_MENU = {"SEGUN_TIPO", "ADULTO", "INFANTIL"}
SI_NO_TRUE = {"SI", "SÍ", "TRUE", "1", "YES", "Y"}
SI_NO_FALSE = {"NO", "FALSE", "0", "N"}

EXPORT_HEADERS = [
    "UUID invitacion",
    "Grupo / Familia",
    "Tipo invitacion",
    "Nombre",
    "Apellidos",
    "Tipo persona",
    "Es acompanante extra",
    "Telefono persona",
    "Correo persona",
    "RSVP",
    "Fecha confirmacion",
    "Menu configurado",
    "Menu efectivo",
    "Mesa",
    "Restricciones alimentarias",
    "Alergias",
    "Notas",
    "Telefono contacto grupo",
    "Correo contacto grupo",
    "Lugares del grupo",
    "RSVP grupo",
    "Enlace invitacion",
]


@dataclass
class ImportIssue:
    row: int
    message: str

    def __str__(self):
        prefix = f"Fila {self.row}: " if self.row else ""
        return prefix + self.message


@dataclass
class ParsedPerson:
    row: int
    nombre: str
    apellidos: str = ""
    tipo_persona: str = "ADULTO"
    telefono: str = ""
    correo: str = ""
    menu: str = "SEGUN_TIPO"
    mesa: str = ""
    restricciones: str = ""
    alergias: str = ""
    notas: str = ""


@dataclass
class ParsedGroup:
    nombre: str
    tipo: str
    telefono: str = ""
    correo: str = ""
    permitir_extras: bool = False
    extras: int = 0
    people: list = field(default_factory=list)


@dataclass
class ParsedImport:
    groups: list
    issues: list
    total_people: int
    adults: int
    children: int
    extra_placeholders: int


def _normalize_header(value):
    text = str(value or "").strip().lower()
    text = (
        text.replace("á", "a")
        .replace("é", "e")
        .replace("í", "i")
        .replace("ó", "o")
        .replace("ú", "u")
        .replace("ñ", "n")
    )
    return re.sub(r"[^a-z0-9]+", "_", text).strip("_")


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _int(value, default=0):
    if value in (None, ""):
        return default
    try:
        number = int(float(value))
    except (TypeError, ValueError):
        raise ValueError("debe ser un número entero")
    if number < 0:
        raise ValueError("no puede ser negativo")
    return number


def _bool(value, default=False):
    text = _text(value).upper()
    if not text:
        return default
    if text in SI_NO_TRUE:
        return True
    if text in SI_NO_FALSE:
        return False
    raise ValueError("usa SI o NO")


def _rsvp_label(value):
    if value is True:
        return "SI"
    if value is False:
        return "NO"
    return "PENDIENTE"


def _group_rsvp_label(group):
    people = list(group.invitados.all())
    if not people:
        return "PENDIENTE"
    yes = sum(1 for p in people if p.asistira is True)
    no = sum(1 for p in people if p.asistira is False)
    pending = len(people) - yes - no
    if pending:
        return "PENDIENTE"
    if yes and not no:
        return "CONFIRMADA"
    if no and not yes:
        return "NO ASISTE"
    return "MIXTA"


def _style_workbook(ws, headers):
    fill = PatternFill("solid", fgColor="1F3A32")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(headers)).column_letter}{max(ws.max_row, 1)}"


def build_import_template_bytes():
    wb = Workbook()
    ws = wb.active
    ws.title = "Invitados"
    ws.append(IMPORT_HEADERS)
    examples = [
        ["Familia Lopez", "FAMILIAR", "Juan", "Lopez", "ADULTO", "", "", "4770000000", "", "SEGUN_TIPO", "", "", "", "", "NO", 0],
        ["Familia Lopez", "FAMILIAR", "Maria", "Lopez", "ADULTO", "", "", "", "", "ADULTO", "", "", "", "", "NO", 0],
        ["Familia Lopez", "FAMILIAR", "Sofia", "Lopez", "NINO", "", "", "", "", "INFANTIL", "", "", "", "", "NO", 0],
        ["Carlos Perez", "PERSONAL", "Carlos", "Perez", "ADULTO", "", "", "4770000001", "", "SEGUN_TIPO", "", "", "", "", "SI", 1],
    ]
    for row in examples:
        ws.append(row)

    widths = [24, 17, 18, 20, 14, 18, 26, 23, 28, 15, 12, 28, 24, 24, 24, 20]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, index).column_letter].width = width

    dv_group = DataValidation(type="list", formula1='"PERSONAL,FAMILIAR"', allow_blank=False)
    dv_person = DataValidation(type="list", formula1='"ADULTO,NINO"', allow_blank=False)
    dv_menu = DataValidation(type="list", formula1='"SEGUN_TIPO,ADULTO,INFANTIL"', allow_blank=False)
    dv_bool = DataValidation(type="list", formula1='"SI,NO"', allow_blank=True)
    for validation in (dv_group, dv_person, dv_menu, dv_bool):
        ws.add_data_validation(validation)
    dv_group.add("B2:B5000")
    dv_person.add("E2:E5000")
    dv_menu.add("J2:J5000")
    dv_bool.add("O2:O5000")

    _style_workbook(ws, IMPORT_HEADERS)

    info = wb.create_sheet("Instrucciones")
    info_rows = [
        ["PLANTILLA OFICIAL DIRTEC EVENT STUDIO - IMPORTACION DE INVITADOS", ""],
        ["Regla", "Descripcion"],
        ["Una persona por fila", "En una familia repite el mismo Grupo / Familia para cada integrante."],
        ["PERSONAL", "Debe tener una sola persona nominal. Los acompañantes extra se crean como lugares autorizados."],
        ["FAMILIAR", "Puede tener varias personas, una por fila."],
        ["Tipo persona", "Usa ADULTO o NINO."],
        ["Menu", "Usa SEGUN_TIPO, ADULTO o INFANTIL."],
        ["RSVP", "No se importa. Nunca se sobrescriben confirmaciones con esta plantilla."],
        ["UUID", "No se importa ni reutiliza. Cada grupo nuevo obtiene su UUID propio."],
        ["Duplicados", "Si un grupo ya existe en el evento o la lista repite una persona, no se guarda ninguna fila."],
        ["Capacidad", "Se valida antes de guardar. La importación es transaccional: todo o nada."],
        ["Ejemplos", "Elimina las filas de ejemplo antes de subir tu lista real."],
    ]
    for row in info_rows:
        info.append(row)
    info.merge_cells("A1:B1")
    info["A1"].fill = PatternFill("solid", fgColor="1F3A32")
    info["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    info["A1"].alignment = Alignment(horizontal="center")
    for cell in info[2]:
        cell.fill = PatternFill("solid", fgColor="1F3A32")
        cell.font = Font(color="FFFFFF", bold=True)
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 88
    for row in info.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def parse_import_workbook(uploaded_file):
    issues = []
    try:
        wb = load_workbook(uploaded_file, data_only=True, read_only=True)
    except Exception as exc:
        return ParsedImport([], [ImportIssue(0, f"No se pudo abrir el Excel: {exc}")], 0, 0, 0, 0)

    if "Invitados" in wb.sheetnames:
        ws = wb["Invitados"]
    else:
        ws = wb.active

    raw_headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    mapped = {}
    for index, value in enumerate(raw_headers):
        alias = HEADER_ALIASES.get(_normalize_header(value))
        if alias and alias not in mapped:
            mapped[alias] = index

    required = {"grupo", "tipo_grupo", "nombre", "tipo_persona"}
    missing = sorted(required - set(mapped))
    if missing:
        labels = ", ".join(missing)
        return ParsedImport([], [ImportIssue(1, f"Faltan columnas obligatorias: {labels}.")], 0, 0, 0, 0)

    groups = OrderedDict()
    seen_people = set()

    def val(row, key):
        idx = mapped.get(key)
        return row[idx] if idx is not None and idx < len(row) else None

    for row_number, cells in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        row = list(cells)
        if not any(_text(value) for value in row):
            continue

        group_name = _text(val(row, "grupo"))
        group_type = _text(val(row, "tipo_grupo")).upper()
        name = _text(val(row, "nombre"))
        person_type = _text(val(row, "tipo_persona")).upper() or "ADULTO"
        menu = _text(val(row, "menu")).upper() or "SEGUN_TIPO"

        if not group_name:
            issues.append(ImportIssue(row_number, "Grupo / Familia es obligatorio."))
        if group_type not in TIPOS_GRUPO:
            issues.append(ImportIssue(row_number, "Tipo invitacion debe ser PERSONAL o FAMILIAR."))
        if not name:
            issues.append(ImportIssue(row_number, "Nombre es obligatorio."))
        if person_type not in TIPOS_PERSONA:
            issues.append(ImportIssue(row_number, "Tipo persona debe ser ADULTO o NINO."))
        if menu not in TIPOS_MENU:
            issues.append(ImportIssue(row_number, "Menu debe ser SEGUN_TIPO, ADULTO o INFANTIL."))
        if not group_name or group_type not in TIPOS_GRUPO or not name or person_type not in TIPOS_PERSONA or menu not in TIPOS_MENU:
            continue

        try:
            allow_extras = _bool(val(row, "permitir_extras"), default=False)
            extras = _int(val(row, "extras"), default=0)
        except ValueError as exc:
            issues.append(ImportIssue(row_number, f"Acompañantes extra: {exc}."))
            continue
        if not allow_extras:
            extras = 0

        normalized_group = group_name.casefold()
        current = groups.get(normalized_group)
        if current is None:
            current = ParsedGroup(
                nombre=group_name,
                tipo=group_type,
                telefono=_text(val(row, "telefono_grupo")),
                correo=_text(val(row, "correo_grupo")),
                permitir_extras=allow_extras,
                extras=extras,
            )
            groups[normalized_group] = current
        else:
            if current.tipo != group_type:
                issues.append(ImportIssue(row_number, f"El grupo '{group_name}' mezcla PERSONAL y FAMILIAR."))
            if current.extras != extras or current.permitir_extras != allow_extras:
                issues.append(ImportIssue(row_number, f"El grupo '{group_name}' tiene configuración de acompañantes inconsistente."))
            if not current.telefono:
                current.telefono = _text(val(row, "telefono_grupo"))
            if not current.correo:
                current.correo = _text(val(row, "correo_grupo"))

        person = ParsedPerson(
            row=row_number,
            nombre=name,
            apellidos=_text(val(row, "apellidos")),
            tipo_persona=person_type,
            telefono=_text(val(row, "telefono")),
            correo=_text(val(row, "correo")),
            menu=menu,
            mesa=_text(val(row, "mesa")),
            restricciones=_text(val(row, "restricciones")),
            alergias=_text(val(row, "alergias")),
            notas=_text(val(row, "notas")),
        )

        duplicate_key = (
            normalized_group,
            person.nombre.casefold(),
            person.apellidos.casefold(),
        )
        if duplicate_key in seen_people:
            issues.append(ImportIssue(row_number, f"Persona duplicada en '{group_name}': {person.nombre} {person.apellidos}".strip()))
        else:
            seen_people.add(duplicate_key)
            current.people.append(person)

    for group in groups.values():
        if group.tipo == "PERSONAL" and len(group.people) != 1:
            row = group.people[0].row if group.people else 0
            issues.append(ImportIssue(row, f"La invitación PERSONAL '{group.nombre}' debe tener exactamente una persona nominal."))
        if not group.people:
            issues.append(ImportIssue(0, f"El grupo '{group.nombre}' no tiene personas."))

    adults = sum(
        1 for group in groups.values() for person in group.people
        if person.tipo_persona == "ADULTO"
    )
    children = sum(
        1 for group in groups.values() for person in group.people
        if person.tipo_persona == "NINO"
    )
    extras = sum(group.extras for group in groups.values())
    total = adults + children + extras

    return ParsedImport(list(groups.values()), issues, total, adults, children, extras)


def validate_import_against_event(evento, parsed):
    issues = list(parsed.issues)
    if issues:
        return issues

    existing_names = {
        value.casefold()
        for value in Grupoinvitacion.objects.filter(evento=evento)
        .values_list("nombre_grupo", flat=True)
    }
    for group in parsed.groups:
        if group.nombre.casefold() in existing_names:
            issues.append(
                ImportIssue(
                    group.people[0].row if group.people else 0,
                    f"Ya existe una invitación llamada '{group.nombre}'. No se sobrescribirá.",
                )
            )

    capacity = resumen_cupo_evento(evento)
    if capacity.capacidad and capacity.reservados + parsed.total_people > capacity.capacidad:
        available = max(capacity.capacidad - capacity.reservados, 0)
        issues.append(
            ImportIssue(
                0,
                f"La importación requiere {parsed.total_people} lugares y solo hay {available} disponibles.",
            )
        )

    if capacity.capacidad_separada:
        adults = capacity.adultos
        children = capacity.ninos
        proposed_adults = parsed.adults + parsed.extra_placeholders
        proposed_children = parsed.children

        if adults.get("capacidad") and adults["reservados"] + proposed_adults > adults["capacidad"]:
            available = max(adults["capacidad"] - adults["reservados"], 0)
            issues.append(
                ImportIssue(
                    0,
                    f"Adultos: se requieren {proposed_adults} lugares nuevos y solo hay {available} disponibles.",
                )
            )
        if children.get("capacidad") and children["reservados"] + proposed_children > children["capacidad"]:
            available = max(children["capacidad"] - children["reservados"], 0)
            issues.append(
                ImportIssue(
                    0,
                    f"Niños: se requieren {proposed_children} lugares nuevos y solo hay {available} disponibles.",
                )
            )

    return issues


@transaction.atomic
def apply_guest_import(evento, parsed, *, user=None, request=None):
    issues = validate_import_against_event(evento, parsed)
    if issues:
        raise ValidationError([str(issue) for issue in issues])

    created_groups = 0
    created_people = 0

    for group_data in parsed.groups:
        group = Grupoinvitacion.objects.create(
            evento=evento,
            nombre_grupo=group_data.nombre[:100],
            tipo=group_data.tipo,
            cantidad_maxima=max(len(group_data.people), 1),
            permitir_acompanantes_extra=group_data.permitir_extras,
            cantidad_extra_permitida=group_data.extras,
            telefono_contacto=group_data.telefono[:20] or None,
            correo_contacto=group_data.correo or None,
        )
        created_groups += 1

        if group.es_personal:
            # Create the canonical principal with guest_domain, then enrich it
            # from the single imported row. This preserves the same authority
            # as manual creation.
            asegurar_roster_grupo(group)
            person_data = group_data.people[0]
            principal = group.invitados.filter(es_acompanante_extra=False).order_by("orden", "id").first()
            principal.nombre = person_data.nombre[:100]
            principal.apellidos = person_data.apellidos[:120] or None
            principal.telefono = person_data.telefono[:30] or None
            principal.correo = person_data.correo or None
            principal.tipo_persona = person_data.tipo_persona
            principal.menu_asignado = person_data.menu
            principal.mesa = person_data.mesa[:50] or None
            principal.restricciones_alimentarias = person_data.restricciones or None
            principal.alergias = person_data.alergias or None
            principal.notas = person_data.notas or None
            principal.save()
            created_people += 1 + group_data.extras
        else:
            for order, person_data in enumerate(group_data.people, 1):
                Invitado.objects.create(
                    grupo=group,
                    nombre=person_data.nombre[:100],
                    apellidos=person_data.apellidos[:120] or None,
                    telefono=person_data.telefono[:30] or None,
                    correo=person_data.correo or None,
                    tipo_persona=person_data.tipo_persona,
                    menu_asignado=person_data.menu,
                    mesa=person_data.mesa[:50] or None,
                    restricciones_alimentarias=person_data.restricciones or None,
                    alergias=person_data.alergias or None,
                    notas=person_data.notas or None,
                    orden=order,
                )
                created_people += 1
            asegurar_roster_grupo(group)
            created_people += group_data.extras

        group.cantidad_maxima = max(
            group.invitados.filter(es_acompanante_extra=False).count(),
            1,
        )
        group.save(update_fields=["cantidad_maxima"])

    registrar_auditoria(
        usuario=user,
        empresa=evento.empresa,
        evento=evento,
        accion="IMPORTAR_INVITADOS_EXCEL_K9",
        modelo="EventoBoda",
        objeto_id=evento.id,
        descripcion=(
            f"Importación Excel: {created_groups} invitaciones y "
            f"{created_people} personas/lugares creados."
        ),
        valores_nuevos={
            "grupos_creados": created_groups,
            "personas_creadas": created_people,
            "adultos_nominales": parsed.adults,
            "ninos_nominales": parsed.children,
            "acompanantes_extra": parsed.extra_placeholders,
        },
        request=request,
    )

    return {
        "groups": created_groups,
        "people": created_people,
    }


def build_export_bytes(evento, *, base_url=""):
    wb = Workbook()
    ws = wb.active
    ws.title = "Invitados"
    ws.append(EXPORT_HEADERS)

    groups = (
        Grupoinvitacion.objects.filter(evento=evento)
        .prefetch_related("invitados")
        .order_by("nombre_grupo", "id")
    )

    for group in groups:
        path = f"/invitacion/{group.codigo}/"
        link = f"{base_url.rstrip('/')}{path}" if base_url else path
        group_rsvp = _group_rsvp_label(group)

        for person in group.invitados.all():
            ws.append([
                str(group.codigo),
                group.nombre_grupo,
                group.tipo,
                person.nombre,
                person.apellidos or "",
                person.tipo_persona,
                "SI" if person.es_acompanante_extra else "NO",
                person.telefono or "",
                person.correo or "",
                _rsvp_label(person.asistira),
                timezone.localtime(person.fecha_confirmacion).strftime("%Y-%m-%d %H:%M")
                if person.fecha_confirmacion else "",
                person.menu_asignado,
                person.menu_buffet_efectivo,
                person.mesa or "",
                person.restricciones_alimentarias or "",
                person.alergias or "",
                person.notas or "",
                group.telefono_contacto or "",
                group.correo_contacto or "",
                group.total_lugares,
                group_rsvp,
                link,
            ])

    widths = [
        38, 24, 17, 18, 20, 14, 18, 18, 27, 14, 20, 18, 16, 12,
        28, 24, 24, 23, 28, 16, 16, 52,
    ]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, index).column_letter].width = width

    _style_workbook(ws, EXPORT_HEADERS)

    summary = wb.create_sheet("Resumen")
    capacity = resumen_cupo_evento(evento)
    summary_rows = [
        ["DIRTEC EVENT STUDIO - RESUMEN DE INVITADOS", ""],
        ["Evento", str(evento)],
        ["Capacidad global", capacity.capacidad or "Sin límite"],
        ["Reservados activos", capacity.reservados],
        ["Disponibles", capacity.disponibles if capacity.disponibles is not None else "Sin límite"],
        ["Liberados por NO", capacity.liberados],
        ["Adultos confirmados", capacity.adultos["confirmados"]],
        ["Adultos pendientes", capacity.adultos["pendientes"]],
        ["Adultos no asisten", capacity.adultos["no_asisten"]],
        ["Niños confirmados", capacity.ninos["confirmados"]],
        ["Niños pendientes", capacity.ninos["pendientes"]],
        ["Niños no asisten", capacity.ninos["no_asisten"]],
    ]
    for row in summary_rows:
        summary.append(row)
    summary.merge_cells("A1:B1")
    summary["A1"].fill = PatternFill("solid", fgColor="1F3A32")
    summary["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    summary["A1"].alignment = Alignment(horizontal="center")
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 42

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
