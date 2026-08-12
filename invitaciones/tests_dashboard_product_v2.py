from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from invitaciones.models import (
    EventoBoda,
    Grupoinvitacion,
    Invitado,
)
from mesas.models import AsignacionMesa, Mesa


class DashboardProductV2Tests(TestCase):
    def setUp(self):
        self.evento = EventoBoda.objects.create(
            novio='Diego',
            novia='Fernanda',
            frase_portada='Test',
            mensaje_general='Test',
            fecha_misa=timezone.now(),
            lugar_misa='Ceremonia',
            fecha_fiesta=timezone.now(),
            lugar_fiesta='Recepción',
        )
        User = get_user_model()
        self.user = User.objects.create_superuser(
            username='dashboard_product_admin',
            password='test123',
        )
        self.client.force_login(self.user)

    def test_dashboard_has_one_canonical_navigation_and_real_operation_panel(self):
        response = self.client.get(
            reverse('dashboard')
            + f'?evento={self.evento.id}'
        )

        self.assertEqual(
            response.status_code,
            200,
        )
        self.assertContains(
            response,
            'data-dashboard-tab="resumen"',
        )
        self.assertContains(
            response,
            'data-tab-panel="resumen"',
        )
        for panel in (
            'servicios',
            'agenda',
            'tareas',
            'finanzas',
            'documentos',
            'invitados',
            'invitacion',
        ):
            self.assertContains(
                response,
                f'data-tab-panel="{panel}"',
            )
        self.assertNotContains(
            response,
            'class="dashboard-tabs"',
        )
        self.assertNotContains(
            response,
            'class="workflow-strip"',
        )

    def test_guest_excel_is_individual_and_uses_real_table_assignment(self):
        grupo = Grupoinvitacion.objects.create(
            evento=self.evento,
            nombre_grupo='Familia Pérez',
            tipo='FAMILIAR',
        )
        ana = Invitado.objects.create(
            grupo=grupo,
            nombre='Ana Pérez',
            tipo_persona='ADULTO',
            menu_asignado='ADULTO',
            asistira=True,
        )
        luis = Invitado.objects.create(
            grupo=grupo,
            nombre='Luis Pérez',
            tipo_persona='NINO',
            menu_asignado='SEGUN_TIPO',
            asistira=False,
        )
        mesa = Mesa.objects.create(
            evento=self.evento,
            nombre='Mesa 7',
            capacidad=10,
        )
        AsignacionMesa.objects.create(
            mesa=mesa,
            invitado=ana,
        )

        response = self.client.get(
            reverse('exportar_excel')
            + f'?evento={self.evento.id}'
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        wb = load_workbook(
            BytesIO(response.content)
        )
        ws = wb['Confirmaciones']
        rows = list(
            ws.iter_rows(values_only=True)
        )

        self.assertEqual(
            rows[0],
            (
                'Evento',
                'Invitación / grupo',
                'Tipo de invitación',
                'Persona',
                'Adulto o niño',
                'Buffet asignado',
                'Asistencia',
                'Mesa',
                'Acompañante extra',
                'Código',
            ),
        )

        data = {
            row[3]: row
            for row in rows[1:]
        }
        self.assertEqual(
            data['Ana Pérez'][6],
            'Sí',
        )
        self.assertEqual(
            data['Ana Pérez'][7],
            'Mesa 7',
        )
        self.assertEqual(
            data['Luis Pérez'][4],
            'Niño',
        )
        self.assertEqual(
            data['Luis Pérez'][5],
            'Infantil',
        )
        self.assertEqual(
            data['Luis Pérez'][6],
            'No',
        )
        self.assertEqual(
            data['Luis Pérez'][7],
            'Sin mesa',
        )
        self.assertNotIn(
            'Cantidad confirmada',
            rows[0],
        )
        self.assertNotIn(
            'Comentarios',
            rows[0],
        )

    def test_personal_export_does_not_recreate_legacy_pass_counts(self):
        grupo = Grupoinvitacion.objects.create(
            evento=self.evento,
            nombre_grupo='Carlos',
            tipo='PERSONAL',
            permitir_acompanantes_extra=True,
            cantidad_extra_permitida=1,
        )
        Invitado.objects.create(
            grupo=grupo,
            nombre='Carlos',
            asistira=True,
            es_acompanante_extra=False,
        )
        Invitado.objects.create(
            grupo=grupo,
            nombre='Acompañante 1',
            asistira=None,
            es_acompanante_extra=True,
        )

        response = self.client.get(
            reverse('exportar_excel')
            + f'?evento={self.evento.id}'
        )

        wb = load_workbook(
            BytesIO(response.content)
        )
        ws = wb['Confirmaciones']
        people = [
            row[3]
            for row in ws.iter_rows(
                min_row=2,
                values_only=True,
            )
        ]

        self.assertEqual(
            sorted(people),
            ['Acompañante 1', 'Carlos'],
        )
        self.assertEqual(
            len(people),
            2,
        )

        rows = {
            row[3]: row
            for row in ws.iter_rows(
                min_row=2,
                values_only=True,
            )
        }
        self.assertEqual(
            rows['Carlos'][6],
            'Sí',
        )
        self.assertEqual(
            rows['Acompañante 1'][6],
            'Pendiente',
        )
        self.assertEqual(
            rows['Acompañante 1'][8],
            'Sí',
        )
