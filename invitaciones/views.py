import csv
import io
import json
import re
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation

from django.contrib import messages
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import transaction
from django.db.models.deletion import ProtectedError
from django.http import HttpResponse, JsonResponse
from django.db.models import Prefetch, Sum
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.views.decorators.clickjacking import xframe_options_sameorigin
from django.views.decorators.http import require_POST
from django.utils import timezone
from django.utils.crypto import get_random_string
from django.utils.text import slugify
from openpyxl import Workbook, load_workbook
from urllib.parse import quote

from .models import (
    AssetInvitacion,
    ComponenteInvitacion,
    DetalleProduccionEvento,
    DisenoInvitacion,
    EnlaceRegalo,
    EventoBoda,
    FotoEvento,
    Grupoinvitacion,
    Invitado,
    ItinerarioEvento,
    MenuBoda,
    PersonaCeremonia,
    PlantillaInvitacion,
    SeccionInvitacion,
    VersionDisenoInvitacion,
    es_video_archivo,
    google_maps_src,
)
from .guest_domain import asegurar_roster_grupo
from .guest_analytics import (
    resumen_invitados_evento,
    resumen_invitados_eventos,
)
from .permissions import (
    empresas_para_usuario,
    eventos_visibles_usuario,
    roles_activos_empresa,
    usuario_es_dirtec,
    usuario_puede_gestionar_catalogos,
    usuario_puede_gestionar_usuarios,
)
from catering.models import Alimento, CateringEvento, PaqueteBuffet
from decoracion.models import ElementoDecoracion
from entretenimiento.models import CancionEvento, EntretenimientoEvento
from itinerario.models import ActividadItinerario
from mesas.models import AsignacionMesa, Mesa
from paquetes.models import PaqueteBoda
from presupuesto.models import CategoriaGasto, GastoEvento, PagoEvento
from proveedores.models import PersonalEvento, Proveedor, ServicioEvento
from suscripciones.models import PagoSuscripcion, PlanSuscripcion, SuscripcionEmpresa
from aprobaciones.models import AprobacionEvento
from documentos.models import DocumentoEvento
from notificaciones.models import Notificacion
from organizaciones.models import EmpresaSuscriptora, MembresiaEmpresa, SedeEvento
from tareas.models import TareaEvento
from eventos.planner_dashboard_v3 import construir_contexto_dashboard_planner_v3
from eventos.company_dashboard_v3 import construir_contexto_dashboard_empresa_v3
from eventos.client_portal_v3 import construir_contexto_portal_cliente_v3
from eventos.client_guests_v3 import construir_contexto_invitados_cliente_v3
from eventos.provider_portal_v3 import construir_contexto_portal_proveedor_v3
from core.services.auditoria import registrar_auditoria
from core.services.authorization import Actions, usuario_puede_evento
from core.services.identity import (
    actualizar_identidad_usuario,
    crear_identidad_usuario,
    perfil_acceso_de,
)
from core.services.tenant_context import (
    exigir_tenant,
    tenant_desde_request,
    validar_slug_tenant,
)
from colaboracion.services import (
    preparar_expedientes_cliente,
    preparar_expedientes_planner,
    preparar_expedientes_proveedor,
    presupuesto_cliente,
)
from colaboracion.models import ExpedienteServicio

from core.services.limites_plan import (
    LimitePlanExcedido,
    resumen_uso_plan,
    validar_limite_eventos_activos,
    validar_limite_clientes,
    validar_limite_planners,
    validar_limite_usuarios,
)
from .shared_dashboard_services import (
    bool_post,
    convertir_decimal,
    convertir_entero,
    limpiar_texto,
    resumen_operativo_eventos,
)
from .company_user_services import (
    aplicar_datos_usuario_empresa,
    actualizar_usuario_empresa_dashboard,
    crear_o_actualizar_usuario_empresa,
    desactivar_usuario_empresa_dashboard,
    eliminar_usuario_empresa_dashboard,
)




LOGIN_DASHBOARD_URL = '/login/'


ACCIONES_PANEL_EMPRESA_USUARIOS = {
    'crear_planner',
    'crear_cliente',
    'editar_usuario_empresa',
    'desactivar_usuario_empresa',
    'eliminar_usuario_empresa',
}
ACCIONES_PANEL_EMPRESA_CATALOGOS = {
    'crear_sede',
    'editar_sede',
    'desactivar_sede',
    'eliminar_sede',
    'crear_proveedor',
    'editar_proveedor',
    'desactivar_proveedor',
    'eliminar_proveedor',
    'crear_paquete',
    'editar_paquete',
    'desactivar_paquete',
    'eliminar_paquete',
}


def bloquear_accion_dashboard(request, empresa=None, evento=None, accion='', permiso=''):
    registrar_auditoria(
        usuario=request.user,
        empresa=empresa,
        evento=evento,
        accion='ACCESO_DENEGADO_DASHBOARD',
        modelo='Dashboard',
        objeto_id=getattr(empresa, 'id', '') or getattr(evento, 'id', ''),
        descripcion=f'Accion bloqueada: {accion}. Permiso requerido: {permiso}.',
        valores_nuevos={'accion_solicitada': accion, 'permiso_requerido': permiso},
        request=request,
    )
    raise PermissionDenied('No tienes permiso para realizar esta accion.')


def slug_empresa_unico(nombre):
    base = slugify(nombre) or 'empresa'
    slug = base
    contador = 2
    while EmpresaSuscriptora.objects.filter(slug=slug).exists():
        slug = f'{base}-{contador}'
        contador += 1
    return slug


def usuario_empresa_existente_activo(empresa, username, rol=None):
    if not username:
        return False
    filtros = {
        'empresa': empresa,
        'usuario__username': username,
        'activo': True,
    }
    if rol:
        filtros['rol'] = rol
    return MembresiaEmpresa.objects.filter(**filtros).exists()


def bloquear_por_limite_plan(request, empresa, mensaje, destino):
    messages.error(request, mensaje)
    registrar_auditoria(
        usuario=request.user,
        empresa=empresa,
        accion='LIMITE_PLAN_EXCEDIDO',
        modelo='EmpresaSuscriptora',
        objeto_id=empresa.id,
        descripcion=mensaje,
        request=request,
    )
    return redirect(f'/dashboard/empresa/?empresa={empresa.id}{destino}')


def fecha_hora_dashboard(request, campo, default=None):
    valor = request.POST.get(campo)
    if not valor:
        return default or timezone.now()
    try:
        fecha = datetime.strptime(valor, '%Y-%m-%dT%H:%M')
        return timezone.make_aware(fecha, timezone.get_current_timezone())
    except ValueError:
        return default or timezone.now()


def fecha_dashboard(request, campo, default=None):
    valor = request.POST.get(campo)
    if not valor:
        return default or timezone.localdate()
    try:
        return datetime.strptime(valor, '%Y-%m-%d').date()
    except ValueError:
        return default or timezone.localdate()


def hora_dashboard(request, campo, default=None):
    valor = request.POST.get(campo)
    if not valor:
        return default
    try:
        return datetime.strptime(valor, '%H:%M').time()
    except ValueError:
        return default


def resumen_eventos(eventos):
    eventos = eventos.select_related('empresa', 'sede', 'wedding_planner')
    hoy = timezone.localdate()
    return {
        'total': eventos.count(),
        'activos': eventos.filter(activo=True).count(),
        'proximos': eventos.filter(fecha_fiesta__date__gte=hoy).count(),
        'sin_planner': eventos.filter(wedding_planner__isnull=True).count(),
        'planeacion': eventos.filter(estado='PLANEACION').count(),
        'confirmados': eventos.filter(estado='CONFIRMADO').count(),
    }


def empresa_dashboard_autorizada(request, roles_permitidos=None):
    context = exigir_tenant(
        request,
        roles=roles_permitidos,
        allow_dirtec_switch=True,
    )

    if context.es_dirtec:
        empresas = EmpresaSuscriptora.objects.filter(
            activo=True
        ).order_by('nombre_comercial')
        return context.empresa or empresas.first(), empresas

    return (
        context.empresa,
        EmpresaSuscriptora.objects.filter(
            id=context.empresa.id
        ),
    )


def crear_evento_para_empresa(request, empresa):
    nombre_evento = limpiar_texto(request, 'nombre_evento') or 'Nuevo evento'
    principal = limpiar_texto(request, 'nombre_principal') or nombre_evento
    secundario = limpiar_texto(request, 'nombre_secundario') or ''
    fecha_evento = fecha_hora_dashboard(request, 'fecha_evento')
    sede = None
    sede_id = request.POST.get('sede_id')
    if sede_id:
        sede = SedeEvento.objects.filter(empresa=empresa, id=sede_id, activa=True).first()

    evento = EventoBoda.objects.create(
        empresa=empresa,
        sede=sede,
        nombre_evento=nombre_evento,
        tipo_evento=request.POST.get('tipo_evento') or 'BODA',
        estado='PLANEACION',
        novio=principal,
        novia=secundario or principal,
        nombre_principal=principal,
        nombre_secundario=secundario,
        mostrar_nombre_secundario=bool(secundario),
        frase_portada=limpiar_texto(request, 'frase_portada') or nombre_evento,
        mensaje_general='Estamos preparando todos los detalles de este evento.',
        fecha_misa=fecha_evento,
        lugar_misa=limpiar_texto(request, 'lugar_ceremonia') or (sede.nombre if sede else 'Ceremonia'),
        fecha_fiesta=fecha_evento,
        lugar_fiesta=limpiar_texto(request, 'lugar_recepcion') or (sede.nombre if sede else 'Recepcion'),
        capacidad_contratada=convertir_entero(request.POST.get('capacidad_contratada'), 0),
    )
    planner_id = request.POST.get('planner_id')
    if planner_id:
        membresia = MembresiaEmpresa.objects.filter(
            empresa=empresa,
            usuario_id=planner_id,
            rol='WEDDING_PLANNER',
            activo=True,
        ).select_related('usuario').first()
        if membresia:
            evento.wedding_planner = membresia.usuario
            evento.save(update_fields=['wedding_planner'])
    return evento


def crear_plan_suscripcion_dirtec(request):
    nombre = limpiar_texto(request, 'nombre_plan')
    if not nombre:
        return None
    plan = PlanSuscripcion.objects.create(
        nombre=nombre,
        descripcion=limpiar_texto(request, 'descripcion_plan') or '',
        precio_mensual=convertir_decimal(request.POST.get('precio_mensual_plan'), 0),
        precio_anual=convertir_decimal(request.POST.get('precio_anual_plan'), 0) or None,
        limite_usuarios=convertir_entero(request.POST.get('limite_usuarios_plan'), 5),
        limite_wedding_planners=convertir_entero(request.POST.get('limite_planners_plan'), 3),
        limite_eventos_activos=convertir_entero(request.POST.get('limite_eventos_plan'), 10),
        limite_clientes=convertir_entero(request.POST.get('limite_clientes_plan'), 50),
        limite_almacenamiento_mb=convertir_entero(request.POST.get('limite_almacenamiento_plan'), 1024),
        permite_proveedores=bool_post(request, 'permite_proveedores_plan'),
        permite_reportes=bool_post(request, 'permite_reportes_plan'),
        permite_api=bool_post(request, 'permite_api_plan'),
        permite_personalizacion=bool_post(request, 'permite_personalizacion_plan'),
        activo=bool_post(request, 'activo_plan'),
    )
    registrar_auditoria(
        usuario=request.user,
        accion='CREAR_PLAN_SUSCRIPCION',
        modelo='PlanSuscripcion',
        objeto_id=plan.id,
        descripcion=f'DIRTEC creo el plan {plan.nombre}.',
        request=request,
    )
    return plan


def actualizar_plan_suscripcion_dirtec(request):
    plan = get_object_or_404(PlanSuscripcion, id=request.POST.get('plan_id'))
    nombre = limpiar_texto(request, 'nombre_plan')
    if nombre:
        plan.nombre = nombre
    plan.descripcion = limpiar_texto(request, 'descripcion_plan') or ''
    plan.precio_mensual = convertir_decimal(request.POST.get('precio_mensual_plan'), plan.precio_mensual)
    plan.precio_anual = convertir_decimal(request.POST.get('precio_anual_plan'), 0) or None
    plan.limite_usuarios = convertir_entero(request.POST.get('limite_usuarios_plan'), plan.limite_usuarios)
    plan.limite_wedding_planners = convertir_entero(request.POST.get('limite_planners_plan'), plan.limite_wedding_planners)
    plan.limite_eventos_activos = convertir_entero(request.POST.get('limite_eventos_plan'), plan.limite_eventos_activos)
    plan.limite_clientes = convertir_entero(request.POST.get('limite_clientes_plan'), plan.limite_clientes)
    plan.limite_almacenamiento_mb = convertir_entero(request.POST.get('limite_almacenamiento_plan'), plan.limite_almacenamiento_mb)
    plan.permite_proveedores = bool_post(request, 'permite_proveedores_plan')
    plan.permite_reportes = bool_post(request, 'permite_reportes_plan')
    plan.permite_api = bool_post(request, 'permite_api_plan')
    plan.permite_personalizacion = bool_post(request, 'permite_personalizacion_plan')
    plan.activo = bool_post(request, 'activo_plan')
    plan.save()

    registrar_auditoria(
        usuario=request.user,
        accion='ACTUALIZAR_PLAN_SUSCRIPCION',
        modelo='PlanSuscripcion',
        objeto_id=plan.id,
        descripcion=f'DIRTEC actualizo el plan {plan.nombre}.',
        request=request,
    )
    messages.success(request, f'Plan actualizado: {plan.nombre}.')
    return plan


def eliminar_plan_suscripcion_dirtec(request):
    plan = get_object_or_404(PlanSuscripcion, id=request.POST.get('plan_id'))
    plan_id = plan.id
    nombre = plan.nombre
    try:
        plan.delete()
        messages.success(request, f'Plan eliminado: {nombre}.')
        accion = 'ELIMINAR_PLAN_SUSCRIPCION'
    except ProtectedError:
        plan.activo = False
        plan.save(update_fields=['activo', 'fecha_actualizacion'])
        messages.error(request, f'El plan {nombre} tiene suscripciones asociadas; se desactivo para no romper historial.')
        accion = 'DESACTIVAR_PLAN_SUSCRIPCION'

    registrar_auditoria(
        usuario=request.user,
        accion=accion,
        modelo='PlanSuscripcion',
        objeto_id=plan_id,
        descripcion=f'DIRTEC proceso baja del plan {nombre}.',
        request=request,
    )
    return plan_id


def crear_empresa_saas_dirtec(request):
    nombre = limpiar_texto(request, 'nombre_empresa')
    if not nombre:
        return None
    plan = get_object_or_404(PlanSuscripcion, id=request.POST.get('plan_suscripcion_id'), activo=True)
    fecha_inicio = fecha_dashboard(request, 'fecha_inicio_suscripcion')
    fecha_vencimiento = fecha_dashboard(request, 'fecha_vencimiento_suscripcion', fecha_inicio)
    estado_suscripcion = request.POST.get('estado_suscripcion') or 'ACTIVA'
    estados_validos = {valor for valor, _ in SuscripcionEmpresa.ESTADOS}
    if estado_suscripcion not in estados_validos:
        estado_suscripcion = 'ACTIVA'

    with transaction.atomic():
        empresa = EmpresaSuscriptora.objects.create(
            nombre_comercial=nombre,
            razon_social=limpiar_texto(request, 'razon_social'),
            slug=slug_empresa_unico(nombre),
            rfc=limpiar_texto(request, 'rfc_empresa') or '',
            contacto_nombre=limpiar_texto(request, 'contacto_nombre'),
            contacto_email=limpiar_texto(request, 'contacto_email'),
            contacto_telefono=limpiar_texto(request, 'contacto_telefono'),
            direccion=limpiar_texto(request, 'direccion_empresa') or '',
            estado='ACTIVA' if bool_post(request, 'empresa_activa') else 'INACTIVA',
            plan='PROFESIONAL',
            max_eventos_activos=plan.limite_eventos_activos,
            max_usuarios=plan.limite_usuarios,
            fecha_inicio=fecha_inicio,
            fecha_fin=fecha_vencimiento,
            activo=bool_post(request, 'empresa_activa'),
        )
        if request.FILES.get('logotipo_empresa'):
            empresa.logotipo = request.FILES['logotipo_empresa']
        colores_marca = {
            'primario': limpiar_texto(request, 'color_primario_empresa'),
            'secundario': limpiar_texto(request, 'color_secundario_empresa'),
        }
        empresa.colores_marca = {clave: valor for clave, valor in colores_marca.items() if valor}
        if empresa.logotipo or empresa.colores_marca:
            empresa.save(update_fields=['logotipo', 'colores_marca', 'fecha_actualizacion'])
        suscripcion = SuscripcionEmpresa.objects.create(
            empresa=empresa,
            plan=plan,
            fecha_inicio=fecha_inicio,
            fecha_vencimiento=fecha_vencimiento,
            fecha_periodo_gracia=fecha_dashboard(request, 'fecha_periodo_gracia_suscripcion', None) if request.POST.get('fecha_periodo_gracia_suscripcion') else None,
            estado=estado_suscripcion,
            renovacion_automatica=bool_post(request, 'renovacion_automatica'),
            observaciones=limpiar_texto(request, 'observaciones_suscripcion') or '',
        )
        monto_inicial = convertir_decimal(request.POST.get('monto_pago_inicial'), 0)
        if monto_inicial > 0:
            PagoSuscripcion.objects.create(
                empresa=empresa,
                suscripcion=suscripcion,
                monto=monto_inicial,
                fecha_vencimiento=fecha_dashboard(request, 'fecha_vencimiento_pago', fecha_vencimiento),
                fecha_pago=fecha_dashboard(request, 'fecha_pago_inicial', None) if request.POST.get('fecha_pago_inicial') else None,
                metodo_pago=limpiar_texto(request, 'metodo_pago_inicial') or '',
                referencia=limpiar_texto(request, 'referencia_pago_inicial') or '',
                estado=request.POST.get('estado_pago_inicial') or 'PENDIENTE',
                registrado_por=request.user,
                notas=limpiar_texto(request, 'notas_pago_inicial') or '',
            )
        if limpiar_texto(request, 'username_usuario'):
            crear_o_actualizar_usuario_empresa(request, empresa, 'ADMIN_EMPRESA')

        registrar_auditoria(
            usuario=request.user,
            empresa=empresa,
            accion='CREAR_EMPRESA_SAAS',
            modelo='EmpresaSuscriptora',
            objeto_id=empresa.id,
            descripcion=f'DIRTEC creo la empresa {empresa.nombre_comercial} con suscripcion {suscripcion.id}.',
            valores_nuevos={'plan': plan.nombre, 'estado_suscripcion': suscripcion.estado},
            request=request,
        )
    return empresa


def actualizar_empresa_dirtec(request):
    empresa = get_object_or_404(EmpresaSuscriptora, id=request.POST.get('empresa_id'))
    estados_validos = {valor for valor, _ in EmpresaSuscriptora.ESTADOS}
    planes_validos = {valor for valor, _ in EmpresaSuscriptora.PLANES}
    estado = request.POST.get('estado_empresa')
    plan = request.POST.get('plan_empresa')

    empresa.nombre_comercial = limpiar_texto(request, 'nombre_empresa') or empresa.nombre_comercial
    empresa.razon_social = limpiar_texto(request, 'razon_social')
    empresa.rfc = limpiar_texto(request, 'rfc_empresa') or ''
    empresa.contacto_nombre = limpiar_texto(request, 'contacto_nombre')
    empresa.contacto_email = limpiar_texto(request, 'contacto_email')
    empresa.contacto_telefono = limpiar_texto(request, 'contacto_telefono')
    empresa.direccion = limpiar_texto(request, 'direccion_empresa') or ''
    empresa.notas = limpiar_texto(request, 'notas_empresa')
    if request.FILES.get('logotipo_empresa'):
        empresa.logotipo = request.FILES['logotipo_empresa']
    colores_marca = {
        'primario': limpiar_texto(request, 'color_primario_empresa'),
        'secundario': limpiar_texto(request, 'color_secundario_empresa'),
    }
    empresa.colores_marca = {clave: valor for clave, valor in colores_marca.items() if valor}
    empresa.max_eventos_activos = convertir_entero(request.POST.get('max_eventos_activos'), empresa.max_eventos_activos)
    empresa.max_usuarios = convertir_entero(request.POST.get('max_usuarios'), empresa.max_usuarios)
    empresa.fecha_inicio = fecha_dashboard(request, 'fecha_inicio_empresa', None) if request.POST.get('fecha_inicio_empresa') else None
    empresa.fecha_fin = fecha_dashboard(request, 'fecha_fin_empresa', None) if request.POST.get('fecha_fin_empresa') else None
    if estado in estados_validos:
        empresa.estado = estado
    if plan in planes_validos:
        empresa.plan = plan
    empresa.activo = bool_post(request, 'empresa_activa')
    empresa.save()

    registrar_auditoria(
        usuario=request.user,
        empresa=empresa,
        accion='ACTUALIZAR_EMPRESA_DIRTEC',
        modelo='EmpresaSuscriptora',
        objeto_id=empresa.id,
        descripcion=f'DIRTEC actualizo la empresa {empresa.nombre_comercial}.',
        request=request,
    )
    return empresa


def desactivar_empresa_dirtec(request):
    empresa = get_object_or_404(EmpresaSuscriptora, id=request.POST.get('empresa_id'))
    empresa.activo = False
    empresa.estado = 'INACTIVA'
    empresa.save(update_fields=['activo', 'estado', 'fecha_actualizacion'])
    MembresiaEmpresa.objects.filter(empresa=empresa).update(activo=False)
    registrar_auditoria(
        usuario=request.user,
        empresa=empresa,
        accion='DESACTIVAR_EMPRESA_DIRTEC',
        modelo='EmpresaSuscriptora',
        objeto_id=empresa.id,
        descripcion=f'DIRTEC desactivo la empresa {empresa.nombre_comercial}.',
        request=request,
    )
    return empresa


def eliminar_empresa_dirtec(request):
    empresa = get_object_or_404(EmpresaSuscriptora, id=request.POST.get('empresa_id'))
    empresa_id = empresa.id
    nombre = empresa.nombre_comercial
    registrar_auditoria(
        usuario=request.user,
        empresa=empresa,
        accion='ELIMINAR_EMPRESA_DIRTEC',
        modelo='EmpresaSuscriptora',
        objeto_id=empresa.id,
        descripcion=f'DIRTEC elimino definitivamente la empresa {nombre}.',
        request=request,
    )

    with transaction.atomic():
        EventoBoda.objects.filter(empresa=empresa).delete()
        Proveedor.objects.filter(empresa=empresa).delete()
        for paquete in PaqueteBoda.objects.filter(empresa=empresa):
            paquete.delete()
        empresa.delete()

    messages.success(request, f'Empresa eliminada: {nombre}.')
    return empresa_id


def actualizar_suscripcion_dirtec(request):
    suscripcion = get_object_or_404(SuscripcionEmpresa.objects.select_related('empresa'), id=request.POST.get('suscripcion_id'))
    plan = get_object_or_404(PlanSuscripcion, id=request.POST.get('plan_suscripcion_id'), activo=True)
    estados_validos = {valor for valor, _ in SuscripcionEmpresa.ESTADOS}
    estado = request.POST.get('estado_suscripcion') if request.POST.get('estado_suscripcion') in estados_validos else suscripcion.estado
    anteriores = {
        'plan': suscripcion.plan_id,
        'estado': suscripcion.estado,
        'fecha_vencimiento': str(suscripcion.fecha_vencimiento),
        'bloqueada_manualmente': suscripcion.bloqueada_manualmente,
    }

    suscripcion.plan = plan
    suscripcion.fecha_inicio = fecha_dashboard(request, 'fecha_inicio_suscripcion', suscripcion.fecha_inicio)
    suscripcion.fecha_vencimiento = fecha_dashboard(request, 'fecha_vencimiento_suscripcion', suscripcion.fecha_vencimiento)
    suscripcion.fecha_periodo_gracia = fecha_dashboard(request, 'fecha_periodo_gracia_suscripcion', None) if request.POST.get('fecha_periodo_gracia_suscripcion') else None
    suscripcion.estado = estado
    suscripcion.renovacion_automatica = bool_post(request, 'renovacion_automatica')
    suscripcion.bloqueada_manualmente = bool_post(request, 'bloqueada_manualmente')
    suscripcion.motivo_bloqueo = limpiar_texto(request, 'motivo_bloqueo') or ''
    suscripcion.observaciones = limpiar_texto(request, 'observaciones_suscripcion') or ''
    suscripcion.save()

    empresa = suscripcion.empresa
    empresa.estado = 'BLOQUEADA' if suscripcion.bloqueada_manualmente else ('SUSPENDIDA' if estado == 'SUSPENDIDA' else 'ACTIVA')
    empresa.activo = estado not in {'SUSPENDIDA', 'CANCELADA'} and not suscripcion.bloqueada_manualmente
    empresa.fecha_inicio = suscripcion.fecha_inicio
    empresa.fecha_fin = suscripcion.fecha_vencimiento
    empresa.max_eventos_activos = plan.limite_eventos_activos
    empresa.max_usuarios = plan.limite_usuarios
    empresa.save(update_fields=['estado', 'activo', 'fecha_inicio', 'fecha_fin', 'max_eventos_activos', 'max_usuarios', 'fecha_actualizacion'])

    registrar_auditoria(
        usuario=request.user,
        empresa=empresa,
        accion='ACTUALIZAR_SUSCRIPCION',
        modelo='SuscripcionEmpresa',
        objeto_id=suscripcion.id,
        descripcion=f'DIRTEC actualizo la suscripcion de {empresa.nombre_comercial}.',
        valores_anteriores=anteriores,
        valores_nuevos={
            'plan': plan.id,
            'estado': suscripcion.estado,
            'fecha_vencimiento': str(suscripcion.fecha_vencimiento),
            'bloqueada_manualmente': suscripcion.bloqueada_manualmente,
        },
        request=request,
    )
    return suscripcion


def registrar_pago_suscripcion_dirtec(request):
    suscripcion = get_object_or_404(SuscripcionEmpresa.objects.select_related('empresa'), id=request.POST.get('suscripcion_id'))
    estado = request.POST.get('estado_pago') or 'PENDIENTE'
    estados_validos = {valor for valor, _ in PagoSuscripcion.ESTADOS}
    if estado not in estados_validos:
        estado = 'PENDIENTE'
    pago = PagoSuscripcion.objects.create(
        empresa=suscripcion.empresa,
        suscripcion=suscripcion,
        monto=convertir_decimal(request.POST.get('monto_pago'), 0),
        fecha_vencimiento=fecha_dashboard(request, 'fecha_vencimiento_pago', timezone.localdate()),
        fecha_pago=fecha_dashboard(request, 'fecha_pago', None) if request.POST.get('fecha_pago') else None,
        metodo_pago=limpiar_texto(request, 'metodo_pago') or '',
        referencia=limpiar_texto(request, 'referencia_pago') or '',
        estado=estado,
        registrado_por=request.user,
        notas=limpiar_texto(request, 'notas_pago') or '',
    )
    if pago.estado == 'PAGADO':
        suscripcion.estado = 'ACTIVA'
        suscripcion.bloqueada_manualmente = False
        suscripcion.motivo_bloqueo = ''
        suscripcion.save(update_fields=['estado', 'bloqueada_manualmente', 'motivo_bloqueo', 'fecha_actualizacion'])
        suscripcion.empresa.estado = 'ACTIVA'
        suscripcion.empresa.activo = True
        suscripcion.empresa.save(update_fields=['estado', 'activo', 'fecha_actualizacion'])
    registrar_auditoria(
        usuario=request.user,
        empresa=suscripcion.empresa,
        accion='REGISTRAR_PAGO_SUSCRIPCION',
        modelo='PagoSuscripcion',
        objeto_id=pago.id,
        descripcion=f'DIRTEC registro pago de suscripcion por {pago.monto}.',
        request=request,
    )
    return pago


def planner_activo_empresa(empresa, planner_id):
    if not planner_id:
        return None
    membresia = MembresiaEmpresa.objects.filter(
        empresa=empresa,
        usuario_id=planner_id,
        rol='WEDDING_PLANNER',
        activo=True,
    ).select_related('usuario').first()
    return membresia.usuario if membresia else None


def actualizar_evento_empresa_dashboard(request, empresa):
    evento = get_object_or_404(EventoBoda, empresa=empresa, id=request.POST.get('evento_id'))
    tipos_validos = {valor for valor, _ in EventoBoda.TIPOS_EVENTO}
    estados_validos = {valor for valor, _ in EventoBoda.ESTADOS_EVENTO}
    tipo_evento = request.POST.get('tipo_evento')
    estado = request.POST.get('estado')
    sede_id = request.POST.get('sede_id')
    principal = limpiar_texto(request, 'nombre_principal') or evento.nombre_principal or evento.novio
    secundario = limpiar_texto(request, 'nombre_secundario') or ''

    evento.nombre_evento = limpiar_texto(request, 'nombre_evento') or evento.nombre_evento
    if tipo_evento in tipos_validos:
        evento.tipo_evento = tipo_evento
    if estado in estados_validos:
        evento.estado = estado
    evento.nombre_principal = principal
    evento.nombre_secundario = secundario
    evento.novio = principal
    evento.novia = secundario or principal
    evento.mostrar_nombre_secundario = bool(secundario)
    evento.fecha_fiesta = fecha_hora_dashboard(request, 'fecha_evento', evento.fecha_fiesta)
    evento.fecha_misa = fecha_hora_dashboard(request, 'fecha_ceremonia', evento.fecha_misa)
    evento.lugar_misa = limpiar_texto(request, 'lugar_ceremonia') or evento.lugar_misa
    evento.lugar_fiesta = limpiar_texto(request, 'lugar_recepcion') or evento.lugar_fiesta
    evento.capacidad_contratada = convertir_entero(request.POST.get('capacidad_contratada'), evento.capacidad_contratada)
    evento.frase_portada = limpiar_texto(request, 'frase_portada') or evento.frase_portada
    evento.activo = bool_post(request, 'activo_evento')

    evento.sede = SedeEvento.objects.filter(empresa=empresa, id=sede_id).first() if sede_id else None
    evento.wedding_planner = planner_activo_empresa(empresa, request.POST.get('planner_id'))
    evento.save()
    return evento


def crear_cliente_empresa_dashboard(request, empresa):
    user = crear_o_actualizar_usuario_empresa(request, empresa, 'CLIENTE')
    evento_id = request.POST.get('evento_cliente_id')
    if user and evento_id:
        evento = get_object_or_404(EventoBoda, empresa=empresa, id=evento_id)
        evento.clientes.add(user)
    if user:
        registrar_auditoria(
            usuario=request.user,
            empresa=empresa,
            accion='CREAR_CLIENTE_EMPRESA',
            modelo='auth.User',
            objeto_id=user.id,
            descripcion=f'Se creo o actualizo cliente {user.username} desde dashboard de empresa.',
            request=request,
        )
    return user


def crear_evento_planner_dashboard(request, empresa):
    post = request.POST.copy()
    post['planner_id'] = str(request.user.id)
    request.POST = post
    evento = crear_evento_para_empresa(request, empresa)
    if evento.wedding_planner_id != request.user.id:
        evento.wedding_planner = request.user
        evento.save(update_fields=['wedding_planner'])
    registrar_auditoria(
        usuario=request.user,
        empresa=empresa,
        evento=evento,
        accion='CREAR_EVENTO_PLANNER',
        modelo='EventoBoda',
        objeto_id=evento.id,
        descripcion=f'Wedding planner creo evento {evento}.',
        request=request,
    )
    return evento


def actualizar_sede_empresa_dashboard(request, empresa):
    sede = get_object_or_404(SedeEvento, empresa=empresa, id=request.POST.get('sede_id'))
    tipos_validos = {valor for valor, _ in SedeEvento.TIPOS}
    tipo = request.POST.get('tipo_sede')
    sede.nombre = limpiar_texto(request, 'nombre_sede') or sede.nombre
    if tipo in tipos_validos:
        sede.tipo = tipo
    sede.capacidad_minima = convertir_entero(request.POST.get('capacidad_minima'), sede.capacidad_minima)
    sede.capacidad_maxima = convertir_entero(request.POST.get('capacidad_maxima'), sede.capacidad_maxima)
    sede.precio_base = convertir_decimal(request.POST.get('precio_base_sede'), sede.precio_base)
    sede.direccion = limpiar_texto(request, 'direccion_sede')
    sede.descripcion = limpiar_texto(request, 'descripcion_sede')
    sede.activa = bool_post(request, 'activa_sede')
    sede.save()
    return sede


def sincronizar_acceso_proveedor(
    request,
    empresa,
    proveedor,
):
    username = limpiar_texto(
        request,
        'username_usuario',
    )

    if not username:
        return proveedor

    user = proveedor.usuario
    if user is None:
        password = (
            request.POST.get(
                'password_usuario'
            )
            or get_random_string(12)
        )
        try:
            user = crear_identidad_usuario(
                username=username,
                email=limpiar_texto(
                    request,
                    'email_usuario',
                ),
                phone=limpiar_texto(
                    request,
                    'telefono_usuario',
                ),
                first_name=(
                    limpiar_texto(
                        request,
                        'first_name_usuario',
                    )
                    or proveedor.nombre_contacto
                    or proveedor.nombre_comercial
                ),
                last_name=limpiar_texto(
                    request,
                    'last_name_usuario',
                ),
                password=password,
                active=True,
            )
        except ValidationError as exc:
            messages.error(
                request,
                '; '.join(exc.messages),
            )
            return proveedor

        proveedor.usuario = user
        proveedor.save(
            update_fields=['usuario']
        )
    else:
        try:
            actualizar_identidad_usuario(
                user,
                username=username,
                email=limpiar_texto(
                    request,
                    'email_usuario',
                ),
                phone=limpiar_texto(
                    request,
                    'telefono_usuario',
                ),
                first_name=(
                    limpiar_texto(
                        request,
                        'first_name_usuario',
                    )
                    or user.first_name
                ),
                last_name=(
                    limpiar_texto(
                        request,
                        'last_name_usuario',
                    )
                    or user.last_name
                ),
                password=(
                    request.POST.get(
                        'password_usuario'
                    )
                    or None
                ),
                active=proveedor.activo,
            )
        except ValidationError as exc:
            messages.error(
                request,
                '; '.join(exc.messages),
            )
            return proveedor

    MembresiaEmpresa.objects.update_or_create(
        empresa=empresa,
        usuario=proveedor.usuario,
        rol='PROVEEDOR',
        defaults={
            'activo': proveedor.activo,
            'puede_gestionar_catalogos': False,
        },
    )
    return proveedor


def actualizar_proveedor_empresa_dashboard(request, empresa):
    proveedor = get_object_or_404(Proveedor, empresa=empresa, id=request.POST.get('proveedor_id'))
    tipos_validos = {valor for valor, _ in Proveedor.TIPOS}
    tipo = request.POST.get('tipo_proveedor')
    proveedor.nombre_comercial = limpiar_texto(request, 'nombre_proveedor') or proveedor.nombre_comercial
    if tipo in tipos_validos:
        proveedor.tipo_proveedor = tipo
    proveedor.nombre_contacto = limpiar_texto(request, 'contacto_proveedor')
    proveedor.telefono = limpiar_texto(request, 'telefono_proveedor')
    proveedor.correo = limpiar_texto(request, 'correo_proveedor')
    proveedor.contacto_operativo = limpiar_texto(request, 'contacto_operativo_proveedor')
    proveedor.telefono_operativo = limpiar_texto(request, 'telefono_operativo_proveedor')
    proveedor.correo_operativo = limpiar_texto(request, 'correo_operativo_proveedor')
    proveedor.visible_para_wedding_planners = bool_post(request, 'visible_wedding_planners_proveedor')
    proveedor.rfc = limpiar_texto(request, 'rfc_proveedor')
    proveedor.datos_bancarios = limpiar_texto(request, 'datos_bancarios_proveedor')
    proveedor.notas_privadas = limpiar_texto(request, 'notas_privadas_proveedor')
    proveedor.descripcion = limpiar_texto(request, 'descripcion_proveedor')
    proveedor.activo = bool_post(request, 'activo_proveedor')
    proveedor.save()
    sincronizar_acceso_proveedor(
        request,
        empresa,
        proveedor,
    )
    return proveedor


def actualizar_paquete_empresa_dashboard(request, empresa):
    paquete = get_object_or_404(PaqueteBoda, empresa=empresa, id=request.POST.get('paquete_id'))
    paquete.nombre = limpiar_texto(request, 'nombre_paquete') or paquete.nombre
    paquete.descripcion = limpiar_texto(request, 'descripcion_paquete')
    paquete.numero_personas_incluidas = convertir_entero(request.POST.get('personas_paquete'), paquete.numero_personas_incluidas)
    paquete.precio_base = convertir_decimal(request.POST.get('precio_base_paquete'), paquete.precio_base)
    paquete.activo = bool_post(request, 'activo_paquete')
    paquete.save()
    return paquete


def eliminar_evento_empresa_dashboard(request, empresa):
    evento = get_object_or_404(EventoBoda, empresa=empresa, id=request.POST.get('evento_id'))
    nombre = str(evento)
    evento.delete()
    registrar_auditoria(
        usuario=request.user,
        empresa=empresa,
        accion='ELIMINAR_EVENTO_EMPRESA',
        modelo='EventoBoda',
        objeto_id=request.POST.get('evento_id'),
        descripcion=f'Se elimino definitivamente el evento {nombre}.',
        request=request,
    )
    messages.success(request, f'Evento eliminado: {nombre}.')


def eliminar_sede_empresa_dashboard(request, empresa):
    sede = get_object_or_404(SedeEvento, empresa=empresa, id=request.POST.get('sede_id'))
    nombre = sede.nombre
    sede.delete()
    registrar_auditoria(
        usuario=request.user,
        empresa=empresa,
        accion='ELIMINAR_SEDE_EMPRESA',
        modelo='SedeEvento',
        objeto_id=request.POST.get('sede_id'),
        descripcion=f'Se elimino definitivamente la sede {nombre}.',
        request=request,
    )
    messages.success(request, f'Sede eliminada: {nombre}.')


def eliminar_proveedor_empresa_dashboard(request, empresa):
    proveedor = get_object_or_404(Proveedor, empresa=empresa, id=request.POST.get('proveedor_id'))
    nombre = proveedor.nombre_comercial
    proveedor.delete()
    registrar_auditoria(
        usuario=request.user,
        empresa=empresa,
        accion='ELIMINAR_PROVEEDOR_EMPRESA',
        modelo='Proveedor',
        objeto_id=request.POST.get('proveedor_id'),
        descripcion=f'Se elimino definitivamente el proveedor {nombre}.',
        request=request,
    )
    messages.success(request, f'Proveedor eliminado: {nombre}.')


def eliminar_paquete_empresa_dashboard(request, empresa):
    paquete = get_object_or_404(PaqueteBoda, empresa=empresa, id=request.POST.get('paquete_id'))
    nombre = paquete.nombre
    try:
        paquete.delete()
    except ProtectedError:
        messages.error(request, f'No se pudo eliminar el paquete {nombre} porque tiene relaciones protegidas.')
        return
    registrar_auditoria(
        usuario=request.user,
        empresa=empresa,
        accion='ELIMINAR_PAQUETE_EMPRESA',
        modelo='PaqueteBoda',
        objeto_id=request.POST.get('paquete_id'),
        descripcion=f'Se elimino definitivamente el paquete {nombre}.',
        request=request,
    )
    messages.success(request, f'Paquete eliminado: {nombre}.')


def empresa_actual_dashboard(request):
    context = tenant_desde_request(
        request,
        allow_dirtec_switch=True,
    )

    if context.es_dirtec:
        empresas = EmpresaSuscriptora.objects.filter(
            activo=True
        ).order_by('nombre_comercial')
        return context.empresa or empresas.first(), empresas

    if not context.empresa:
        return None, EmpresaSuscriptora.objects.none()

    return (
        context.empresa,
        EmpresaSuscriptora.objects.filter(
            id=context.empresa.id
        ),
    )


def empresa_autorizada_para_evento(request, evento, empresa_id):
    if not empresa_id:
        return evento.empresa
    if usuario_es_dirtec(request.user):
        return get_object_or_404(EmpresaSuscriptora, id=empresa_id)

    if evento.empresa_id and str(evento.empresa_id) != str(empresa_id):
        registrar_auditoria(
            usuario=request.user,
            empresa=evento.empresa,
            evento=evento,
            accion='INTENTO_CAMBIO_EMPRESA_EVENTO',
            modelo='EventoBoda',
            objeto_id=evento.id,
            descripcion=f'Intento de reasignar evento a empresa {empresa_id} sin permiso DIRTEC.',
            request=request,
        )
        raise PermissionDenied('No puedes cambiar la empresa de este evento.')
    empresa = get_object_or_404(empresas_para_usuario(request.user), id=empresa_id)
    return empresa


def membresia_empresa_usuario(user, empresa):
    if not user.is_authenticated or not empresa:
        return None
    return MembresiaEmpresa.objects.filter(usuario=user, empresa=empresa, activo=True).order_by('rol').first()


def eventos_dashboard_queryset(request):
    empresa, _ = empresa_actual_dashboard(request)
    eventos = eventos_visibles_usuario(request.user)
    if empresa:
        eventos = eventos.filter(empresa=empresa)
    return eventos


def obtener_evento_dashboard(request):
    evento_id = request.GET.get('evento')
    eventos = eventos_dashboard_queryset(request)

    if evento_id:
        evento = get_object_or_404(eventos, id=evento_id)
    else:
        evento = eventos.filter(activo=True).first() or eventos.first()

    return evento, eventos


def exigir_permiso_evento(request, evento, action):
    if not evento or not usuario_puede_evento(request.user, evento, action):
        raise PermissionDenied('No tienes permiso para realizar esta acción en el evento.')
    return evento

def eventos_para_cliente(user):
    if not getattr(user, 'is_authenticated', False):
        return EventoBoda.objects.none()
    if usuario_es_dirtec(user):
        return EventoBoda.objects.none()
    return (
        EventoBoda.objects
        .filter(clientes=user)
        .filter(
            empresa__membresias__usuario=user,
            empresa__membresias__activo=True,
            empresa__membresias__rol='CLIENTE',
        )
        .distinct()
        .order_by('-activo', '-fecha_fiesta')
    )


def obtener_evento_portal(request, eventos):
    evento_id = request.GET.get('evento')
    if evento_id:
        return get_object_or_404(eventos, id=evento_id)
    return eventos.filter(activo=True).first() or eventos.first()


def usuario_puede_ver_evento_cliente(user, evento):
    if not getattr(user, 'is_authenticated', False) or not evento:
        return False
    return eventos_para_cliente(user).filter(id=evento.id).exists()


def proveedor_de_usuario(user):
    if not user.is_authenticated:
        return None
    return Proveedor.objects.filter(usuario=user, activo=True).first()


def usuarios_notificacion_evento(evento, incluir_clientes=False, usuarios_extra=None):
    usuarios = []
    usuarios_extra = usuarios_extra or []

    if evento and evento.wedding_planner:
        usuarios.append(evento.wedding_planner)

    if incluir_clientes and evento:
        usuarios.extend(list(evento.clientes.all()))

    usuarios.extend([usuario for usuario in usuarios_extra if usuario])

    unicos = {}
    for usuario in usuarios:
        if usuario and usuario.pk:
            unicos[usuario.pk] = usuario
    return unicos.values()


def crear_notificaciones_evento(evento, titulo, mensaje, tipo='INFO', enlace=None, incluir_clientes=False, usuarios_extra=None):
    for usuario in usuarios_notificacion_evento(evento, incluir_clientes=incluir_clientes, usuarios_extra=usuarios_extra):
        Notificacion.objects.create(
            usuario=usuario,
            evento=evento,
            titulo=titulo,
            mensaje=mensaje,
            tipo=tipo,
            enlace=enlace,
        )


@login_required(login_url=LOGIN_DASHBOARD_URL)
def dashboard_profesional(request):
    if usuario_es_dirtec(request.user):
        return redirect('/dirtec/dashboard/')

    empresas = empresas_para_usuario(request.user)
    for empresa in empresas:
        roles = roles_activos_empresa(request.user, empresa)
        if roles.intersection({'ADMIN_EMPRESA', 'VENTAS'}):
            return redirect(f'/empresa/{empresa.slug}/dashboard/')
        if 'WEDDING_PLANNER' in roles:
            return redirect(f'/empresa/{empresa.slug}/planner/dashboard/')

    return redirect('dashboard')


@login_required(login_url=LOGIN_DASHBOARD_URL)
def dashboard_dirtec(request):
    if not usuario_es_dirtec(request.user):
        return redirect('dashboard_profesional')

    if request.method == 'POST':
        accion = request.POST.get('accion')
        if accion == 'crear_plan_suscripcion':
            crear_plan_suscripcion_dirtec(request)
            return redirect('/dirtec/dashboard/#planes')
        if accion == 'editar_plan_suscripcion':
            actualizar_plan_suscripcion_dirtec(request)
            return redirect('/dirtec/dashboard/#planes')
        if accion == 'eliminar_plan_suscripcion':
            eliminar_plan_suscripcion_dirtec(request)
            return redirect('/dirtec/dashboard/#planes')
        if accion == 'crear_empresa':
            empresa = crear_empresa_saas_dirtec(request)
            return redirect(f'/dirtec/dashboard/?empresa={empresa.id}#empresas' if empresa else '/dirtec/dashboard/#empresas')
        if accion == 'editar_empresa':
            empresa = actualizar_empresa_dirtec(request)
            return redirect(f'/dirtec/dashboard/?empresa={empresa.id}#empresas')
        if accion == 'desactivar_empresa':
            empresa = desactivar_empresa_dirtec(request)
            return redirect(f'/dirtec/dashboard/?empresa={empresa.id}#empresas')
        if accion == 'eliminar_empresa':
            eliminar_empresa_dirtec(request)
            return redirect('/dirtec/dashboard/#empresas')
        if accion == 'crear_admin_empresa':
            empresa = get_object_or_404(EmpresaSuscriptora, id=request.POST.get('empresa_id'))
            crear_o_actualizar_usuario_empresa(request, empresa, 'ADMIN_EMPRESA')
            return redirect(f'/dirtec/dashboard/?empresa={empresa.id}#usuarios')
        if accion == 'editar_usuario_empresa':
            empresa = get_object_or_404(EmpresaSuscriptora, id=request.POST.get('empresa_id'))
            actualizar_usuario_empresa_dashboard(request, empresa)
            return redirect(f'/dirtec/dashboard/?empresa={empresa.id}#usuarios')
        if accion == 'desactivar_usuario_empresa':
            empresa = get_object_or_404(EmpresaSuscriptora, id=request.POST.get('empresa_id'))
            desactivar_usuario_empresa_dashboard(request, empresa)
            return redirect(f'/dirtec/dashboard/?empresa={empresa.id}#usuarios')
        if accion == 'eliminar_usuario_empresa':
            empresa = get_object_or_404(EmpresaSuscriptora, id=request.POST.get('empresa_id'))
            eliminar_usuario_empresa_dashboard(request, empresa)
            return redirect(f'/dirtec/dashboard/?empresa={empresa.id}#usuarios')
        if accion == 'actualizar_suscripcion':
            suscripcion = actualizar_suscripcion_dirtec(request)
            return redirect(f'/dirtec/dashboard/?empresa={suscripcion.empresa_id}#suscripciones')
        if accion == 'registrar_pago_suscripcion':
            pago = registrar_pago_suscripcion_dirtec(request)
            return redirect(f'/dirtec/dashboard/?empresa={pago.empresa_id}#pagos')

    empresas = EmpresaSuscriptora.objects.all().order_by('nombre_comercial')
    empresa_actual = None
    empresa_id = request.GET.get('empresa')
    if empresa_id:
        empresa_actual = get_object_or_404(empresas, id=empresa_id)
    else:
        empresa_actual = empresas.first()

    eventos = EventoBoda.objects.all()
    membresias = MembresiaEmpresa.objects.select_related('empresa', 'usuario')
    planes_suscripcion = PlanSuscripcion.objects.all().order_by('-activo', 'nombre')
    suscripciones = SuscripcionEmpresa.objects.select_related('empresa', 'plan').order_by('fecha_vencimiento')
    pagos_suscripcion = PagoSuscripcion.objects.select_related('empresa', 'suscripcion', 'registrado_por').order_by('-fecha_registro')
    hoy = timezone.localdate()
    empresas_detalle = []
    for empresa in empresas:
        eventos_empresa = eventos.filter(empresa=empresa)
        suscripcion = getattr(empresa, 'suscripcion', None)
        empresas_detalle.append({
            'empresa': empresa,
            'suscripcion': suscripcion,
            'eventos': eventos_empresa.count(),
            'usuarios': membresias.filter(empresa=empresa, activo=True).values('usuario_id').distinct().count(),
            'planners': membresias.filter(empresa=empresa, rol='WEDDING_PLANNER', activo=True).count(),
            'pendientes': resumen_operativo_eventos(eventos_empresa)['tareas_pendientes'],
        })

    context = {
        'titulo_dashboard': 'DIRTEC Control',
        'empresas': empresas,
        'empresa_actual': empresa_actual,
        'empresas_detalle': empresas_detalle,
        'planes': EmpresaSuscriptora.PLANES,
        'estados_empresa': EmpresaSuscriptora.ESTADOS,
        'planes_suscripcion': planes_suscripcion,
        'suscripciones': suscripciones,
        'suscripcion_actual': getattr(empresa_actual, 'suscripcion', None) if empresa_actual else None,
        'pagos_suscripcion': pagos_suscripcion.filter(empresa=empresa_actual)[:12] if empresa_actual else [],
        'estados_suscripcion': SuscripcionEmpresa.ESTADOS,
        'estados_pago_suscripcion': PagoSuscripcion.ESTADOS,
        'estados_evento': EventoBoda.ESTADOS_EVENTO,
        'roles_empresa': MembresiaEmpresa.ROLES,
        'usuarios_empresa': membresias.filter(empresa=empresa_actual).order_by('-activo', 'rol') if empresa_actual else [],
        'eventos_recientes': eventos.select_related('empresa', 'sede', 'wedding_planner').order_by('-fecha_creacion')[:10],
        'resumen': {
            'empresas': empresas.count(),
            'empresas_activas': empresas.filter(activo=True).count(),
            'empresas_suspendidas': empresas.filter(estado='SUSPENDIDA').count(),
            'empresas_bloqueadas': empresas.filter(estado='BLOQUEADA').count(),
            'eventos': eventos.count(),
            'usuarios': membresias.values('usuario_id').distinct().count(),
            'suscripciones_vigentes': suscripciones.filter(estado__in=['PRUEBA', 'ACTIVA', 'PROXIMA_A_VENCER'], fecha_vencimiento__gte=hoy).count(),
            'suscripciones_por_vencer': suscripciones.filter(fecha_vencimiento__gte=hoy, fecha_vencimiento__lte=hoy + timedelta(days=7)).count(),
            'suscripciones_vencidas': suscripciones.filter(fecha_vencimiento__lt=hoy).count(),
            'pagos_pendientes': pagos_suscripcion.filter(estado='PENDIENTE').count(),
        },
    }
    return render(request, 'invitaciones/dashboard_dirtec.html', context)


@login_required(login_url=LOGIN_DASHBOARD_URL)
def dashboard_empresa(request):
    empresa, empresas = empresa_dashboard_autorizada(request, ['ADMIN_EMPRESA', 'VENTAS'])
    if not empresa:
        return redirect('dashboard_profesional')

    puede_catalogos = usuario_puede_gestionar_catalogos(request.user, empresa)
    puede_usuarios = usuario_puede_gestionar_usuarios(request.user, empresa)
    roles_usuario_actual = roles_activos_empresa(request.user, empresa)

    if request.method == 'POST':
        accion = request.POST.get('accion')
        if accion in ACCIONES_PANEL_EMPRESA_USUARIOS and not puede_usuarios:
            bloquear_accion_dashboard(request, empresa=empresa, accion=accion, permiso='gestionar usuarios')
        if accion in ACCIONES_PANEL_EMPRESA_CATALOGOS and not puede_catalogos:
            bloquear_accion_dashboard(request, empresa=empresa, accion=accion, permiso='gestionar catalogos')
        if accion == 'crear_evento':
            try:
                validar_limite_eventos_activos(empresa)
            except LimitePlanExcedido as exc:
                return bloquear_por_limite_plan(request, empresa, str(exc), '#eventos')
            evento = crear_evento_para_empresa(request, empresa)
            return redirect(f'/dashboard/empresa/?empresa={empresa.id}#eventos')
        if accion == 'editar_evento':
            actualizar_evento_empresa_dashboard(request, empresa)
            return redirect(f'/dashboard/empresa/?empresa={empresa.id}#eventos')
        if accion == 'desactivar_evento':
            evento = get_object_or_404(EventoBoda, empresa=empresa, id=request.POST.get('evento_id'))
            evento.activo = False
            evento.save(update_fields=['activo'])
            return redirect(f'/dashboard/empresa/?empresa={empresa.id}#eventos')
        if accion == 'eliminar_evento':
            eliminar_evento_empresa_dashboard(request, empresa)
            return redirect(f'/dashboard/empresa/?empresa={empresa.id}#eventos')
        if accion == 'crear_planner' and puede_usuarios:
            username = limpiar_texto(request, 'username_usuario')
            if not usuario_empresa_existente_activo(empresa, username):
                try:
                    validar_limite_usuarios(empresa)
                except LimitePlanExcedido as exc:
                    return bloquear_por_limite_plan(request, empresa, str(exc), '#equipo')
            if not usuario_empresa_existente_activo(empresa, username, 'WEDDING_PLANNER'):
                try:
                    validar_limite_planners(empresa)
                except LimitePlanExcedido as exc:
                    return bloquear_por_limite_plan(request, empresa, str(exc), '#equipo')
            crear_o_actualizar_usuario_empresa(request, empresa, 'WEDDING_PLANNER')
            return redirect(f'/dashboard/empresa/?empresa={empresa.id}#equipo')
        if accion == 'editar_usuario_empresa' and puede_usuarios:
            membresia = actualizar_usuario_empresa_dashboard(request, empresa)
            origen = request.POST.get('origen') or ('clientes' if membresia.rol == 'CLIENTE' else 'equipo')
            return redirect(f'/dashboard/empresa/?empresa={empresa.id}#{origen}')
        if accion == 'desactivar_usuario_empresa' and puede_usuarios:
            membresia = desactivar_usuario_empresa_dashboard(request, empresa)
            origen = request.POST.get('origen') or ('clientes' if membresia.rol == 'CLIENTE' else 'equipo')
            return redirect(f'/dashboard/empresa/?empresa={empresa.id}#{origen}')
        if accion == 'eliminar_usuario_empresa' and puede_usuarios:
            rol = eliminar_usuario_empresa_dashboard(request, empresa)
            origen = request.POST.get('origen') or ('clientes' if rol == 'CLIENTE' else 'equipo')
            return redirect(f'/dashboard/empresa/?empresa={empresa.id}#{origen}')
        if accion == 'crear_cliente' and puede_usuarios:
            username = limpiar_texto(request, 'username_usuario')
            if not usuario_empresa_existente_activo(empresa, username):
                try:
                    validar_limite_usuarios(empresa)
                except LimitePlanExcedido as exc:
                    return bloquear_por_limite_plan(request, empresa, str(exc), '#clientes')
            if not usuario_empresa_existente_activo(empresa, username, 'CLIENTE'):
                try:
                    validar_limite_clientes(empresa)
                except LimitePlanExcedido as exc:
                    return bloquear_por_limite_plan(request, empresa, str(exc), '#clientes')
            crear_cliente_empresa_dashboard(request, empresa)
            return redirect(f'/dashboard/empresa/?empresa={empresa.id}#clientes')
        if accion == 'crear_sede' and puede_catalogos:
            nombre = limpiar_texto(request, 'nombre_sede')
            if nombre:
                SedeEvento.objects.create(
                    empresa=empresa,
                    nombre=nombre,
                    tipo=request.POST.get('tipo_sede') or 'SALON',
                    capacidad_minima=convertir_entero(request.POST.get('capacidad_minima'), 0),
                    capacidad_maxima=convertir_entero(request.POST.get('capacidad_maxima'), 0),
                    precio_base=convertir_decimal(request.POST.get('precio_base_sede'), 0),
                    direccion=limpiar_texto(request, 'direccion_sede'),
                    activa=True,
                )
            return redirect(f'/dashboard/empresa/?empresa={empresa.id}#catalogos')
        if accion == 'editar_sede' and puede_catalogos:
            actualizar_sede_empresa_dashboard(request, empresa)
            return redirect(f'/dashboard/empresa/?empresa={empresa.id}#catalogos')
        if accion == 'desactivar_sede' and puede_catalogos:
            sede = get_object_or_404(SedeEvento, empresa=empresa, id=request.POST.get('sede_id'))
            sede.activa = False
            sede.save(update_fields=['activa'])
            return redirect(f'/dashboard/empresa/?empresa={empresa.id}#catalogos')
        if accion == 'eliminar_sede' and puede_catalogos:
            eliminar_sede_empresa_dashboard(request, empresa)
            return redirect(f'/dashboard/empresa/?empresa={empresa.id}#catalogos')
        if accion == 'crear_proveedor' and puede_catalogos:
            nombre = limpiar_texto(request, 'nombre_proveedor')
            if nombre:
                proveedor = Proveedor.objects.create(
                    empresa=empresa,
                    nombre_comercial=nombre,
                    tipo_proveedor=request.POST.get('tipo_proveedor') or 'OTROS',
                    nombre_contacto=limpiar_texto(request, 'contacto_proveedor'),
                    telefono=limpiar_texto(request, 'telefono_proveedor'),
                    correo=limpiar_texto(request, 'correo_proveedor'),
                    contacto_operativo=limpiar_texto(request, 'contacto_operativo_proveedor'),
                    telefono_operativo=limpiar_texto(request, 'telefono_operativo_proveedor'),
                    correo_operativo=limpiar_texto(request, 'correo_operativo_proveedor'),
                    visible_para_wedding_planners=request.POST.get('visible_wedding_planners_proveedor', 'on') == 'on',
                    rfc=limpiar_texto(request, 'rfc_proveedor'),
                    datos_bancarios=limpiar_texto(request, 'datos_bancarios_proveedor'),
                    notas_privadas=limpiar_texto(request, 'notas_privadas_proveedor'),
                    activo=True,
                )
                sincronizar_acceso_proveedor(
                    request,
                    empresa,
                    proveedor,
                )
            return redirect(f'/dashboard/empresa/?empresa={empresa.id}#proveedores')
        if accion == 'editar_proveedor' and puede_catalogos:
            actualizar_proveedor_empresa_dashboard(request, empresa)
            return redirect(f'/dashboard/empresa/?empresa={empresa.id}#proveedores')
        if accion == 'desactivar_proveedor' and puede_catalogos:
            proveedor = get_object_or_404(Proveedor, empresa=empresa, id=request.POST.get('proveedor_id'))
            proveedor.activo = False
            proveedor.save(update_fields=['activo'])
            return redirect(f'/dashboard/empresa/?empresa={empresa.id}#proveedores')
        if accion == 'eliminar_proveedor' and puede_catalogos:
            eliminar_proveedor_empresa_dashboard(request, empresa)
            return redirect(f'/dashboard/empresa/?empresa={empresa.id}#proveedores')
        if accion == 'crear_paquete' and puede_catalogos:
            nombre = limpiar_texto(request, 'nombre_paquete')
            if nombre:
                PaqueteBoda.objects.create(
                    empresa=empresa,
                    nombre=nombre,
                    descripcion=limpiar_texto(request, 'descripcion_paquete'),
                    precio_base=convertir_decimal(request.POST.get('precio_base_paquete'), 0),
                    numero_personas_incluidas=convertir_entero(request.POST.get('personas_paquete'), 0),
                    activo=True,
                )
            return redirect(f'/dashboard/empresa/?empresa={empresa.id}#catalogos')
        if accion == 'editar_paquete' and puede_catalogos:
            actualizar_paquete_empresa_dashboard(request, empresa)
            return redirect(f'/dashboard/empresa/?empresa={empresa.id}#catalogos')
        if accion == 'desactivar_paquete' and puede_catalogos:
            paquete = get_object_or_404(PaqueteBoda, empresa=empresa, id=request.POST.get('paquete_id'))
            paquete.activo = False
            paquete.save(update_fields=['activo'])
            return redirect(f'/dashboard/empresa/?empresa={empresa.id}#catalogos')
        if accion == 'eliminar_paquete' and puede_catalogos:
            eliminar_paquete_empresa_dashboard(request, empresa)
            return redirect(f'/dashboard/empresa/?empresa={empresa.id}#catalogos')

    eventos = EventoBoda.objects.filter(empresa=empresa).select_related('sede', 'wedding_planner').order_by('-activo', '-fecha_fiesta')
    membresias = MembresiaEmpresa.objects.filter(empresa=empresa).select_related('usuario').order_by('-activo', 'rol', 'usuario__username')
    roles_equipo_valores = {
        'ADMIN_EMPRESA',
        'VENTAS',
        'WEDDING_PLANNER',
    }
    membresias_equipo = membresias.filter(
        rol__in=roles_equipo_valores
    )
    planners = membresias.filter(
        rol='WEDDING_PLANNER',
        activo=True,
    )
    clientes_empresa = membresias.filter(
        rol='CLIENTE'
    )
    roles_equipo = [
        item
        for item in MembresiaEmpresa.ROLES
        if item[0] in roles_equipo_valores
    ]
    context = {
        'empresa': empresa,
        'empresas': empresas,
        'eventos': eventos,
        'resumen_eventos': resumen_eventos(eventos),
        'resumen_operativo': resumen_operativo_eventos(eventos),
        'uso_plan': resumen_uso_plan(empresa),
        'suscripcion_empresa': getattr(empresa, 'suscripcion', None),
        'membresias': membresias,
        'membresias_equipo': membresias_equipo,
        'planners': planners,
        'clientes_empresa': clientes_empresa,
        'roles_equipo': roles_equipo,
        'eventos_activos_empresa': eventos.filter(activo=True),
        'sedes': empresa.sedes.all(),
        'proveedores': empresa.proveedores.all(),
        'paquetes': empresa.paquetes.all(),
        'puede_catalogos': puede_catalogos,
        'puede_usuarios': puede_usuarios,
        'roles_usuario_actual': roles_usuario_actual,
        'es_planner_empresa': 'WEDDING_PLANNER' in roles_usuario_actual,
        'tipos_evento': EventoBoda.TIPOS_EVENTO,
        'estados_evento': EventoBoda.ESTADOS_EVENTO,
        'roles_empresa': MembresiaEmpresa.ROLES,
        'tipos_sede': SedeEvento.TIPOS,
        'tipos_proveedor': Proveedor.TIPOS,
    }
    context.update(construir_contexto_dashboard_empresa_v3(
        empresa=empresa,
        eventos=eventos,
    ))
    return render(request, 'invitaciones/dashboard_empresa.html', context)


@login_required(login_url=LOGIN_DASHBOARD_URL)
def dashboard_planner(request):
    tenant = exigir_tenant(
        request,
        roles={'WEDDING_PLANNER'},
        allow_dirtec_switch=False,
    )
    empresa_actual = tenant.empresa
    eventos = (
        eventos_visibles_usuario(request.user)
        .filter(
            wedding_planner=request.user,
            empresa=empresa_actual,
        )
        .select_related('empresa', 'sede')
        .order_by('-activo', 'fecha_fiesta')
    )

    if request.method == 'POST' and request.POST.get('accion') == 'crear_evento_planner':
        if not empresa_actual or 'WEDDING_PLANNER' not in roles_activos_empresa(request.user, empresa_actual):
            bloquear_accion_dashboard(
                request,
                empresa=empresa_actual,
                accion='crear_evento_planner',
                permiso='wedding planner activo de la empresa',
            )
        try:
            validar_limite_eventos_activos(empresa_actual)
        except LimitePlanExcedido as exc:
            messages.error(request, str(exc))
            registrar_auditoria(
                usuario=request.user,
                empresa=empresa_actual,
                accion='LIMITE_PLAN_EXCEDIDO',
                modelo='EmpresaSuscriptora',
                objeto_id=empresa_actual.id,
                descripcion=str(exc),
                request=request,
            )
            return redirect(
                f'/empresa/{empresa_actual.slug}/planner/dashboard/#eventos'
            )
        crear_evento_planner_dashboard(request, empresa_actual)
        return redirect(
            f'/empresa/{empresa_actual.slug}/planner/dashboard/#eventos'
        )

    if request.method == 'POST' and request.POST.get('accion') == 'actualizar_estado_evento':
        evento = get_object_or_404(eventos, id=request.POST.get('evento_id'))
        estado = request.POST.get('estado')
        estados_validos = {valor for valor, _ in EventoBoda.ESTADOS_EVENTO}
        if estado in estados_validos:
            evento.estado = estado
            evento.save(update_fields=['estado'])
        return redirect(
            f'/empresa/{evento.empresa.slug}/planner/dashboard/#eventos'
            if evento.empresa
            else '/dashboard/planner/#eventos'
        )

    eventos_lista = list(eventos)
    context = {
        'eventos': eventos,
        'eventos_lista': eventos_lista,
        'empresa_actual': empresa_actual,
        'puede_crear_eventos_planner': True,
        'tipos_evento': EventoBoda.TIPOS_EVENTO,
        'estados_evento': EventoBoda.ESTADOS_EVENTO,
    }
    context.update(
        construir_contexto_dashboard_planner_v3(
            eventos=eventos_lista,
            planner=request.user,
            empresa=empresa_actual,
        )
    )
    return render(request, 'invitaciones/dashboard_planner.html', context)


@login_required(login_url=LOGIN_DASHBOARD_URL)
def dashboard_planner_slug(request, empresa_slug):
    validar_slug_tenant(
        request,
        empresa_slug,
        roles={'WEDDING_PLANNER'},
    )
    return dashboard_planner(request)


def inicio(request):
    return render(request, 'invitaciones/inicio.html')


@login_required(login_url=LOGIN_DASHBOARD_URL)
def portal_cliente(request):
    eventos = eventos_para_cliente(
        request.user
    )
    evento = obtener_evento_portal(
        request,
        eventos,
    )

    if not evento:
        return render(
            request,
            'invitaciones/portal_cliente.html',
            {
                'evento': None,
                'eventos': eventos,
            },
        )

    # Client portal intentionally avoids exposing internal provider costs,
    # margins and company-operation tasks. Only client-facing information is
    # projected here.
    invitados_metricas = resumen_invitados_evento(
        evento
    )

    aprobaciones = (
        AprobacionEvento.objects
        .filter(evento=evento)
        .order_by(
            'estado',
            '-fecha_solicitud',
        )
    )
    aprobaciones_pendientes_qs = (
        aprobaciones.filter(
            estado='PENDIENTE'
        )
    )

    documentos = (
        DocumentoEvento.objects
        .filter(
            evento=evento,
            visible_cliente=True,
        )
        .order_by('-fecha_carga')
    )

    tareas_cliente_qs = (
        TareaEvento.objects
        .filter(
            evento=evento,
            responsable=request.user,
        )
        .exclude(
            estado__in=[
                'COMPLETADA',
                'CANCELADA',
            ]
        )
        .order_by(
            'fecha_limite',
            'fecha_inicio',
            'prioridad',
        )
    )
    tareas_cliente = list(
        tareas_cliente_qs[:12]
    )

    hoy = timezone.localdate()
    ahora = timezone.localtime()

    proximos_hitos = []
    for tarea in tareas_cliente:
        fecha = (
            tarea.fecha_inicio
            or tarea.fecha_limite
        )
        if not fecha:
            continue
        if fecha < hoy:
            continue

        proximos_hitos.append({
            'tipo': 'TAREA',
            'titulo': tarea.titulo,
            'fecha': fecha,
            'hora_inicio': None,
            'hora_fin': None,
            'estado': tarea.get_estado_display(),
            'prioridad': tarea.get_prioridad_display(),
            'descripcion': tarea.descripcion or '',
        })

    # Ceremony/reception are always client-facing milestones.
    if evento.fecha_misa and evento.fecha_misa >= ahora:
        proximos_hitos.append({
            'tipo': 'EVENTO',
            'titulo': (
                evento.lugar_misa
                or 'Ceremonia'
            ),
            'fecha': evento.fecha_misa.date(),
            'hora_inicio': evento.fecha_misa.time(),
            'hora_fin': None,
            'estado': 'Ceremonia',
            'prioridad': '',
            'descripcion': (
                evento.direccion_ceremonia
                or ''
            ),
        })

    if evento.fecha_fiesta and evento.fecha_fiesta >= ahora:
        proximos_hitos.append({
            'tipo': 'EVENTO',
            'titulo': (
                evento.lugar_fiesta
                or 'Recepcion'
            ),
            'fecha': evento.fecha_fiesta.date(),
            'hora_inicio': evento.fecha_fiesta.time(),
            'hora_fin': None,
            'estado': 'Evento',
            'prioridad': '',
            'descripcion': (
                evento.direccion_recepcion
                or ''
            ),
        })

    proximos_hitos.sort(
        key=lambda item: (
            item['fecha'],
            item['hora_inicio']
            or datetime.min.time(),
        )
    )

    diseno = (
        DisenoInvitacion.objects
        .filter(evento=evento)
        .first()
    )
    invitacion_publicada = bool(
        diseno
        and diseno.documento_builder_publicado
    )

    total_personas = (
        invitados_metricas[
            'total_personas'
        ]
    )
    total_confirmados = (
        invitados_metricas[
            'confirmados'
        ]
    )
    total_pendientes = (
        invitados_metricas[
            'pendientes'
        ]
    )

    porcentaje_rsvp = (
        round(
            (
                total_confirmados
                / total_personas
            )
            * 100
        )
        if total_personas
        else 0
    )

    # K.8.4.1: el workspace del cliente nace de ServicioEvento, no de
    # ExpedienteServicio legacy. Esto permite conversar sobre cualquier
    # servicio del evento aunque nunca haya existido una solicitud legacy.
    servicios_cliente_workspace = list(
        ServicioEvento.objects
        .filter(evento=evento)
        .exclude(estado="CANCELADO")
        .select_related("proveedor")
        .order_by("fecha_servicio", "nombre_servicio", "id")
    )

    contexto_evento = {
        'tipo': evento.get_tipo_evento_display(),
        'estado': evento.get_estado_display(),
        'fecha_principal': evento.fecha_fiesta,
        'sede': (
            evento.sede.nombre
            if evento.sede
            else evento.lugar_fiesta
        ),
        'planner': (
            evento.wedding_planner.get_full_name()
            or evento.wedding_planner.username
            if evento.wedding_planner
            else ''
        ),
    }

    cliente_v3 = construir_contexto_portal_cliente_v3(evento, request.user)
    cliente_invitados_v3 = construir_contexto_invitados_cliente_v3(evento, request)

    context = {
        'evento': evento,
        'eventos': eventos,
        'contexto_evento': contexto_evento,
        'total_lugares': total_personas,
        'total_confirmados': total_confirmados,
        'total_pendientes': total_pendientes,
        'porcentaje_rsvp': porcentaje_rsvp,
        'aprobaciones': aprobaciones[:12],
        'aprobaciones_pendientes': (
            aprobaciones_pendientes_qs.count()
        ),
        'documentos': documentos[:12],
        'documentos_total': documentos.count(),
        'tareas_cliente': tareas_cliente,
        'tareas_cliente_total': (
            tareas_cliente_qs.count()
        ),
        'proximos_hitos': proximos_hitos[:8],
        'invitacion_publicada': invitacion_publicada,
        'invitacion_publicada_en': (
            diseno.publicado_en
            if diseno
            else None
        ),
        'servicios_cliente_workspace': servicios_cliente_workspace,
        **cliente_v3,
        **cliente_invitados_v3,
    }
    return render(
        request,
        'invitaciones/portal_cliente.html',
        context,
    )


@login_required(login_url=LOGIN_DASHBOARD_URL)
@require_POST
def responder_aprobacion_cliente(request, aprobacion_id):
    aprobacion = get_object_or_404(
        AprobacionEvento.objects.select_related('evento'),
        id=aprobacion_id,
    )

    if not usuario_puede_ver_evento_cliente(request.user, aprobacion.evento):
        return JsonResponse({'ok': False, 'error': 'No autorizado'}, status=403)

    accion = request.POST.get('accion')
    estados = {
        'aprobar': 'APROBADO',
        'rechazar': 'RECHAZADO',
        'cambios': 'CAMBIOS',
    }
    if accion not in estados:
        return JsonResponse({'ok': False, 'error': 'Accion invalida'}, status=400)

    aprobacion.estado = estados[accion]
    aprobacion.aprobado_por = request.user
    aprobacion.comentario = request.POST.get('comentario', '').strip()
    aprobacion.save()
    crear_notificaciones_evento(
        aprobacion.evento,
        f'Aprobacion respondida: {aprobacion.titulo}',
        f'{request.user} marco la solicitud como {aprobacion.get_estado_display()}.',
        tipo='APROBACION',
        enlace=f'/admin/aprobaciones/aprobacionevento/{aprobacion.id}/change/',
        usuarios_extra=[aprobacion.solicitado_por],
    )

    return redirect(f'/cliente/dashboard/?evento={aprobacion.evento_id}#aprobaciones')


@login_required(login_url=LOGIN_DASHBOARD_URL)
def portal_proveedor(request):
    proveedor = proveedor_de_usuario(request.user)

    if not proveedor:
        return render(
            request,
            'invitaciones/portal_proveedor.html',
            {
                'proveedor': None,
                'evento': None,
                'eventos': EventoBoda.objects.none(),
            },
        )

    eventos = (
        EventoBoda.objects
        .filter(servicios_contratados__proveedor=proveedor)
        .distinct()
        .order_by('fecha_fiesta')
    )
    evento = obtener_evento_portal(request, eventos)

    if not evento:
        return render(
            request,
            'invitaciones/portal_proveedor.html',
            {
                'proveedor': proveedor,
                'evento': None,
                'eventos': eventos,
            },
        )

    context = {
        'proveedor': proveedor,
        'evento': evento,
        'eventos': eventos,
        **construir_contexto_portal_proveedor_v3(
            evento,
            proveedor,
            request.user,
        ),
    }
    return render(
        request,
        'invitaciones/portal_proveedor.html',
        context,
    )


@login_required(login_url=LOGIN_DASHBOARD_URL)
@require_POST
def subir_documento_proveedor(request):
    proveedor = proveedor_de_usuario(request.user)
    if not proveedor:
        raise PermissionDenied('No autorizado.')

    servicio = get_object_or_404(
        ServicioEvento.objects.select_related('evento'),
        id=request.POST.get('servicio_evento_id'),
        proveedor=proveedor,
    )
    archivo = request.FILES.get('archivo')
    if not archivo:
        messages.error(request, 'Selecciona un archivo.')
        return redirect(f"{reverse('portal_proveedor')}?evento={servicio.evento_id}#documentos")
    if not archivo_pasa_validadores(DocumentoEvento, 'archivo', archivo):
        messages.error(request, 'Archivo no permitido.')
        return redirect(f"{reverse('portal_proveedor')}?evento={servicio.evento_id}#documentos")

    documento = DocumentoEvento(
        evento=servicio.evento,
        servicio_evento=servicio,
        proveedor=proveedor,
        tipo_documento=request.POST.get('tipo_documento') or 'OTRO',
        titulo=(request.POST.get('titulo') or '').strip() or archivo.name,
        archivo=archivo,
        descripcion=(request.POST.get('descripcion') or '').strip() or None,
        cargado_por=request.user,
        visible_cliente=False,
        visible_proveedor=True,
    )
    try:
        documento.full_clean()
        documento.save()
        messages.success(request, 'Documento compartido con el equipo.')
    except ValidationError as exc:
        messages.error(request, ' · '.join(exc.messages))

    return redirect(f"{reverse('portal_proveedor')}?evento={servicio.evento_id}#documentos")


SECCIONES_INVITACION_DEFAULTS = {
    'PORTADA': {'orden': 5, 'titulo_attr': 'frase_portada', 'descripcion_attr': 'mensaje_general', 'fondo_attr': 'foto_portada', 'activa': True},
    'PADRES_PADRINOS': {'orden': 12, 'titulo': 'Nuestros padres y padrinos', 'descripcion': '', 'activa': True},
    'CUENTA_REGRESIVA': {'orden': 20, 'titulo_attr': 'titulo_cuenta_regresiva', 'descripcion': '', 'fondo_attr': 'fondo_cuenta_regresiva', 'activa': True},
    'DETALLES': {'orden': 30, 'titulo_attr': 'titulo_detalles', 'descripcion_attr': 'texto_detalles', 'fondo_attr': 'fondo_detalles', 'activa': True},
    'DRESS_CODE': {'orden': 35, 'titulo': 'Dress Code', 'descripcion_attr': 'dress_code_descripcion', 'fondo_attr': 'fondo_detalles', 'activa': True},
    'ITINERARIO': {'orden': 40, 'titulo': 'Itinerario', 'descripcion': '', 'activa': True},
    'ALBUM': {'orden': 50, 'titulo_attr': 'titulo_album', 'descripcion': '', 'fondo_attr': 'fondo_album', 'activa': True},
    'MENU': {'orden': 60, 'titulo_attr': 'titulo_menu', 'descripcion': '', 'fondo_attr': 'fondo_menu', 'activa': True},
    'REGALOS': {'orden': 70, 'titulo_attr': 'titulo_regalos', 'descripcion': '', 'fondo_attr': 'fondo_regalos', 'activa': True},
    'ALBUM_COMPARTIDO': {
        'orden': 75,
        'titulo_attr': 'titulo_album_compartido',
        'descripcion_attr': 'texto_album_compartido',
        'activa': True,
    },
    'RSVP': {'orden': 80, 'titulo_attr': 'titulo_rsvp', 'descripcion_attr': 'texto_rsvp', 'fondo_attr': 'fondo_rsvp', 'activa': True},
}


def datos_default_seccion(evento, tipo):
    defaults = SECCIONES_INVITACION_DEFAULTS[tipo]
    titulo = defaults.get('titulo')
    descripcion = defaults.get('descripcion', '')
    if defaults.get('titulo_attr'):
        titulo = getattr(evento, defaults['titulo_attr'], '') or titulo
    if defaults.get('descripcion_attr'):
        descripcion = getattr(evento, defaults['descripcion_attr'], '') or descripcion
    fondo = getattr(evento, defaults['fondo_attr'], None) if defaults.get('fondo_attr') else None
    return {
        'tipo': tipo,
        'titulo': titulo or dict(SeccionInvitacion.TIPOS).get(tipo, tipo.title()),
        'descripcion': descripcion or '',
        'orden': defaults.get('orden', 0),
        'activa': defaults.get('activa', True),
        'fondo': fondo,
        'fondo_es_video': es_video_archivo(fondo),
        'imagen_titulo': None,
        'imagen_titulo_es_video': False,
        'mostrar_titulo_texto': True,
        'posicion_fondo': 'center center',
        'opacidad_fondo': '0.18',
        'color_texto': '',
    }


EDITOR_ANIMACIONES = {'none', 'fade', 'slide', 'soft'}
EDITOR_ESPACIADOS = {'compact', 'soft', 'wide'}
EDITOR_ALINEACIONES = {'left', 'center', 'right'}
EDITOR_TAMANOS_TITULO = {'small', 'medium', 'large'}
EDITOR_BG_FIT = {'contain', 'cover', 'repeat', 'free'}
EDITOR_LAYER_TYPES = {'background', 'title', 'text', 'decor'}
EDITOR_DECOR_STYLES = {'line', 'flourish', 'rings', 'dots'}
EDITOR_CUSTOM_LAYER_TYPES = {'text', 'image', 'video'}
EDITOR_MEDIA_FIT = {'contain', 'cover'}
EDITOR_REAL_CONTENT_SECTIONS = {
    'DETALLES',
    'REGALOS',
    'ALBUM',
    'ALBUM_COMPARTIDO',
}
EDITOR_REAL_CONTENT_LAYOUTS = {'grid', 'stack'}
EDITOR_REAL_MEDIA_POSITIONS = {'top', 'bottom', 'hidden'}
EDITOR_REAL_MAP_DISPLAY = {'button-map', 'button-only', 'map-only', 'hidden'}
HEX_COLOR_RE = re.compile(r'^#[0-9a-fA-F]{6}$')


PLANTILLAS_EVENTO = [
    ('BODA', 'Boda elegante'),
    ('XV', 'XV años'),
    ('BABY_SHOWER', 'Baby shower'),
    ('BAUTIZO', 'Bautizo'),
    ('CUMPLEANOS', 'Cumpleanos'),
    ('CORPORATIVO', 'Corporativo'),
]

CONFIGURACION_PLANTILLAS_EVENTO = {
    'BODA': {
        'campos': {
            'tipo_evento': 'BODA',
            'paleta_colores': 'TINTA_MARFIL',
            'paleta_sobre': 'TINTA_MARFIL',
            'estilo_letra': 'CURSIVA_ELEGANTE',
            'etiqueta_principal': 'Novia',
            'etiqueta_secundario': 'Novio',
            'mostrar_nombre_secundario': True,
            'mostrar_album': True,
            'mostrar_menu': True,
            'mostrar_regalos': True,
            'mostrar_mapa': True,
            'mostrar_album_compartido': True,
            'frase_portada': 'Nuestra Boda',
            'titulo_invitacion': 'Con mucha alegria te invitamos',
            'texto_invitacion': 'Tu presencia hara este dia todavia mas especial.',
            'titulo_cuenta_regresiva': 'Faltan',
            'titulo_detalles': 'Ceremonia y recepcion',
            'texto_detalles': 'Ceremonia, recepcion, ubicacion y dress code reunidos para acompanarnos sin complicaciones.',
            'titulo_album': 'Nuestra historia',
            'titulo_menu': 'Menu de la boda',
            'titulo_regalos': 'Mesa de regalos',
            'titulo_album_compartido': 'Comparte tus fotos',
            'texto_album_compartido': 'Ayudanos a guardar tus mejores momentos. Sube tus fotos y videos al album compartido.',
            'titulo_rsvp': 'Confirma tu asistencia',
            'texto_rsvp': 'Tu respuesta nos ayuda a organizar lugares, mesas y buffet.',
        },
        'secciones': {
            'PORTADA': {'orden': 5, 'activa': True},
            'PADRES_PADRINOS': {'orden': 12, 'activa': True, 'titulo': 'Nuestros padres y padrinos'},
            'CUENTA_REGRESIVA': {'orden': 20, 'activa': True},
            'DETALLES': {'orden': 30, 'activa': True},
            'DRESS_CODE': {'orden': 35, 'activa': True},
            'ITINERARIO': {'orden': 40, 'activa': True, 'titulo': 'Itinerario'},
            'ALBUM': {'orden': 50, 'activa': True},
            'MENU': {'orden': 60, 'activa': True},
            'REGALOS': {'orden': 70, 'activa': True},
            'ALBUM_COMPARTIDO': {'orden': 75, 'activa': True},
            'RSVP': {'orden': 80, 'activa': True},
        },
    },
    'XV': {
        'campos': {
            'tipo_evento': 'XV',
            'paleta_colores': 'ROSA',
            'paleta_sobre': 'ROSA',
            'estilo_letra': 'ROMANTICA',
            'etiqueta_principal': 'Festejada',
            'etiqueta_secundario': 'Familia',
            'mostrar_nombre_secundario': False,
            'mostrar_album': True,
            'mostrar_menu': True,
            'mostrar_regalos': True,
            'mostrar_mapa': True,
            'mostrar_album_compartido': True,
            'frase_portada': 'Mis XV años',
            'titulo_invitacion': 'Celebra conmigo',
            'texto_invitacion': 'Me encantaria compartir esta noche tan especial contigo.',
            'titulo_cuenta_regresiva': 'Faltan',
            'titulo_detalles': 'Misa y fiesta',
            'texto_detalles': 'Aqui encontraras horarios, ubicacion, dress code y todo lo necesario para celebrar juntos.',
            'titulo_album': 'Momentos especiales',
            'titulo_menu': 'Menu de la fiesta',
            'titulo_regalos': 'Regalos',
            'titulo_album_compartido': 'Sube tus fotos',
            'texto_album_compartido': 'Comparte fotos y videos de la fiesta para guardar todos los recuerdos.',
            'titulo_rsvp': 'Confirma tu asistencia',
            'texto_rsvp': 'Tu confirmacion nos ayuda a preparar tu lugar en la fiesta.',
        },
        'secciones': {
            'PORTADA': {'orden': 5, 'activa': True},
            'PADRES_PADRINOS': {'orden': 12, 'activa': True, 'titulo': 'Mis padres y padrinos'},
            'CUENTA_REGRESIVA': {'orden': 20, 'activa': True},
            'DETALLES': {'orden': 30, 'activa': True},
            'DRESS_CODE': {'orden': 35, 'activa': True},
            'ITINERARIO': {'orden': 40, 'activa': True, 'titulo': 'Programa'},
            'ALBUM': {'orden': 50, 'activa': True},
            'MENU': {'orden': 60, 'activa': True},
            'REGALOS': {'orden': 70, 'activa': True},
            'ALBUM_COMPARTIDO': {'orden': 75, 'activa': True},
            'RSVP': {'orden': 80, 'activa': True},
        },
    },
    'BABY_SHOWER': {
        'campos': {
            'tipo_evento': 'BABY_SHOWER',
            'paleta_colores': 'SALVIA_PERLA',
            'paleta_sobre': 'SALVIA_PERLA',
            'estilo_letra': 'MANUSCRITA',
            'etiqueta_principal': 'Bebe',
            'etiqueta_secundario': 'Familia',
            'mostrar_nombre_secundario': False,
            'mostrar_album': True,
            'mostrar_menu': True,
            'mostrar_regalos': True,
            'mostrar_mapa': True,
            'mostrar_album_compartido': True,
            'frase_portada': 'Baby Shower',
            'titulo_invitacion': 'Celebremos esta dulce espera',
            'texto_invitacion': 'Acompananos a celebrar la llegada de nuestro bebe.',
            'titulo_cuenta_regresiva': 'Faltan',
            'titulo_detalles': 'Detalles del baby shower',
            'texto_detalles': 'Te compartimos horario, direccion, mesa de regalos y detalles para celebrar esta llegada especial.',
            'titulo_album': 'Momentos del baby shower',
            'titulo_menu': 'Bocadillos y bebidas',
            'titulo_regalos': 'Mesa de regalos',
            'titulo_album_compartido': 'Comparte tus fotos',
            'texto_album_compartido': 'Sube fotos y videos del baby shower para guardar este recuerdo.',
            'titulo_rsvp': 'Confirma tu asistencia',
            'texto_rsvp': 'Tu confirmacion nos ayuda a preparar cada detalle.',
        },
        'secciones': {
            'PORTADA': {'orden': 5, 'activa': True},
            'PADRES_PADRINOS': {'orden': 12, 'activa': False, 'titulo': 'Familia'},
            'CUENTA_REGRESIVA': {'orden': 20, 'activa': True},
            'DETALLES': {'orden': 30, 'activa': True},
            'DRESS_CODE': {'orden': 35, 'activa': False},
            'ITINERARIO': {'orden': 40, 'activa': False, 'titulo': 'Actividades'},
            'ALBUM': {'orden': 50, 'activa': True},
            'MENU': {'orden': 60, 'activa': True},
            'REGALOS': {'orden': 70, 'activa': True},
            'ALBUM_COMPARTIDO': {'orden': 75, 'activa': True},
            'RSVP': {'orden': 80, 'activa': True},
        },
    },
    'BAUTIZO': {
        'campos': {
            'tipo_evento': 'BAUTIZO',
            'paleta_colores': 'DORADO',
            'paleta_sobre': 'DORADO',
            'estilo_letra': 'CLASICA',
            'etiqueta_principal': 'Bautizado',
            'etiqueta_secundario': 'Familia',
            'mostrar_nombre_secundario': False,
            'mostrar_album': True,
            'mostrar_menu': True,
            'mostrar_regalos': False,
            'mostrar_mapa': True,
            'mostrar_album_compartido': True,
            'frase_portada': 'Mi Bautizo',
            'titulo_invitacion': 'Con alegria te invitamos',
            'texto_invitacion': 'Acompanamos en este dia de bendicion y celebracion familiar.',
            'titulo_cuenta_regresiva': 'Faltan',
            'titulo_detalles': 'Ceremonia y celebracion',
            'texto_detalles': 'Aqui tienes los horarios, ubicaciones y detalles para acompanarnos en este dia especial.',
            'titulo_album': 'Recuerdos del bautizo',
            'titulo_menu': 'Menu de la celebracion',
            'titulo_regalos': 'Detalles',
            'titulo_album_compartido': 'Comparte tus fotos',
            'texto_album_compartido': 'Ayudanos a guardar fotos y videos de este dia tan especial.',
            'titulo_rsvp': 'Confirma tu asistencia',
            'texto_rsvp': 'Tu respuesta nos ayuda a organizar la celebracion.',
        },
        'secciones': {
            'PORTADA': {'orden': 5, 'activa': True},
            'PADRES_PADRINOS': {'orden': 12, 'activa': True, 'titulo': 'Padres y padrinos'},
            'CUENTA_REGRESIVA': {'orden': 20, 'activa': True},
            'DETALLES': {'orden': 30, 'activa': True},
            'DRESS_CODE': {'orden': 35, 'activa': False},
            'ITINERARIO': {'orden': 40, 'activa': True, 'titulo': 'Itinerario'},
            'ALBUM': {'orden': 50, 'activa': True},
            'MENU': {'orden': 60, 'activa': True},
            'REGALOS': {'orden': 70, 'activa': False},
            'ALBUM_COMPARTIDO': {'orden': 75, 'activa': True},
            'RSVP': {'orden': 80, 'activa': True},
        },
    },
    'CUMPLEANOS': {
        'campos': {
            'tipo_evento': 'CUMPLEANOS',
            'paleta_colores': 'TERRACOTA',
            'paleta_sobre': 'TERRACOTA',
            'estilo_letra': 'MODERNA',
            'etiqueta_principal': 'Festejado',
            'etiqueta_secundario': 'Familia',
            'mostrar_nombre_secundario': False,
            'mostrar_album': True,
            'mostrar_menu': True,
            'mostrar_regalos': True,
            'mostrar_mapa': True,
            'mostrar_album_compartido': True,
            'frase_portada': 'Mi Cumpleanos',
            'titulo_invitacion': 'Celebremos juntos',
            'texto_invitacion': 'Tu presencia hara que este festejo sea mucho mas especial.',
            'titulo_cuenta_regresiva': 'Faltan',
            'titulo_detalles': 'Detalles de la fiesta',
            'texto_detalles': 'Aqui tienes horario, ubicacion, regalos y detalles para llegar sin complicaciones.',
            'titulo_album': 'Momentos especiales',
            'titulo_menu': 'Menu del festejo',
            'titulo_regalos': 'Mesa de regalos',
            'titulo_album_compartido': 'Comparte tus fotos',
            'texto_album_compartido': 'Sube tus fotos y videos para guardar juntos los mejores momentos.',
            'titulo_rsvp': 'Confirma tu asistencia',
            'texto_rsvp': 'Tu respuesta nos ayuda a organizar lugares, alimentos y detalles.',
        },
        'secciones': {
            'PORTADA': {'orden': 5, 'activa': True},
            'PADRES_PADRINOS': {'orden': 12, 'activa': False, 'titulo': 'Familia'},
            'CUENTA_REGRESIVA': {'orden': 20, 'activa': True},
            'DETALLES': {'orden': 30, 'activa': True},
            'DRESS_CODE': {'orden': 35, 'activa': True, 'titulo': 'Dress code'},
            'ITINERARIO': {'orden': 40, 'activa': True, 'titulo': 'Programa'},
            'ALBUM': {'orden': 50, 'activa': True},
            'MENU': {'orden': 60, 'activa': True},
            'REGALOS': {'orden': 70, 'activa': True},
            'ALBUM_COMPARTIDO': {'orden': 75, 'activa': True},
            'RSVP': {'orden': 80, 'activa': True},
        },
    },
    'CORPORATIVO': {
        'campos': {
            'tipo_evento': 'CORPORATIVO',
            'paleta_colores': 'AZUL',
            'paleta_sobre': 'AZUL',
            'estilo_letra': 'MODERNA',
            'etiqueta_principal': 'Evento',
            'etiqueta_secundario': 'Empresa',
            'mostrar_nombre_secundario': False,
            'mostrar_album': False,
            'mostrar_menu': True,
            'mostrar_regalos': False,
            'mostrar_mapa': True,
            'mostrar_album_compartido': False,
            'frase_portada': 'Evento Corporativo',
            'titulo_invitacion': 'Tenemos el gusto de invitarte',
            'texto_invitacion': 'Te compartimos los detalles para acompanar este encuentro profesional.',
            'titulo_cuenta_regresiva': 'Faltan',
            'titulo_detalles': 'Agenda y ubicacion',
            'texto_detalles': 'Consulta horarios, sede, accesos y detalles importantes del evento.',
            'titulo_album': 'Galeria',
            'titulo_menu': 'Menu',
            'titulo_regalos': 'Informacion adicional',
            'titulo_album_compartido': 'Material del evento',
            'texto_album_compartido': 'Accede al espacio compartido del evento cuando este disponible.',
            'titulo_rsvp': 'Confirma tu asistencia',
            'texto_rsvp': 'Tu confirmacion nos ayuda a preparar accesos, lugares y alimentos.',
        },
        'secciones': {
            'PORTADA': {'orden': 5, 'activa': True},
            'PADRES_PADRINOS': {'orden': 12, 'activa': False, 'titulo': 'Anfitriones'},
            'CUENTA_REGRESIVA': {'orden': 20, 'activa': True},
            'DETALLES': {'orden': 30, 'activa': True},
            'DRESS_CODE': {'orden': 35, 'activa': True, 'titulo': 'Codigo de vestimenta'},
            'ITINERARIO': {'orden': 40, 'activa': True, 'titulo': 'Agenda'},
            'ALBUM': {'orden': 50, 'activa': False},
            'MENU': {'orden': 60, 'activa': True},
            'REGALOS': {'orden': 70, 'activa': False},
            'ALBUM_COMPARTIDO': {'orden': 75, 'activa': False},
            'RSVP': {'orden': 80, 'activa': True},
        },
    },
}


def aplicar_plantilla_evento(evento, codigo):
    configuracion = CONFIGURACION_PLANTILLAS_EVENTO.get(codigo)
    if not configuracion:
        return

    for campo, valor in configuracion['campos'].items():
        setattr(evento, campo, valor)
    return

    plantillas = {
        'BODA': {
            'tipo_evento': 'BODA',
            'paleta_colores': 'TINTA_MARFIL',
            'paleta_sobre': 'TINTA_MARFIL',
            'estilo_letra': 'CURSIVA_ELEGANTE',
            'etiqueta_principal': 'Novia',
            'etiqueta_secundario': 'Novio',
            'frase_portada': 'Nuestra Boda',
            'titulo_detalles': 'Detalles de la celebración',
            'texto_detalles': 'Ceremonia, recepción, ubicación y dress code reunidos para acompañarnos sin complicaciones.',
            'titulo_regalos': 'Mesa de regalos',
            'titulo_rsvp': 'Confirma tu asistencia',
        },
        'XV': {
            'tipo_evento': 'XV',
            'paleta_colores': 'ROSA',
            'paleta_sobre': 'ROSA',
            'estilo_letra': 'ROMANTICA',
            'etiqueta_principal': 'Festejada',
            'etiqueta_secundario': 'Familia',
            'frase_portada': 'Mis XV años',
            'titulo_detalles': 'Detalles de la fiesta',
            'texto_detalles': 'Aquí encontrarás horarios, ubicación, dress code y todo lo necesario para celebrar juntos.',
            'titulo_regalos': 'Regalos',
            'titulo_rsvp': 'Confirma tu asistencia',
        },
        'BABY_SHOWER': {
            'tipo_evento': 'BABY_SHOWER',
            'paleta_colores': 'SALVIA_PERLA',
            'paleta_sobre': 'SALVIA_PERLA',
            'estilo_letra': 'MANUSCRITA',
            'etiqueta_principal': 'Bebé',
            'etiqueta_secundario': 'Familia',
            'frase_portada': 'Baby Shower',
            'titulo_detalles': 'Detalles del baby shower',
            'texto_detalles': 'Te compartimos horario, dirección, mesa de regalos y detalles para celebrar esta llegada especial.',
            'titulo_regalos': 'Mesa de regalos',
            'titulo_rsvp': 'Confirma tu asistencia',
        },
        'BAUTIZO': {
            'tipo_evento': 'BAUTIZO',
            'paleta_colores': 'DORADO',
            'paleta_sobre': 'DORADO',
            'estilo_letra': 'CLASICA',
            'etiqueta_principal': 'Bautizado',
            'etiqueta_secundario': 'Familia',
            'frase_portada': 'Mi Bautizo',
            'titulo_detalles': 'Ceremonia y celebración',
            'texto_detalles': 'Aquí tienes los horarios, ubicaciones y detalles para acompañarnos en este día especial.',
            'titulo_regalos': 'Detalles',
            'titulo_rsvp': 'Confirma tu asistencia',
        },
    }
    valores = plantillas.get(codigo)
    if not valores:
        return

    for campo, valor in valores.items():
        setattr(evento, campo, valor)


def aplicar_estructura_plantilla_evento(evento, codigo):
    configuracion = CONFIGURACION_PLANTILLAS_EVENTO.get(codigo)
    if not configuracion:
        return

    for tipo, ajustes in configuracion.get('secciones', {}).items():
        datos = datos_default_seccion(evento, tipo)
        defaults = {
            'titulo': ajustes.get('titulo', datos['titulo']),
            'descripcion': ajustes.get('descripcion', datos['descripcion']),
            'orden': ajustes.get('orden', datos['orden']),
            'activa': ajustes.get('activa', datos['activa']),
        }
        seccion, creada = SeccionInvitacion.objects.get_or_create(
            evento=evento,
            tipo=tipo,
            defaults=defaults,
        )
        if not creada:
            seccion.titulo = defaults['titulo']
            seccion.descripcion = defaults['descripcion']
            seccion.orden = defaults['orden']
            seccion.activa = defaults['activa']
            seccion.save(update_fields=['titulo', 'descripcion', 'orden', 'activa', 'fecha_actualizacion'])


def convertir_datetime_local(valor):
    if not valor:
        return None
    try:
        parsed = datetime.fromisoformat(valor)
    except (TypeError, ValueError):
        return None

    if timezone.is_naive(parsed):
        parsed = timezone.make_aware(
            parsed,
            timezone.get_current_timezone(),
        )
    return parsed


def archivo_pasa_validadores(modelo, campo, archivo):
    field = modelo._meta.get_field(campo)
    try:
        for validator in field.validators:
            validator(archivo)
    except ValidationError:
        return False
    return True


def guardar_campos_regalo(regalo, request):
    tipo = request.POST.get('tipo_regalo') or 'TIENDA'
    tiene_datos_bancarios = any(
        request.POST.get(campo)
        for campo in ['banco', 'titular', 'numero_cuenta', 'clabe']
    )
    if tiene_datos_bancarios and not request.POST.get('url_regalo'):
        tipo = 'DEPOSITO'

    regalo.tipo = tipo
    regalo.nombre = limpiar_texto(request, 'nombre_regalo') or ('Deposito bancario' if tipo == 'DEPOSITO' else 'Mesa de regalos')
    regalo.url = limpiar_texto(request, 'url_regalo')
    regalo.banco = limpiar_texto(request, 'banco')
    regalo.titular = limpiar_texto(request, 'titular')
    regalo.numero_cuenta = limpiar_texto(request, 'numero_cuenta')
    regalo.clabe = limpiar_texto(request, 'clabe')
    regalo.instrucciones = limpiar_texto(request, 'instrucciones')
    regalo.visible = bool_post(request, 'visible')
    regalo.save()


def agregar_persona_ceremonia_dashboard(request, evento):
    nombre = limpiar_texto(request, 'nombre_persona')
    if not nombre:
        return 'operacion'

    PersonaCeremonia.objects.create(
        evento=evento,
        seccion=request.POST.get('seccion_persona') or 'PADRINOS',
        etiqueta=limpiar_texto(request, 'etiqueta_persona') or 'Nombre',
        nombre=nombre,
        orden=convertir_entero(request.POST.get('orden_persona'), 0),
        visible=bool_post(request, 'visible_persona'),
    )
    return 'operacion'


def editar_persona_ceremonia_dashboard(request, evento):
    persona = get_object_or_404(PersonaCeremonia, id=request.POST.get('persona_id'), evento=evento)
    persona.seccion = request.POST.get('seccion_persona') or persona.seccion
    persona.etiqueta = limpiar_texto(request, 'etiqueta_persona') or persona.etiqueta
    persona.nombre = limpiar_texto(request, 'nombre_persona') or persona.nombre
    persona.orden = convertir_entero(request.POST.get('orden_persona'), persona.orden)
    persona.visible = bool_post(request, 'visible_persona')
    persona.save()
    return 'operacion'


def eliminar_persona_ceremonia_dashboard(request, evento):
    persona = get_object_or_404(PersonaCeremonia, id=request.POST.get('persona_id'), evento=evento)
    persona.delete()
    return 'operacion'


def editar_regalo_dashboard(request, evento):
    regalo = get_object_or_404(EnlaceRegalo, id=request.POST.get('regalo_id'), evento=evento)
    guardar_campos_regalo(regalo, request)
    return 'operacion'


def eliminar_regalo_dashboard(request, evento):
    regalo = get_object_or_404(EnlaceRegalo, id=request.POST.get('regalo_id'), evento=evento)
    regalo.delete()
    return 'operacion'


def editar_album_dashboard(request, evento):
    foto = get_object_or_404(FotoEvento, id=request.POST.get('foto_id'), evento=evento)
    foto.titulo = limpiar_texto(request, 'titulo_album_media')
    foto.orden = convertir_entero(request.POST.get('orden_album_media'), foto.orden)
    foto.visible = bool_post(request, 'visible_album_media')
    archivo = request.FILES.get('archivo_album')
    if archivo and archivo_pasa_validadores(FotoEvento, 'imagen', archivo):
        if foto.imagen:
            foto.imagen.delete(save=False)
        foto.imagen = archivo
    foto.save()
    return 'contenido'


def eliminar_album_dashboard(request, evento):
    foto = get_object_or_404(FotoEvento, id=request.POST.get('foto_id'), evento=evento)
    if foto.imagen:
        foto.imagen.delete(save=False)
    foto.delete()
    return 'contenido'


def proveedor_empresa_para_evento(evento, proveedor_id, *, user=None):
    """Resolve a provider inside the event tenant.

    Wedding Planners may only select providers explicitly visible to planners.
    DIRTEC/Admin Empresa keep access to the complete active tenant catalog.
    Supplying an invalid/foreign/hidden provider id is rejected instead of
    silently converting the relationship to ``None``.
    """
    if not proveedor_id:
        return None

    queryset = Proveedor.objects.filter(activo=True)
    if evento.empresa:
        queryset = queryset.filter(empresa=evento.empresa)

    if user and not usuario_es_dirtec(user):
        roles = roles_activos_empresa(user, evento.empresa) if evento.empresa else set()
        if 'WEDDING_PLANNER' in roles and 'ADMIN_EMPRESA' not in roles:
            queryset = queryset.filter(visible_para_wedding_planners=True)

    return get_object_or_404(queryset, id=proveedor_id)


def usuarios_empresa_para_evento(evento):
    if not evento.empresa:
        return get_user_model().objects.none()
    return get_user_model().objects.filter(
        membresias_empresa__empresa=evento.empresa,
        membresias_empresa__activo=True,
    ).distinct()


def categoria_gasto_dashboard(request):
    categoria_id = request.POST.get('categoria_gasto_id')
    nombre = limpiar_texto(request, 'categoria_gasto_nombre')
    if categoria_id:
        categoria = CategoriaGasto.objects.filter(id=categoria_id, activo=True).first()
        if categoria:
            return categoria
    if nombre:
        categoria, _ = CategoriaGasto.objects.get_or_create(nombre=nombre, defaults={'activo': True})
        return categoria
    categoria, _ = CategoriaGasto.objects.get_or_create(nombre='General', defaults={'activo': True})
    return categoria


def guardar_campos_servicio_evento(servicio, request, evento):
    estados_validos = {valor for valor, _ in ServicioEvento.ESTADOS}
    servicio.proveedor = proveedor_empresa_para_evento(evento, request.POST.get('proveedor_id'), user=request.user)
    servicio.nombre_servicio = limpiar_texto(request, 'nombre_servicio') or servicio.nombre_servicio or 'Servicio'
    servicio.descripcion = limpiar_texto(request, 'descripcion_servicio')
    servicio.fecha_servicio = fecha_dashboard(request, 'fecha_servicio', servicio.fecha_servicio)
    servicio.hora_inicio = hora_dashboard(request, 'hora_inicio_servicio', servicio.hora_inicio)
    servicio.hora_fin = hora_dashboard(request, 'hora_fin_servicio', servicio.hora_fin)
    servicio.lugar = limpiar_texto(request, 'lugar_servicio')
    servicio.costo_proveedor = convertir_decimal(request.POST.get('costo_proveedor_servicio'), servicio.costo_proveedor)
    estado = request.POST.get('estado_servicio')
    if estado in estados_validos:
        servicio.estado = estado
    servicio.notas = limpiar_texto(request, 'notas_servicio')
    for campo in ('contrato', 'cotizacion'):
        if request.POST.get(f'eliminar_{campo}_servicio') == 'on':
            archivo_actual = getattr(servicio, campo, None)
            if archivo_actual:
                archivo_actual.delete(save=False)
            setattr(servicio, campo, None)
        archivo = request.FILES.get(f'{campo}_servicio')
        if archivo and archivo_pasa_validadores(ServicioEvento, campo, archivo):
            archivo_actual = getattr(servicio, campo, None)
            if archivo_actual:
                archivo_actual.delete(save=False)
            setattr(servicio, campo, archivo)
    servicio.save()
    return servicio


def agregar_servicio_evento_dashboard(request, evento):
    servicio = ServicioEvento(evento=evento, nombre_servicio='Servicio')
    guardar_campos_servicio_evento(servicio, request, evento)
    return 'servicios'


def editar_servicio_evento_dashboard(request, evento):
    servicio = get_object_or_404(ServicioEvento, evento=evento, id=request.POST.get('servicio_id'))
    guardar_campos_servicio_evento(servicio, request, evento)
    return 'servicios'


def eliminar_servicio_evento_dashboard(request, evento):
    servicio = get_object_or_404(ServicioEvento, evento=evento, id=request.POST.get('servicio_id'))
    servicio.delete()
    return 'servicios'


def guardar_campos_personal_evento(personal, request, evento):
    tipos_validos = {valor for valor, _ in PersonalEvento.TIPOS}
    estados_validos = {valor for valor, _ in PersonalEvento.ESTADOS}
    personal.proveedor = proveedor_empresa_para_evento(evento, request.POST.get('proveedor_id'), user=request.user)
    personal.nombre = limpiar_texto(request, 'nombre_personal') or personal.nombre or 'Personal'
    if request.POST.get('tipo_personal') in tipos_validos:
        personal.tipo_personal = request.POST.get('tipo_personal')
    personal.telefono = limpiar_texto(request, 'telefono_personal')
    personal.hora_entrada = hora_dashboard(request, 'hora_entrada_personal', personal.hora_entrada)
    personal.hora_salida = hora_dashboard(request, 'hora_salida_personal', personal.hora_salida)
    personal.area_asignada = limpiar_texto(request, 'area_personal')
    personal.mesas_asignadas = limpiar_texto(request, 'mesas_personal')
    personal.costo = convertir_decimal(request.POST.get('costo_personal'), personal.costo)
    personal.uniforme = limpiar_texto(request, 'uniforme_personal')
    if request.POST.get('estado_personal') in estados_validos:
        personal.estado = request.POST.get('estado_personal')
    personal.notas = limpiar_texto(request, 'notas_personal')
    personal.save()
    return personal


def agregar_personal_evento_dashboard(request, evento):
    personal = PersonalEvento(evento=evento, nombre='Personal')
    guardar_campos_personal_evento(personal, request, evento)
    return 'operacion'


def editar_personal_evento_dashboard(request, evento):
    personal = get_object_or_404(PersonalEvento, evento=evento, id=request.POST.get('personal_id'))
    guardar_campos_personal_evento(personal, request, evento)
    return 'operacion'


def eliminar_personal_evento_dashboard(request, evento):
    personal = get_object_or_404(PersonalEvento, evento=evento, id=request.POST.get('personal_id'))
    personal.delete()
    return 'operacion'


def guardar_campos_tarea_evento(tarea, request, evento):
    estados_validos = {valor for valor, _ in TareaEvento.ESTADOS}
    prioridades_validas = {valor for valor, _ in TareaEvento.PRIORIDADES}
    categorias_validas = {valor for valor, _ in TareaEvento.CATEGORIAS}
    tarea.titulo = limpiar_texto(request, 'titulo_tarea') or tarea.titulo or 'Tarea'
    tarea.descripcion = limpiar_texto(request, 'descripcion_tarea')
    responsable_id = request.POST.get('responsable_id')
    tarea.responsable = usuarios_empresa_para_evento(evento).filter(id=responsable_id).first() if responsable_id else None
    servicio_tarea_id = request.POST.get('servicio_evento_id')
    tarea.servicio_evento = ServicioEvento.objects.filter(evento=evento, id=servicio_tarea_id).first() if servicio_tarea_id else None
    tarea.fecha_inicio = (
        fecha_dashboard(
            request,
            'fecha_inicio_tarea',
            tarea.fecha_inicio,
        )
        if request.POST.get(
            'fecha_inicio_tarea'
        )
        else None
    )
    # K.8.7.1.2: una tarea ya no representa citas ni rangos horarios.
    tarea.hora_inicio = None
    tarea.fecha_limite = (
        fecha_dashboard(
            request,
            'fecha_limite_tarea',
            tarea.fecha_limite,
        )
        if request.POST.get(
            'fecha_limite_tarea'
        )
        else None
    )
    tarea.hora_fin = None
    if request.POST.get('prioridad_tarea') in prioridades_validas:
        tarea.prioridad = request.POST.get('prioridad_tarea')
    if request.POST.get('estado_tarea') in estados_validos:
        tarea.estado = request.POST.get('estado_tarea')
    if request.POST.get('categoria_tarea') in categorias_validas:
        tarea.categoria = request.POST.get('categoria_tarea')
    tarea.porcentaje_avance = min(convertir_entero(request.POST.get('porcentaje_avance_tarea'), tarea.porcentaje_avance), 100)
    tarea.notas = limpiar_texto(request, 'notas_tarea')
    if request.POST.get('eliminar_evidencia_tarea') == 'on' and tarea.evidencia:
        tarea.evidencia.delete(save=False)
        tarea.evidencia = None
    evidencia = request.FILES.get('evidencia_tarea')
    if evidencia and archivo_pasa_validadores(TareaEvento, 'evidencia', evidencia):
        if tarea.evidencia:
            tarea.evidencia.delete(save=False)
        tarea.evidencia = evidencia
    tarea.save()
    return tarea


def agregar_tarea_evento_dashboard(request, evento):
    tarea = TareaEvento(evento=evento, titulo='Tarea')
    guardar_campos_tarea_evento(tarea, request, evento)
    return 'tareas'


def editar_tarea_evento_dashboard(request, evento):
    tarea = get_object_or_404(TareaEvento, evento=evento, id=request.POST.get('tarea_id'))
    guardar_campos_tarea_evento(tarea, request, evento)
    return 'tareas'


def eliminar_tarea_evento_dashboard(request, evento):
    tarea = get_object_or_404(TareaEvento, evento=evento, id=request.POST.get('tarea_id'))
    tarea.delete()
    return 'tareas'


def guardar_campos_gasto_evento(gasto, request, evento):
    estados_validos = {valor for valor, _ in GastoEvento.ESTADOS}
    gasto.categoria = categoria_gasto_dashboard(request)
    gasto.proveedor = proveedor_empresa_para_evento(evento, request.POST.get('proveedor_id'), user=request.user)
    gasto.concepto = limpiar_texto(request, 'concepto_gasto') or gasto.concepto or 'Gasto'
    gasto.monto_estimado = convertir_decimal(request.POST.get('monto_estimado_gasto'), gasto.monto_estimado)
    gasto.monto_real = convertir_decimal(request.POST.get('monto_real_gasto'), gasto.monto_real)
    gasto.fecha_limite = fecha_dashboard(request, 'fecha_limite_gasto', gasto.fecha_limite)
    if request.POST.get('estado_gasto') in estados_validos:
        gasto.estado = request.POST.get('estado_gasto')
    gasto.notas = limpiar_texto(request, 'notas_gasto')
    gasto.save()
    return gasto


def agregar_gasto_evento_dashboard(request, evento):
    gasto = GastoEvento(evento=evento, categoria=categoria_gasto_dashboard(request), concepto='Gasto')
    guardar_campos_gasto_evento(gasto, request, evento)
    return 'operacion'


def editar_gasto_evento_dashboard(request, evento):
    gasto = get_object_or_404(GastoEvento, evento=evento, id=request.POST.get('gasto_id'))
    guardar_campos_gasto_evento(gasto, request, evento)
    return 'operacion'


def eliminar_gasto_evento_dashboard(request, evento):
    gasto = get_object_or_404(GastoEvento, evento=evento, id=request.POST.get('gasto_id'))
    gasto.delete()
    return 'operacion'


def guardar_campos_pago_evento(pago, request):
    metodos_validos = {valor for valor, _ in PagoEvento.METODOS}
    pago.monto = convertir_decimal(request.POST.get('monto_pago'), pago.monto)
    pago.fecha_pago = fecha_dashboard(request, 'fecha_pago', pago.fecha_pago)
    if request.POST.get('metodo_pago') in metodos_validos:
        pago.metodo_pago = request.POST.get('metodo_pago')
    pago.referencia = limpiar_texto(request, 'referencia_pago')
    pago.notas = limpiar_texto(request, 'notas_pago')
    if request.POST.get('eliminar_comprobante_pago') == 'on' and pago.comprobante:
        pago.comprobante.delete(save=False)
        pago.comprobante = None
    comprobante = request.FILES.get('comprobante_pago')
    if comprobante and archivo_pasa_validadores(PagoEvento, 'comprobante', comprobante):
        if pago.comprobante:
            pago.comprobante.delete(save=False)
        pago.comprobante = comprobante
    pago.save()
    return pago


def agregar_pago_evento_dashboard(request, evento):
    gasto = get_object_or_404(GastoEvento, evento=evento, id=request.POST.get('gasto_id'))
    pago = PagoEvento(gasto=gasto, monto=0)
    guardar_campos_pago_evento(pago, request)
    return 'operacion'


def editar_pago_evento_dashboard(request, evento):
    pago = get_object_or_404(PagoEvento, gasto__evento=evento, id=request.POST.get('pago_id'))
    guardar_campos_pago_evento(pago, request)
    return 'operacion'


def eliminar_pago_evento_dashboard(request, evento):
    pago = get_object_or_404(PagoEvento, gasto__evento=evento, id=request.POST.get('pago_id'))
    pago.delete()
    return 'operacion'


def guardar_campos_documento_evento(documento, request, evento):
    tipos_validos = {valor for valor, _ in DocumentoEvento.TIPOS}
    if request.POST.get('tipo_documento') in tipos_validos:
        documento.tipo_documento = request.POST.get('tipo_documento')
    documento.titulo = limpiar_texto(request, 'titulo_documento') or documento.titulo or 'Documento'
    documento.descripcion = limpiar_texto(request, 'descripcion_documento')
    documento.proveedor = proveedor_empresa_para_evento(evento, request.POST.get('proveedor_id'), user=request.user)
    documento.visible_cliente = bool_post(request, 'visible_cliente_documento')
    archivo = request.FILES.get('archivo_documento')
    if archivo and archivo_pasa_validadores(DocumentoEvento, 'archivo', archivo):
        if documento.archivo:
            documento.archivo.delete(save=False)
        documento.archivo = archivo
    documento.save()
    return documento


def agregar_documento_evento_dashboard(request, evento):
    archivo = request.FILES.get('archivo_documento')
    if not archivo or not archivo_pasa_validadores(DocumentoEvento, 'archivo', archivo):
        return 'operacion'
    documento = DocumentoEvento(evento=evento, titulo='Documento', cargado_por=request.user)
    guardar_campos_documento_evento(documento, request, evento)
    return 'operacion'


def editar_documento_evento_dashboard(request, evento):
    documento = get_object_or_404(DocumentoEvento, evento=evento, id=request.POST.get('documento_id'))
    guardar_campos_documento_evento(documento, request, evento)
    return 'operacion'


def eliminar_documento_evento_dashboard(request, evento):
    documento = get_object_or_404(DocumentoEvento, evento=evento, id=request.POST.get('documento_id'))
    if documento.archivo:
        documento.archivo.delete(save=False)
    documento.delete()
    return 'operacion'


def guardar_campos_aprobacion_evento(aprobacion, request):
    tipos_validos = {valor for valor, _ in AprobacionEvento.TIPOS}
    estados_validos = {valor for valor, _ in AprobacionEvento.ESTADOS}
    if request.POST.get('tipo_aprobacion') in tipos_validos:
        aprobacion.tipo = request.POST.get('tipo_aprobacion')
    aprobacion.titulo = limpiar_texto(request, 'titulo_aprobacion') or aprobacion.titulo or 'Aprobacion'
    aprobacion.descripcion = limpiar_texto(request, 'descripcion_aprobacion')
    if request.POST.get('estado_aprobacion') in estados_validos:
        aprobacion.estado = request.POST.get('estado_aprobacion')
    aprobacion.comentario = limpiar_texto(request, 'comentario_aprobacion')
    aprobacion.save()
    return aprobacion


def agregar_aprobacion_evento_dashboard(request, evento):
    aprobacion = AprobacionEvento(evento=evento, tipo='OTRO', titulo='Aprobacion', solicitado_por=request.user)
    guardar_campos_aprobacion_evento(aprobacion, request)
    return 'operacion'


def editar_aprobacion_evento_dashboard(request, evento):
    aprobacion = get_object_or_404(AprobacionEvento, evento=evento, id=request.POST.get('aprobacion_id'))
    guardar_campos_aprobacion_evento(aprobacion, request)
    return 'operacion'


def eliminar_aprobacion_evento_dashboard(request, evento):
    aprobacion = get_object_or_404(AprobacionEvento, evento=evento, id=request.POST.get('aprobacion_id'))
    aprobacion.delete()
    return 'operacion'


def agregar_grupo_dashboard(request, evento):
    nombre = limpiar_texto(request, 'nombre_grupo')
    tipo = request.POST.get('tipo_grupo') or 'PERSONAL'
    destino = request.POST.get('destino') if request.POST.get('destino') in {'personalizacion', 'invitados'} else 'invitados'
    if not nombre:
        return destino

    permitir_extras = bool_post(request, 'permitir_acompanantes_extra')
    cantidad_extras = (
        max(convertir_entero(request.POST.get('cantidad_extra_permitida'), 0), 0)
        if permitir_extras
        else 0
    )
    grupo = Grupoinvitacion.objects.create(
        evento=evento,
        nombre_grupo=nombre,
        tipo=tipo,
        cantidad_extra_permitida=cantidad_extras,
        cantidad_maxima=1,
        telefono_contacto=limpiar_texto(request, 'telefono_contacto'),
        correo_contacto=limpiar_texto(request, 'correo_contacto'),
        permitir_acompanantes_extra=permitir_extras,
    )
    if grupo.es_familiar:
        crear_invitados_desde_textarea(grupo, request.POST.get('invitados_familia', ''))
    asegurar_roster_grupo(grupo)
    return destino


def editar_grupo_dashboard(request, evento):
    grupo = get_object_or_404(Grupoinvitacion, id=request.POST.get('grupo_id'), evento=evento)
    grupo.nombre_grupo = limpiar_texto(request, 'nombre_grupo') or grupo.nombre_grupo
    permitir_extras = bool_post(request, 'permitir_acompanantes_extra')
    grupo.cantidad_extra_permitida = (
        max(convertir_entero(request.POST.get('cantidad_extra_permitida'), 0), 0)
        if permitir_extras
        else 0
    )
    grupo.telefono_contacto = limpiar_texto(request, 'telefono_contacto')
    grupo.correo_contacto = limpiar_texto(request, 'correo_contacto')
    grupo.permitir_acompanantes_extra = permitir_extras
    grupo.save()
    asegurar_roster_grupo(grupo)
    return 'invitados'


def eliminar_grupo_dashboard(request, evento):
    grupo = get_object_or_404(Grupoinvitacion, id=request.POST.get('grupo_id'), evento=evento)
    grupo.delete()
    return 'invitados'


def crear_invitados_desde_textarea(grupo, texto):
    for orden, linea in enumerate(texto.splitlines(), 1):
        partes = [parte.strip() for parte in linea.split('|')]
        nombre = partes[0] if partes else ''
        if not nombre:
            continue
        tipo_persona = 'NINO' if len(partes) > 1 and partes[1].upper().startswith('N') else 'ADULTO'
        Invitado.objects.create(
            grupo=grupo,
            nombre=nombre,
            tipo_persona=tipo_persona,
            orden=orden,
        )


def agregar_invitado_dashboard(request, evento):
    grupo = get_object_or_404(Grupoinvitacion, id=request.POST.get('grupo_id'), evento=evento, tipo='FAMILIAR')
    nombre = limpiar_texto(request, 'nombre_invitado')
    if not nombre:
        return 'invitados'

    Invitado.objects.create(
        grupo=grupo,
        nombre=nombre,
        tipo_persona=request.POST.get('tipo_persona') or 'ADULTO',
        menu_asignado=request.POST.get('menu_asignado') or 'SEGUN_TIPO',
        orden=convertir_entero(request.POST.get('orden_invitado'), 0),
    )
    return 'invitados'


def editar_invitado_dashboard(request, evento):
    invitado = get_object_or_404(Invitado, id=request.POST.get('invitado_id'), grupo__evento=evento)
    invitado.nombre = limpiar_texto(request, 'nombre_invitado') or invitado.nombre
    invitado.tipo_persona = request.POST.get('tipo_persona') or invitado.tipo_persona
    invitado.orden = convertir_entero(request.POST.get('orden_invitado'), invitado.orden)
    invitado.menu_asignado = request.POST.get('menu_asignado') or invitado.menu_asignado
    invitado.save()
    return 'invitados'


def eliminar_invitado_dashboard(request, evento):
    invitado = get_object_or_404(Invitado, id=request.POST.get('invitado_id'), grupo__evento=evento)
    if invitado.es_acompanante_extra:
        return 'invitados'
    if invitado.grupo.es_personal:
        return 'invitados'
    invitado.delete()
    return 'invitados'


def empresa_operativa_dashboard(request, evento):
    empresa, _ = empresa_actual_dashboard(request)
    return evento.empresa or empresa


def agregar_sede_dashboard(request, evento):
    empresa = empresa_operativa_dashboard(request, evento)
    nombre = limpiar_texto(request, 'nombre_sede')
    if not empresa or not nombre:
        return 'operacion'

    SedeEvento.objects.create(
        empresa=empresa,
        nombre=nombre,
        tipo=request.POST.get('tipo_sede') or 'SALON',
        capacidad_minima=convertir_entero(request.POST.get('capacidad_minima'), 0),
        capacidad_maxima=convertir_entero(request.POST.get('capacidad_maxima'), 0),
        precio_base=convertir_decimal(request.POST.get('precio_base_sede'), 0),
        direccion=limpiar_texto(request, 'direccion_sede'),
        descripcion=limpiar_texto(request, 'descripcion_sede'),
        activa=True,
    )
    return 'operacion'


def agregar_proveedor_empresa_dashboard(request, evento):
    empresa = empresa_operativa_dashboard(request, evento)
    nombre = limpiar_texto(request, 'nombre_proveedor')
    if not empresa or not nombre:
        return 'operacion'

    Proveedor.objects.create(
        empresa=empresa,
        nombre_comercial=nombre,
        tipo_proveedor=request.POST.get('tipo_proveedor') or 'OTROS',
        nombre_contacto=limpiar_texto(request, 'contacto_proveedor'),
        telefono=limpiar_texto(request, 'telefono_proveedor'),
        correo=limpiar_texto(request, 'correo_proveedor'),
        contacto_operativo=limpiar_texto(request, 'contacto_operativo_proveedor'),
        telefono_operativo=limpiar_texto(request, 'telefono_operativo_proveedor'),
        correo_operativo=limpiar_texto(request, 'correo_operativo_proveedor'),
        visible_para_wedding_planners=request.POST.get('visible_wedding_planners_proveedor', 'on') == 'on',
        rfc=limpiar_texto(request, 'rfc_proveedor'),
        datos_bancarios=limpiar_texto(request, 'datos_bancarios_proveedor'),
        notas_privadas=limpiar_texto(request, 'notas_privadas_proveedor'),
        descripcion=limpiar_texto(request, 'descripcion_proveedor'),
        activo=True,
    )
    return 'operacion'


def agregar_paquete_empresa_dashboard(request, evento):
    empresa = empresa_operativa_dashboard(request, evento)
    nombre = limpiar_texto(request, 'nombre_paquete')
    if not empresa or not nombre:
        return 'operacion'

    PaqueteBoda.objects.create(
        empresa=empresa,
        nombre=nombre,
        descripcion=limpiar_texto(request, 'descripcion_paquete'),
        precio_base=convertir_decimal(request.POST.get('precio_base_paquete'), 0),
        numero_personas_incluidas=convertir_entero(request.POST.get('personas_paquete'), 0),
        activo=True,
    )
    return 'operacion'


def editar_sede_empresa_avanzado_dashboard(request, evento):
    empresa = empresa_operativa_dashboard(request, evento)
    if empresa:
        actualizar_sede_empresa_dashboard(request, empresa)
    return 'contenido'


def eliminar_sede_empresa_avanzado_dashboard(request, evento):
    empresa = empresa_operativa_dashboard(request, evento)
    if empresa:
        eliminar_sede_empresa_dashboard(request, empresa)
    return 'contenido'


def editar_proveedor_empresa_avanzado_dashboard(request, evento):
    empresa = empresa_operativa_dashboard(request, evento)
    if empresa:
        actualizar_proveedor_empresa_dashboard(request, empresa)
    return 'contenido'


def eliminar_proveedor_empresa_avanzado_dashboard(request, evento):
    empresa = empresa_operativa_dashboard(request, evento)
    if empresa:
        eliminar_proveedor_empresa_dashboard(request, empresa)
    return 'contenido'


def editar_paquete_empresa_avanzado_dashboard(request, evento):
    empresa = empresa_operativa_dashboard(request, evento)
    if empresa:
        actualizar_paquete_empresa_dashboard(request, empresa)
    return 'contenido'


def eliminar_paquete_empresa_avanzado_dashboard(request, evento):
    empresa = empresa_operativa_dashboard(request, evento)
    if empresa:
        eliminar_paquete_empresa_dashboard(request, empresa)
    return 'contenido'


def crear_usuario_empresa_dashboard(request, evento):
    empresa = empresa_operativa_dashboard(request, evento)
    username = limpiar_texto(request, 'username_usuario')
    if not empresa or not username:
        return 'usuarios'

    usuarios_activos = empresa.membresias.filter(activo=True).values('usuario_id').distinct().count()
    if usuarios_activos >= empresa.max_usuarios:
        return 'usuarios'

    crear_o_actualizar_usuario_empresa(request, empresa, 'WEDDING_PLANNER')
    return 'usuarios'


def editar_membresia_empresa_dashboard(request, evento):
    empresa = empresa_operativa_dashboard(request, evento)
    membresia = get_object_or_404(
        MembresiaEmpresa.objects.select_related('usuario'),
        id=request.POST.get('membresia_id'),
        empresa=empresa,
    )
    roles_validos = {valor for valor, _ in MembresiaEmpresa.ROLES}
    rol = request.POST.get('rol_usuario') if request.POST.get('rol_usuario') in roles_validos else membresia.rol
    user = membresia.usuario
    user.first_name = limpiar_texto(request, 'first_name_usuario')
    user.last_name = limpiar_texto(request, 'last_name_usuario')
    user.email = limpiar_texto(request, 'email_usuario')
    user.is_active = bool_post(request, 'activo_usuario')
    password = request.POST.get('password_usuario')
    if password:
        user.set_password(password)
    user.save()

    datos_membresia = {
        'activo': bool_post(request, 'activo_usuario'),
        'puede_gestionar_catalogos': bool_post(request, 'puede_gestionar_catalogos_usuario'),
    }
    if rol != membresia.rol:
        nueva_membresia, _ = MembresiaEmpresa.objects.update_or_create(
            empresa=empresa,
            usuario=user,
            rol=rol,
            defaults=datos_membresia,
        )
        MembresiaEmpresa.objects.filter(empresa=empresa, usuario=user).exclude(id=nueva_membresia.id).update(activo=False)
    else:
        for campo, valor in datos_membresia.items():
            setattr(membresia, campo, valor)
        membresia.save()
    return 'usuarios'


def desactivar_membresia_empresa_dashboard(request, evento):
    empresa = empresa_operativa_dashboard(request, evento)
    membresia = get_object_or_404(MembresiaEmpresa, id=request.POST.get('membresia_id'), empresa=empresa)
    if membresia.usuario_id == request.user.id:
        messages.error(request, 'No puedes desactivar tu propio acceso desde este panel.')
        return 'usuarios'
    membresia.activo = False
    membresia.save(update_fields=['activo'])
    if not MembresiaEmpresa.objects.filter(usuario=membresia.usuario, activo=True).exists():
        membresia.usuario.is_active = False
        membresia.usuario.save(update_fields=['is_active'])
    return 'usuarios'


def eliminar_membresia_empresa_dashboard(request, evento):
    empresa = empresa_operativa_dashboard(request, evento)
    membresia = get_object_or_404(
        MembresiaEmpresa.objects.select_related('usuario'),
        id=request.POST.get('membresia_id'),
        empresa=empresa,
    )
    if membresia.usuario_id == request.user.id:
        messages.error(request, 'No puedes eliminar tu propio acceso desde este panel.')
        return 'usuarios'

    user = membresia.usuario
    username = user.username
    membresia_id = membresia.id
    membresia.delete()
    if not user.is_staff and not user.is_superuser and not MembresiaEmpresa.objects.filter(usuario=user).exists():
        user.delete()

    registrar_auditoria(
        usuario=request.user,
        empresa=empresa,
        evento=evento,
        accion='ELIMINAR_MEMBRESIA_EMPRESA',
        modelo='MembresiaEmpresa',
        objeto_id=membresia_id,
        descripcion=f'Se elimino el acceso de {username} desde el dashboard avanzado.',
        request=request,
    )
    messages.success(request, f'Acceso eliminado: {username}.')
    return 'usuarios'


def asignar_planner_evento_dashboard(request, evento):
    empresa = empresa_operativa_dashboard(request, evento)
    planner_id = request.POST.get('planner_id')
    if not empresa:
        return 'usuarios'

    if not planner_id:
        evento.wedding_planner = None
        evento.save(update_fields=['wedding_planner'])
        return 'usuarios'

    membresia = get_object_or_404(
        MembresiaEmpresa,
        empresa=empresa,
        usuario_id=planner_id,
        rol='WEDDING_PLANNER',
        activo=True,
    )
    evento.wedding_planner = membresia.usuario
    if not evento.empresa:
        evento.empresa = empresa
        evento.save(update_fields=['wedding_planner', 'empresa'])
    else:
        evento.save(update_fields=['wedding_planner'])
    return 'usuarios'


def agregar_actividad_agenda_v3_dashboard(request, evento):
    from itinerario.services import asegurar_participantes_cita

    titulo = limpiar_texto(request, 'titulo_agenda')
    fecha = fecha_dashboard(request, 'fecha_agenda')
    hora_inicio = hora_dashboard(request, 'hora_inicio_agenda')
    if not titulo or not fecha or not hora_inicio:
        messages.error(request, 'Agenda: título, fecha y hora de inicio son obligatorios.')
        return 'agenda'

    tipo = request.POST.get('tipo_agenda')
    if tipo not in dict(ActividadItinerario.TIPOS):
        tipo = 'CITA'
    categoria = request.POST.get('categoria_agenda')
    if categoria not in dict(ActividadItinerario.CATEGORIAS):
        categoria = 'OTRO'
    prioridad = request.POST.get('prioridad_agenda')
    if prioridad not in dict(ActividadItinerario.PRIORIDADES):
        prioridad = 'MEDIA'
    servicio_id = request.POST.get('servicio_evento_id')
    servicio = ServicioEvento.objects.filter(evento=evento, id=servicio_id).select_related('proveedor').first() if servicio_id else None

    actividad = ActividadItinerario(
        evento=evento,
        servicio_evento=servicio,
        tipo=tipo,
        titulo=titulo,
        descripcion=limpiar_texto(request, 'descripcion_agenda'),
        categoria=categoria,
        fecha=fecha,
        hora_inicio=hora_inicio,
        hora_fin=hora_dashboard(request, 'hora_fin_agenda'),
        ubicacion=limpiar_texto(request, 'ubicacion_agenda'),
        responsable=request.user,
        proveedor=servicio.proveedor if servicio else None,
        prioridad=prioridad,
        estado='PENDIENTE',
        notas=limpiar_texto(request, 'notas_agenda'),
    )
    try:
        actividad.full_clean()
        actividad.save()
        if actividad.tipo == 'CITA':
            asegurar_participantes_cita(actividad, creador=request.user)
        messages.success(request, f'{actividad.get_tipo_display()} agregada a la agenda.')
    except ValidationError as exc:
        messages.error(request, 'No se pudo crear el elemento de agenda: ' + ' · '.join(exc.messages))
    return 'agenda'


ACCIONES_DASHBOARD_AVANZADO = {
    'agregar_actividad_agenda_v3': agregar_actividad_agenda_v3_dashboard,
    'agregar_persona_ceremonia': agregar_persona_ceremonia_dashboard,
    'editar_persona_ceremonia': editar_persona_ceremonia_dashboard,
    'eliminar_persona_ceremonia': eliminar_persona_ceremonia_dashboard,
    'editar_regalo': editar_regalo_dashboard,
    'eliminar_regalo': eliminar_regalo_dashboard,
    'editar_album': editar_album_dashboard,
    'eliminar_album': eliminar_album_dashboard,
    'agregar_servicio_evento': agregar_servicio_evento_dashboard,
    'editar_servicio_evento': editar_servicio_evento_dashboard,
    'eliminar_servicio_evento': eliminar_servicio_evento_dashboard,
    'agregar_personal_evento': agregar_personal_evento_dashboard,
    'editar_personal_evento': editar_personal_evento_dashboard,
    'eliminar_personal_evento': eliminar_personal_evento_dashboard,
    'agregar_tarea_evento': agregar_tarea_evento_dashboard,
    'editar_tarea_evento': editar_tarea_evento_dashboard,
    'eliminar_tarea_evento': eliminar_tarea_evento_dashboard,
    'agregar_gasto_evento': agregar_gasto_evento_dashboard,
    'editar_gasto_evento': editar_gasto_evento_dashboard,
    'eliminar_gasto_evento': eliminar_gasto_evento_dashboard,
    'agregar_pago_evento': agregar_pago_evento_dashboard,
    'editar_pago_evento': editar_pago_evento_dashboard,
    'eliminar_pago_evento': eliminar_pago_evento_dashboard,
    'agregar_documento_evento': agregar_documento_evento_dashboard,
    'editar_documento_evento': editar_documento_evento_dashboard,
    'eliminar_documento_evento': eliminar_documento_evento_dashboard,
    'agregar_aprobacion_evento': agregar_aprobacion_evento_dashboard,
    'editar_aprobacion_evento': editar_aprobacion_evento_dashboard,
    'eliminar_aprobacion_evento': eliminar_aprobacion_evento_dashboard,
    'agregar_grupo': agregar_grupo_dashboard,
    'editar_grupo': editar_grupo_dashboard,
    'eliminar_grupo': eliminar_grupo_dashboard,
    'agregar_invitado': agregar_invitado_dashboard,
    'editar_invitado': editar_invitado_dashboard,
    'eliminar_invitado': eliminar_invitado_dashboard,
    'agregar_sede_empresa': agregar_sede_dashboard,
    'editar_sede_empresa': editar_sede_empresa_avanzado_dashboard,
    'eliminar_sede_empresa': eliminar_sede_empresa_avanzado_dashboard,
    'agregar_proveedor_empresa': agregar_proveedor_empresa_dashboard,
    'editar_proveedor_empresa': editar_proveedor_empresa_avanzado_dashboard,
    'eliminar_proveedor_empresa': eliminar_proveedor_empresa_avanzado_dashboard,
    'agregar_paquete_empresa': agregar_paquete_empresa_dashboard,
    'editar_paquete_empresa': editar_paquete_empresa_avanzado_dashboard,
    'eliminar_paquete_empresa': eliminar_paquete_empresa_avanzado_dashboard,
    'crear_usuario_empresa': crear_usuario_empresa_dashboard,
    'editar_membresia_empresa': editar_membresia_empresa_dashboard,
    'desactivar_membresia_empresa': desactivar_membresia_empresa_dashboard,
    'eliminar_membresia_empresa': eliminar_membresia_empresa_dashboard,
    'asignar_planner_evento': asignar_planner_evento_dashboard,
}

ACCIONES_CATALOGOS_EMPRESA = {
    'agregar_sede_empresa',
    'editar_sede_empresa',
    'eliminar_sede_empresa',
    'agregar_proveedor_empresa',
    'editar_proveedor_empresa',
    'eliminar_proveedor_empresa',
    'agregar_paquete_empresa',
    'editar_paquete_empresa',
    'eliminar_paquete_empresa',
}

ACCIONES_USUARIOS_EMPRESA = {
    'crear_usuario_empresa',
    'editar_membresia_empresa',
    'desactivar_membresia_empresa',
    'eliminar_membresia_empresa',
    'asignar_planner_evento',
}


@login_required(login_url=LOGIN_DASHBOARD_URL)
def dashboard(request):
    evento, eventos = obtener_evento_dashboard(request)
    empresa_dashboard, empresas_dashboard = empresa_actual_dashboard(request)
    if not empresa_dashboard and evento and evento.empresa:
        empresa_dashboard = evento.empresa
    membresia_actual = membresia_empresa_usuario(request.user, empresa_dashboard)
    roles_dashboard = roles_activos_empresa(request.user, empresa_dashboard) if empresa_dashboard else set()
    roles_internos_dashboard = {'ADMIN_EMPRESA', 'WEDDING_PLANNER', 'VENTAS'}
    if evento and not usuario_es_dirtec(request.user) and not roles_dashboard.intersection(roles_internos_dashboard):
        if 'PROVEEDOR' in roles_dashboard and proveedor_de_usuario(request.user):
            return redirect(f"{reverse('portal_proveedor')}?evento={evento.id}")
        if 'CLIENTE' in roles_dashboard or evento.clientes.filter(id=request.user.id).exists():
            return redirect(f"{reverse('portal_cliente')}?evento={evento.id}")
    puede_gestionar_catalogos = usuario_puede_gestionar_catalogos(request.user, empresa_dashboard)
    puede_gestionar_usuarios = usuario_puede_gestionar_usuarios(request.user, empresa_dashboard)

    if request.method == 'POST' and request.POST.get('accion') in ACCIONES_DASHBOARD_AVANZADO:
        accion = request.POST.get('accion')
        evento = get_object_or_404(EventoBoda, id=request.POST.get('evento_id'))
        exigir_permiso_evento(request, evento, Actions.EVENT_EDIT)
        empresa_accion = empresa_operativa_dashboard(request, evento)
        if accion in ACCIONES_CATALOGOS_EMPRESA and not usuario_puede_gestionar_catalogos(request.user, empresa_accion):
            bloquear_accion_dashboard(request, empresa=empresa_accion, evento=evento, accion=accion, permiso='gestionar catalogos')
        if accion in ACCIONES_USUARIOS_EMPRESA and not usuario_puede_gestionar_usuarios(request.user, empresa_accion):
            bloquear_accion_dashboard(request, empresa=empresa_accion, evento=evento, accion=accion, permiso='gestionar usuarios')
        destino = ACCIONES_DASHBOARD_AVANZADO[accion](request, evento)
        return redirect(f'/dashboard/?evento={evento.id}#{destino}')

    if request.method == 'POST' and request.POST.get('accion') == 'personalizar_evento':
        evento = get_object_or_404(EventoBoda, id=request.POST.get('evento_id'))
        exigir_permiso_evento(request, evento, Actions.EVENT_EDIT)
        guardar_personalizacion_evento(request, evento)
        return redirect(f'/dashboard/?evento={evento.id}#invitacion')

    if request.method == 'POST' and request.POST.get('accion') == 'agregar_regalo':
        evento = get_object_or_404(EventoBoda, id=request.POST.get('evento_id'))
        exigir_permiso_evento(request, evento, Actions.EVENT_EDIT)
        agregar_regalo_dashboard(request, evento)
        return redirect(f'/dashboard/?evento={evento.id}#contenido')

    if request.method == 'POST' and request.POST.get('accion') == 'agregar_album':
        evento = get_object_or_404(EventoBoda, id=request.POST.get('evento_id'))
        exigir_permiso_evento(request, evento, Actions.EVENT_EDIT)
        agregar_album_dashboard(request, evento)
        return redirect(f'/dashboard/?evento={evento.id}#contenido')

    if request.method == 'POST' and request.POST.get('accion') == 'agregar_itinerario':
        evento = get_object_or_404(EventoBoda, id=request.POST.get('evento_id'))
        exigir_permiso_evento(request, evento, Actions.EVENT_EDIT)
        agregar_itinerario_dashboard(request, evento)
        return redirect(f'/dashboard/?evento={evento.id}#contenido')

    grupos = list(
        Grupoinvitacion.objects.filter(evento=evento)
        .select_related('evento')
        .prefetch_related(
            'invitados__asignaciones_mesa__mesa',
        )
        .order_by('tipo', 'nombre_grupo')
    ) if evento else []

    total_grupos = len(grupos)
    invitados_metricas = resumen_invitados_evento(
        evento,
        incluir_mesas=True,
    )
    total_lugares = invitados_metricas['total_personas']
    total_asistiran = invitados_metricas['confirmados']
    total_no_asistiran = invitados_metricas['no_asisten']
    total_pendientes = invitados_metricas['pendientes']
    total_adultos = invitados_metricas['buffet_adultos']
    total_ninos = invitados_metricas['buffet_infantiles']
    total_adultos_persona = invitados_metricas['adultos_confirmados_persona']
    total_ninos_persona = invitados_metricas['ninos_confirmados_persona']
    porcentaje_asistencia = invitados_metricas['porcentaje_asistencia']

    servicios_evento = ServicioEvento.objects.filter(evento=evento).select_related('proveedor') if evento else ServicioEvento.objects.none()
    personal_evento = PersonalEvento.objects.filter(evento=evento).select_related('proveedor') if evento else PersonalEvento.objects.none()
    actividades_evento = ActividadItinerario.objects.filter(evento=evento).select_related('servicio_evento', 'proveedor', 'responsable') if evento else ActividadItinerario.objects.none()
    mesas_evento = Mesa.objects.filter(evento=evento) if evento else Mesa.objects.none()
    asignaciones_mesa = AsignacionMesa.objects.filter(mesa__evento=evento) if evento else AsignacionMesa.objects.none()
    gastos_evento = GastoEvento.objects.filter(evento=evento).select_related('categoria', 'proveedor', 'servicio_evento').prefetch_related('pagos') if evento else GastoEvento.objects.none()
    pagos_evento = PagoEvento.objects.filter(gasto__evento=evento).select_related('gasto') if evento else PagoEvento.objects.none()
    tareas_evento = TareaEvento.objects.filter(evento=evento).select_related('responsable', 'servicio_evento') if evento else TareaEvento.objects.none()
    documentos_evento = DocumentoEvento.objects.filter(evento=evento).select_related('proveedor', 'cargado_por', 'servicio_evento') if evento else DocumentoEvento.objects.none()
    aprobaciones_evento = AprobacionEvento.objects.filter(evento=evento).select_related('solicitado_por', 'aprobado_por') if evento else AprobacionEvento.objects.none()
    notificaciones_evento = Notificacion.objects.filter(evento=evento) if evento else Notificacion.objects.none()
    usuarios_empresa = (
        MembresiaEmpresa.objects.filter(empresa=empresa_dashboard)
        .select_related('usuario')
        .order_by('-activo', 'rol', 'usuario__username')
    ) if empresa_dashboard and puede_gestionar_usuarios else MembresiaEmpresa.objects.none()
    planners_empresa = usuarios_empresa.filter(rol='WEDDING_PLANNER', activo=True) if empresa_dashboard and puede_gestionar_usuarios else MembresiaEmpresa.objects.none()
    responsables_evento = usuarios_empresa_para_evento(evento) if evento else get_user_model().objects.none()

    proveedores_empresa_dashboard = (
        empresa_dashboard.proveedores.filter(activo=True)
        if empresa_dashboard else Proveedor.objects.none()
    )
    if (
        empresa_dashboard
        and not usuario_es_dirtec(request.user)
        and 'WEDDING_PLANNER' in roles_dashboard
        and 'ADMIN_EMPRESA' not in roles_dashboard
    ):
        proveedores_empresa_dashboard = proveedores_empresa_dashboard.filter(
            visible_para_wedding_planners=True
        )

    total_proveedores_evento = servicios_evento.values('proveedor').distinct().count()
    total_servicios_evento = servicios_evento.count()
    proveedores_pendientes = servicios_evento.exclude(
        estado__in=['CONTRATADO', 'ANTICIPO_PAGADO', 'LIQUIDADO', 'SERVICIO_COMPLETADO']
    ).count()
    total_personal_evento = personal_evento.count()
    total_actividades_evento = actividades_evento.count()
    actividades_pendientes = actividades_evento.exclude(estado__in=['COMPLETADA', 'CANCELADA']).count()
    total_mesas_evento = mesas_evento.count()
    mesas_excedidas = sum(1 for mesa in mesas_evento if mesa.excedida)
    invitados_confirmados_sin_mesa = invitados_metricas['confirmados_sin_mesa']
    total_gastos_evento = gastos_evento.count()
    total_gasto_estimado = gastos_evento.aggregate(total=Sum('monto_estimado'))['total'] or 0
    total_gasto_real = gastos_evento.aggregate(total=Sum('monto_real'))['total'] or 0
    total_pagos_evento = pagos_evento.aggregate(total=Sum('monto'))['total'] or 0
    gastos_vencidos = sum(1 for gasto in gastos_evento if gasto.esta_vencido)
    total_tareas_evento = tareas_evento.count()
    tareas_pendientes = tareas_evento.exclude(estado__in=['COMPLETADA', 'CANCELADA']).count()
    tareas_vencidas = sum(1 for tarea in tareas_evento if tarea.esta_vencida)
    total_documentos_evento = documentos_evento.count()
    documentos_cliente = documentos_evento.filter(visible_cliente=True).count()
    aprobaciones_pendientes = aprobaciones_evento.filter(estado='PENDIENTE').count()
    total_aprobaciones_evento = aprobaciones_evento.count()
    notificaciones_no_leidas = notificaciones_evento.filter(leida=False).count()

    base_url = request.build_absolute_uri('/')[:-1]
    for grupo in grupos:
        preparar_links_grupo(grupo, base_url)

    context = {
        'evento': evento,
        'eventos': eventos,
        'empresa_dashboard': empresa_dashboard,
        'empresas_dashboard': empresas_dashboard,
        'membresia_actual': membresia_actual,
        'es_dirtec': usuario_es_dirtec(request.user),
        'puede_gestionar_catalogos': puede_gestionar_catalogos,
        'puede_gestionar_usuarios': puede_gestionar_usuarios,
        'puede_ver_finanzas_dashboard_v3': bool(
            usuario_es_dirtec(request.user)
            or roles_dashboard.intersection({'ADMIN_EMPRESA', 'WEDDING_PLANNER', 'VENTAS'})
        ),
        'usuarios_empresa': usuarios_empresa,
        'planners_empresa': planners_empresa,
        'roles_empresa': MembresiaEmpresa.ROLES,
        'sedes_empresa': empresa_dashboard.sedes.filter(activa=True) if empresa_dashboard else [],
        'proveedores_empresa': proveedores_empresa_dashboard,
        'paquetes_empresa': empresa_dashboard.paquetes.filter(activo=True) if empresa_dashboard else [],
        'tipos_sede': SedeEvento.TIPOS,
        'tipos_proveedor': Proveedor.TIPOS,
        'catalogo_sedes_empresa': empresa_dashboard.sedes.all() if empresa_dashboard and puede_gestionar_catalogos else [],
        'catalogo_proveedores_empresa': empresa_dashboard.proveedores.all() if empresa_dashboard and puede_gestionar_catalogos else [],
        'catalogo_paquetes_empresa': empresa_dashboard.paquetes.all() if empresa_dashboard and puede_gestionar_catalogos else [],
        'servicios_evento': servicios_evento,
        'personal_evento': personal_evento,
        'tareas_evento': tareas_evento,
        'gastos_evento': gastos_evento,
        'pagos_evento': pagos_evento,
        'documentos_evento': documentos_evento,
        'aprobaciones_evento': aprobaciones_evento,
        'responsables_evento': responsables_evento,
        'categorias_gasto': CategoriaGasto.objects.filter(activo=True).order_by('nombre'),
        'estados_servicio': ServicioEvento.ESTADOS,
        'tipos_personal': PersonalEvento.TIPOS,
        'estados_personal': PersonalEvento.ESTADOS,
        'estados_tarea': TareaEvento.ESTADOS,
        'prioridades_tarea': TareaEvento.PRIORIDADES,
        'categorias_tarea': TareaEvento.CATEGORIAS,
        'estados_gasto': GastoEvento.ESTADOS,
        'metodos_pago': PagoEvento.METODOS,
        'tipos_documento': DocumentoEvento.TIPOS,
        'tipos_aprobacion': AprobacionEvento.TIPOS,
        'estados_aprobacion': AprobacionEvento.ESTADOS,
        'preview_grupo': grupos[0] if grupos else None,
        'total_grupos': total_grupos,
        'total_lugares': total_lugares,
        'total_asistiran': total_asistiran,
        'total_no_asistiran': total_no_asistiran,
        'total_pendientes': total_pendientes,
        'total_adultos': total_adultos,
        'total_ninos': total_ninos,
        'total_adultos_persona': total_adultos_persona,
        'total_ninos_persona': total_ninos_persona,
        'porcentaje_asistencia': porcentaje_asistencia,
        'total_proveedores_evento': total_proveedores_evento,
        'total_servicios_evento': total_servicios_evento,
        'proveedores_pendientes': proveedores_pendientes,
        'total_personal_evento': total_personal_evento,
        'total_actividades_evento': total_actividades_evento,
        'actividades_pendientes': actividades_pendientes,
        'total_mesas_evento': total_mesas_evento,
        'mesas_excedidas': mesas_excedidas,
        'invitados_confirmados_sin_mesa': invitados_confirmados_sin_mesa,
        'total_gastos_evento': total_gastos_evento,
        'total_gasto_estimado': total_gasto_estimado,
        'total_gasto_real': total_gasto_real,
        'total_pagos_evento': total_pagos_evento,
        'gastos_vencidos': gastos_vencidos,
        'total_tareas_evento': total_tareas_evento,
        'tareas_pendientes': tareas_pendientes,
        'tareas_vencidas': tareas_vencidas,
        'total_documentos_evento': total_documentos_evento,
        'documentos_cliente': documentos_cliente,
        'total_aprobaciones_evento': total_aprobaciones_evento,
        'aprobaciones_pendientes': aprobaciones_pendientes,
        'notificaciones_no_leidas': notificaciones_no_leidas,
        'presupuesto_total': evento.presupuesto_total if evento else 0,
        'monto_pagado': evento.monto_pagado if evento else 0,
        'monto_pendiente': evento.monto_pendiente if evento else 0,
        'grupos': grupos,
        'base_url': base_url,
        'paletas': EventoBoda.PALETAS,
        'estilos_letra': EventoBoda.ESTILOS_LETRA,
        'tipos_evento': EventoBoda.TIPOS_EVENTO,
        'plantillas_evento': PLANTILLAS_EVENTO,
        'estados_evento': EventoBoda.ESTADOS_EVENTO,
        'tipos_regalo': EnlaceRegalo.TIPOS,
        'secciones_persona': PersonaCeremonia.SECCIONES,
        'tipos_invitacion': Grupoinvitacion.TIPO_INVITACION,
        'tipos_persona': Invitado.TIPO_PERSONA,
        'tipos_menu_invitado': Invitado.TIPO_MENU,
        'iconos_itinerario': ItinerarioEvento.ICONOS,
        'editor_invitacion_url': f'/dashboard/editor-invitacion/{evento.id}/' if evento else '',
        'api_dashboard_url': f'/api/dashboard/metricas/?evento={evento.id}' if evento else '',
    }

    if evento:
        from eventos.dashboard_v3 import construir_contexto_dashboard_evento_v3
        context.update(construir_contexto_dashboard_evento_v3(
            evento,
            servicios_evento=servicios_evento,
            tareas_evento=tareas_evento,
            actividades_evento=actividades_evento,
            gastos_evento=gastos_evento,
            documentos_evento=documentos_evento,
            invitados_confirmados_sin_mesa=invitados_confirmados_sin_mesa,
        ))

    return render(request, 'invitaciones/dashboard.html', context)


def construir_metricas_dashboard(evento):
    invitados_metricas = resumen_invitados_evento(evento)
    total_lugares = invitados_metricas['total_personas']
    total_asistiran = invitados_metricas['confirmados']
    total_no_asistiran = invitados_metricas['no_asisten']
    total_pendientes = invitados_metricas['pendientes']
    total_adultos = invitados_metricas['buffet_adultos']
    total_ninos = invitados_metricas['buffet_infantiles']

    gastos = GastoEvento.objects.filter(evento=evento) if evento else GastoEvento.objects.none()
    pagos = PagoEvento.objects.filter(gasto__evento=evento) if evento else PagoEvento.objects.none()
    tareas = TareaEvento.objects.filter(evento=evento) if evento else TareaEvento.objects.none()
    servicios = ServicioEvento.objects.filter(evento=evento) if evento else ServicioEvento.objects.none()

    gasto_estimado = gastos.aggregate(total=Sum('monto_estimado'))['total'] or 0
    gasto_real = gastos.aggregate(total=Sum('monto_real'))['total'] or 0
    total_pagado = pagos.aggregate(total=Sum('monto'))['total'] or 0
    gasto_objetivo = gasto_real or gasto_estimado
    saldo = gasto_objetivo - total_pagado

    return {
        'evento': str(evento) if evento else 'Sin evento',
        'resumen': {
            'total_lugares': total_lugares,
            'confirmados': total_asistiran,
            'pendientes': total_pendientes,
            'rechazados': total_no_asistiran,
            'adultos': total_adultos,
            'ninos': total_ninos,
            'gasto_estimado': float(gasto_estimado),
            'gasto_real': float(gasto_real),
            'pagado': float(total_pagado),
            'saldo': float(saldo if saldo > 0 else 0),
        },
        'charts': {
            'invitados': {
                'labels': ['Confirmados', 'Pendientes', 'No asisten'],
                'values': [total_asistiran, total_pendientes, total_no_asistiran],
            },
            'buffet': {
                'labels': ['Adultos', 'Ninos'],
                'values': [total_adultos, total_ninos],
            },
            'presupuesto': {
                'labels': ['Estimado', 'Real', 'Pagado', 'Saldo'],
                'values': [
                    float(gasto_estimado),
                    float(gasto_real),
                    float(total_pagado),
                    float(saldo if saldo > 0 else 0),
                ],
            },
            'tareas': {
                'labels': ['Pendiente', 'En proceso', 'En revision', 'Completada', 'Cancelada'],
                'values': [
                    tareas.filter(estado='PENDIENTE').count(),
                    tareas.filter(estado='EN_PROCESO').count(),
                    tareas.filter(estado='EN_REVISION').count(),
                    tareas.filter(estado='COMPLETADA').count(),
                    tareas.filter(estado='CANCELADA').count(),
                ],
            },
            'proveedores': {
                'labels': ['Pendientes', 'Contratados', 'Pagados/Listos', 'Cancelados'],
                'values': [
                    servicios.filter(estado__in=['SOLICITADO', 'COTIZADO', 'PENDIENTE_APROBACION', 'APROBADO']).count(),
                    servicios.filter(estado='CONTRATADO').count(),
                    servicios.filter(estado__in=['ANTICIPO_PAGADO', 'LIQUIDADO', 'SERVICIO_COMPLETADO']).count(),
                    servicios.filter(estado='CANCELADO').count(),
                ],
            },
        },
    }


def enlace_operativo_calendario(request, evento, admin_path, fragmento='resumen'):
    if usuario_es_dirtec(request.user):
        return admin_path
    return f'/dashboard/?evento={evento.id}#{fragmento}' if evento else ''


@login_required(login_url=LOGIN_DASHBOARD_URL)
def dashboard_metricas_json(request):
    evento, _ = obtener_evento_dashboard(request)
    if evento:
        exigir_permiso_evento(request, evento, Actions.EVENT_OPERATIONS)
    return JsonResponse(construir_metricas_dashboard(evento))

COMPONENTE_INVITACION_TIPOS = {'TEXTO', 'IMAGEN', 'BOTON'}


@login_required(login_url=LOGIN_DASHBOARD_URL)
def calendario_operativo(request):
    evento, eventos = obtener_evento_dashboard(request)
    if evento:
        exigir_permiso_evento(request, evento, Actions.EVENT_OPERATIONS)
    context = {
        'evento': evento,
        'eventos': eventos,
        'calendar_events_url': f'/api/calendario/eventos/?evento={evento.id}' if evento else '',
        'es_dirtec': usuario_es_dirtec(request.user),
    }
    return render(request, 'invitaciones/calendario_operativo.html', context)


@login_required(login_url=LOGIN_DASHBOARD_URL)
def calendario_eventos_json(request):
    evento, _ = obtener_evento_dashboard(request)
    eventos_calendario = []

    if not evento:
        return JsonResponse(eventos_calendario, safe=False)
    exigir_permiso_evento(request, evento, Actions.EVENT_OPERATIONS)

    for actividad in ActividadItinerario.objects.filter(evento=evento).select_related('proveedor', 'responsable'):
        inicio = datetime.combine(actividad.fecha, actividad.hora_inicio)
        fin = datetime.combine(actividad.fecha, actividad.hora_fin) if actividad.hora_fin else None
        eventos_calendario.append({
            'id': f'actividad-{actividad.id}',
            'title': actividad.titulo,
            'start': inicio.isoformat(),
            'end': fin.isoformat() if fin else None,
            'url': (f'/admin/itinerario/actividaditinerario/{actividad.id}/change/' if usuario_es_dirtec(request.user) else f'/dashboard/?evento={evento.id}#agenda'),
            'backgroundColor': color_evento_calendario(actividad.estado, actividad.prioridad),
            'borderColor': color_evento_calendario(actividad.estado, actividad.prioridad),
            'extendedProps': {
                'tipo': actividad.get_tipo_display(),
                'estado': actividad.get_estado_display(),
                'categoria': actividad.get_categoria_display(),
                'ubicacion': actividad.ubicacion or '',
                'responsable': str(actividad.responsable) if actividad.responsable else '',
                'proveedor': actividad.proveedor.nombre_comercial if actividad.proveedor else '',
            },
        })

    for tarea in TareaEvento.objects.filter(evento=evento):
        fecha_base = tarea.fecha_limite or tarea.fecha_inicio
        if not fecha_base:
            continue
        eventos_calendario.append({
            'id': f'tarea-{tarea.id}',
            'title': tarea.titulo,
            'start': fecha_base.isoformat(),
            'end': None,
            'allDay': True,
            'url': (
                f'/admin/tareas/tareaevento/{tarea.id}/change/'
                if usuario_es_dirtec(request.user)
                else f'/dashboard/?evento={evento.id}#tareas'
            ),
            'backgroundColor': '#9b4d36' if tarea.esta_vencida else '#d2b074',
            'borderColor': '#9b4d36' if tarea.esta_vencida else '#d2b074',
            'extendedProps': {
                'tipo': 'Tarea',
                'estado': tarea.get_estado_display(),
                'categoria': tarea.get_categoria_display(),
                'ubicacion': '',
                'responsable': str(tarea.responsable) if tarea.responsable else '',
                'proveedor': '',
                'descripcion': tarea.descripcion or '',
                'agenda': False,
            },
        })

    for gasto in GastoEvento.objects.filter(evento=evento).exclude(fecha_limite__isnull=True).select_related('categoria', 'proveedor'):
        eventos_calendario.append({
            'id': f'gasto-{gasto.id}',
            'title': f'Pago: {gasto.concepto}',
            'start': gasto.fecha_limite.isoformat(),
            'url': enlace_operativo_calendario(request, evento, f'/admin/presupuesto/gastoevento/{gasto.id}/change/', 'finanzas'),
            'backgroundColor': '#9b4d36' if gasto.esta_vencido else '#52627d',
            'borderColor': '#9b4d36' if gasto.esta_vencido else '#52627d',
            'extendedProps': {
                'tipo': 'Pago',
                'estado': gasto.get_estado_display(),
                'categoria': gasto.categoria.nombre,
                'ubicacion': '',
                'responsable': '',
                'proveedor': gasto.proveedor.nombre_comercial if gasto.proveedor else '',
            },
        })

    return JsonResponse(eventos_calendario, safe=False)


def color_evento_calendario(estado, prioridad):
    if estado in {'COMPLETADA', 'CANCELADA'}:
        return '#93a79b'
    if estado == 'RETRASADA' or prioridad == 'URGENTE':
        return '#9b4d36'
    if prioridad == 'ALTA':
        return '#d2b074'
    if estado in {'CONFIRMADA', 'EN_PROGRESO'}:
        return '#2f7a4d'
    return '#52627d'


@login_required(login_url=LOGIN_DASHBOARD_URL)
def mesas_visual(request):
    evento, eventos = obtener_evento_dashboard(request)
    if evento:
        exigir_permiso_evento(request, evento, Actions.EVENT_TABLES)
    if request.method == 'POST':
        evento = get_object_or_404(EventoBoda, id=request.POST.get('evento_id'))
        exigir_permiso_evento(request, evento, Actions.EVENT_TABLES)
        accion = request.POST.get('accion')
        tipos_mesa_validos = {valor for valor, _ in Mesa.TIPOS}
        tipo_mesa = request.POST.get('tipo_mesa')
        if tipo_mesa not in tipos_mesa_validos:
            tipo_mesa = 'REDONDA'

        if accion == 'crear_mesa':
            Mesa.objects.create(
                evento=evento,
                nombre=limpiar_texto(request, 'nombre_mesa') or 'Mesa',
                numero=convertir_entero(request.POST.get('numero_mesa'), 0) or None,
                tipo=tipo_mesa,
                capacidad=max(convertir_entero(request.POST.get('capacidad_mesa'), 10), 1),
                zona=limpiar_texto(request, 'zona_mesa'),
                notas=limpiar_texto(request, 'notas_mesa'),
            )
            messages.success(request, 'Mesa creada.')
        elif accion == 'editar_mesa':
            mesa = get_object_or_404(Mesa, evento=evento, id=request.POST.get('mesa_id'))
            mesa.nombre = limpiar_texto(request, 'nombre_mesa') or mesa.nombre
            mesa.numero = convertir_entero(request.POST.get('numero_mesa'), 0) or None
            mesa.tipo = tipo_mesa
            mesa.capacidad = max(convertir_entero(request.POST.get('capacidad_mesa'), mesa.capacidad), 1)
            mesa.zona = limpiar_texto(request, 'zona_mesa')
            mesa.notas = limpiar_texto(request, 'notas_mesa')
            mesa.save()
            messages.success(request, 'Mesa actualizada.')
        elif accion == 'eliminar_mesa':
            mesa = get_object_or_404(Mesa, evento=evento, id=request.POST.get('mesa_id'))
            nombre = mesa.nombre
            mesa.delete()
            messages.success(request, f'Mesa eliminada: {nombre}.')
        elif accion == 'asignar_mesa':
            mesa = get_object_or_404(
                Mesa,
                evento=evento,
                id=request.POST.get('mesa_id'),
            )
            invitado = get_object_or_404(
                Invitado,
                id=request.POST.get('invitado_id'),
                grupo__evento=evento,
            )
            try:
                AsignacionMesa.objects.create(
                    mesa=mesa,
                    invitado=invitado,
                    numero_asiento=convertir_entero(
                        request.POST.get('numero_asiento'),
                        0,
                    ) or None,
                    notas=limpiar_texto(
                        request,
                        'notas_asignacion',
                    ),
                )
                messages.success(
                    request,
                    f'{invitado} asignado a {mesa}.',
                )
            except ValidationError as exc:
                messages.error(
                    request,
                    ' '.join(exc.messages),
                )
        elif accion == 'eliminar_asignacion':
            asignacion = get_object_or_404(AsignacionMesa, id=request.POST.get('asignacion_id'), mesa__evento=evento)
            asignacion.delete()
            messages.success(request, 'Asignación eliminada.')

        return redirect(f'/dashboard/mesas/?evento={evento.id}')

    mesas = (
        Mesa.objects.filter(evento=evento)
        .select_related('evento', 'mesero')
        .prefetch_related('asignaciones__invitado__grupo')
        .order_by('numero', 'nombre')
    ) if evento else []
    invitados_sin_mesa = Invitado.objects.none()
    if evento:
        invitados_sin_mesa = (
            Invitado.objects
            .filter(
                grupo__evento=evento,
                asignaciones_mesa__isnull=True,
            )
            .select_related('grupo')
            .order_by(
                'grupo__nombre_grupo',
                'orden',
                'nombre',
            )
        )

    context = {
        'evento': evento,
        'eventos': eventos,
        'mesas': mesas,
        'tipos_mesa': Mesa.TIPOS,
        'invitados_sin_mesa': invitados_sin_mesa,
        'es_dirtec': usuario_es_dirtec(request.user),
    }
    return render(request, 'invitaciones/mesas_visual.html', context)


@require_POST
@login_required(login_url=LOGIN_DASHBOARD_URL)
def guardar_posiciones_mesas(request):
    evento = get_object_or_404(EventoBoda, id=request.POST.get('evento_id'))
    exigir_permiso_evento(request, evento, Actions.EVENT_TABLES)
    try:
        posiciones = json.loads(request.POST.get('posiciones', '[]'))
    except json.JSONDecodeError:
        return JsonResponse({'ok': False, 'error': 'Datos invalidos'}, status=400)

    mesas_actualizadas = 0
    for item in posiciones:
        mesa_id = item.get('id')
        if not mesa_id:
            continue

        mesa = Mesa.objects.filter(id=mesa_id, evento=evento).first()
        if not mesa:
            continue

        mesa.posicion_x = convertir_entero(item.get('x'), mesa.posicion_x)
        mesa.posicion_y = convertir_entero(item.get('y'), mesa.posicion_y)
        mesa.ancho = max(convertir_entero(item.get('ancho'), mesa.ancho), 90)
        mesa.alto = max(convertir_entero(item.get('alto'), mesa.alto), 90)
        mesa.save(update_fields=['posicion_x', 'posicion_y', 'ancho', 'alto'])
        mesas_actualizadas += 1

    return JsonResponse({'ok': True, 'mesas_actualizadas': mesas_actualizadas})


def guardar_personalizacion_evento(request, evento):
    campos_texto = [
        'paleta_colores',
        'paleta_sobre',
        'estilo_letra',
        'tipo_evento',
        'estado',
        'nombre_evento',
        'nombre_principal',
        'nombre_secundario',
        'etiqueta_principal',
        'etiqueta_secundario',
        'frase_portada',
        'titulo_invitacion',
        'texto_invitacion',
        'titulo_cuenta_regresiva',
        'titulo_detalles',
        'texto_detalles',
        'titulo_album',
        'titulo_menu',
        'titulo_regalos',
        'titulo_rsvp',
        'texto_rsvp',
        'lugar_misa',
        'direccion_ceremonia',
        'link_mapa_misa',
        'mapa_misa_embed',
        'lugar_fiesta',
        'direccion_recepcion',
        'link_mapa_fiesta',
        'mapa_fiesta_embed',
        'link_whatsapp',
        'dress_code',
        'dress_code_descripcion',
        'monograma_sobre',
        'texto_boton_sobre',
        'titulo_album_compartido',
        'texto_album_compartido',
        'link_album_compartido',
    ]
    campos_archivo = [
        'foto_portada',
        'foto_ceremonia',
        'foto_recepcion',
        'logo_portada',
        'sello_sobre',
        'cancion',
        'dress_code_permitido_imagen',
        'dress_code_prohibido_imagen',
        'fondo_invitacion',
        'fondo_cuenta_regresiva',
        'fondo_detalles',
        'fondo_album',
        'fondo_menu',
        'fondo_regalos',
        'fondo_rsvp',
    ]

    if 'actualizar_nombre_secundario' in request.POST:
        evento.mostrar_nombre_secundario = (
            request.POST.get(
                'mostrar_nombre_secundario'
            ) == 'on'
        )

    # Legacy visual toggles are intentionally not reset by the modern
    # Dashboard. Builder is the visual authority; these fields remain only
    # for migration/backward compatibility.
    empresa_id = request.POST.get('empresa_id')
    sede_id = request.POST.get('sede_id')
    if empresa_id:
        evento.empresa = empresa_autorizada_para_evento(request, evento, empresa_id)
    if sede_id:
        evento.sede = get_object_or_404(SedeEvento, id=sede_id, empresa=evento.empresa)
    elif 'sede_id' in request.POST:
        evento.sede = None
    if 'capacidad_contratada' in request.POST:
        evento.capacidad_contratada = convertir_entero(request.POST.get('capacidad_contratada'), 0)
    if 'precio_por_persona' in request.POST:
        evento.precio_por_persona = convertir_decimal(request.POST.get('precio_por_persona'), 0)

    if 'fecha_misa' in request.POST:
        fecha_misa = convertir_datetime_local(
            request.POST.get('fecha_misa')
        )
        if fecha_misa:
            evento.fecha_misa = fecha_misa

    if 'fecha_fiesta' in request.POST:
        fecha_fiesta = convertir_datetime_local(
            request.POST.get('fecha_fiesta')
        )
        if fecha_fiesta:
            evento.fecha_fiesta = fecha_fiesta

    for campo in campos_texto:
        if campo in request.POST:
            valor = request.POST.get(campo)
            defaults = {
                'frase_portada': 'Nuestra celebracion',
                'titulo_invitacion': 'Tu invitacion',
                'texto_invitacion': 'Hemos reservado este espacio para celebrar juntos este dia tan especial.',
                'titulo_cuenta_regresiva': 'Faltan',
                'titulo_detalles': 'Detalles del evento',
                'texto_detalles': 'Aqui tienes ceremonia, fiesta, ubicacion y mapas en un solo lugar.',
                'titulo_album': 'Album',
                'titulo_menu': 'Menu',
                'titulo_regalos': 'Regalos',
                'titulo_rsvp': 'Confirma tu asistencia',
                'texto_rsvp': 'Tu respuesta nos ayuda a organizar lugares, mesas y buffet.',
                'lugar_misa': 'Ceremonia por confirmar',
                'lugar_fiesta': 'Recepcion por confirmar',
            }
            if campo in defaults:
                valor = valor or defaults[campo]
            if campo == 'texto_boton_sobre':
                valor = valor or 'Abrir invitación'
            if campo == 'titulo_album_compartido':
                valor = valor or 'Comparte tus fotos'
            if campo == 'texto_album_compartido':
                valor = valor or 'Ayúdanos a guardar tus mejores momentos. Sube tus fotos y videos al álbum compartido.'
            if campo in {'mapa_misa_embed', 'mapa_fiesta_embed'}:
                valor = google_maps_src(valor)
            setattr(evento, campo, valor or None)

    for campo in campos_archivo:
        if request.POST.get(f'eliminar_{campo}') == 'on':
            archivo_actual = getattr(evento, campo, None)
            if archivo_actual:
                archivo_actual.delete(save=False)
            setattr(evento, campo, None)

    for campo in campos_archivo:
        archivo = request.FILES.get(campo)
        if archivo:
            if not archivo_pasa_validadores(EventoBoda, campo, archivo):
                continue
            setattr(evento, campo, archivo)

    plantilla_evento = request.POST.get('plantilla_evento')
    aplicar_plantilla_evento(evento, plantilla_evento)
    evento.save()
    aplicar_estructura_plantilla_evento(evento, plantilla_evento)

    if evento.mostrar_album_compartido:
        SeccionInvitacion.objects.update_or_create(
            evento=evento,
            tipo='ALBUM_COMPARTIDO',
            defaults={
                'titulo': evento.titulo_album_compartido,
                'descripcion': evento.texto_album_compartido,
                'orden': SECCIONES_INVITACION_DEFAULTS['ALBUM_COMPARTIDO']['orden'],
                'activa': True,
            },
        )


def agregar_regalo_dashboard(request, evento):
    tipo = request.POST.get('tipo_regalo') or 'TIENDA'
    tiene_datos_bancarios = any(
        request.POST.get(campo)
        for campo in ['banco', 'titular', 'numero_cuenta', 'clabe']
    )
    if tiene_datos_bancarios and not request.POST.get('url_regalo'):
        tipo = 'DEPOSITO'
    nombre = request.POST.get('nombre_regalo') or ('Deposito bancario' if tipo == 'DEPOSITO' else 'Mesa de regalos')
    EnlaceRegalo.objects.create(
        evento=evento,
        tipo=tipo,
        nombre=nombre,
        url=request.POST.get('url_regalo') or None,
        banco=request.POST.get('banco') or None,
        titular=request.POST.get('titular') or None,
        numero_cuenta=request.POST.get('numero_cuenta') or None,
        clabe=request.POST.get('clabe') or None,
        instrucciones=request.POST.get('instrucciones') or None,
        visible=True,
    )


def agregar_album_dashboard(request, evento):
    archivo = request.FILES.get('archivo_album')
    if not archivo:
        return
    if not archivo_pasa_validadores(FotoEvento, 'imagen', archivo):
        return

    FotoEvento.objects.create(
        evento=evento,
        titulo=request.POST.get('titulo_album_media') or None,
        imagen=archivo,
        orden=convertir_entero(request.POST.get('orden_album_media'), 0),
        visible=True,
    )


def agregar_itinerario_dashboard(request, evento):
    hora = request.POST.get('hora_itinerario')
    titulo = request.POST.get('titulo_itinerario')
    if not hora or not titulo:
        return

    ItinerarioEvento.objects.create(
        evento=evento,
        hora=hora,
        titulo=titulo,
        descripcion=request.POST.get('descripcion_itinerario') or None,
        icono=request.POST.get('icono_itinerario') or 'general',
        orden=convertir_entero(request.POST.get('orden_itinerario'), 0),
        visible=True,
    )


def preparar_links_grupo(grupo, base_url):
    link_invitacion = f"{base_url}/invitacion/{grupo.codigo}/"
    saludo = f"Hola {grupo.nombre_grupo}"
    if grupo.es_familiar:
        saludo = f"Hola Familia {grupo.nombre_grupo}"

    mensaje_invitacion = (
        f"{saludo}, te compartimos tu invitación digital.\n\n"
        f"Puedes verla aquí:\n{link_invitacion}"
    )
    mensaje_recordatorio = (
        f"{saludo}, este es un recordatorio de tu invitación digital.\n\n"
        f"Puedes revisarla aquí:\n{link_invitacion}"
    )
    telefono = (grupo.telefono_contacto or "").replace(" ", "").replace("+", "").replace("-", "")

    grupo.link_invitacion = link_invitacion
    grupo.whatsapp_invitacion_url = ""
    grupo.whatsapp_recordatorio_url = ""
    if telefono:
        grupo.whatsapp_invitacion_url = f"https://wa.me/{telefono}?text={quote(mensaje_invitacion)}"
        grupo.whatsapp_recordatorio_url = f"https://wa.me/{telefono}?text={quote(mensaje_recordatorio)}"


@login_required(login_url=LOGIN_DASHBOARD_URL)
def exportar_excel(request):
    evento, _ = obtener_evento_dashboard(request)
    if not evento:
        raise PermissionDenied('No hay un evento autorizado para exportar.')
    exigir_permiso_evento(request, evento, Actions.EVENT_GUESTS)

    asignaciones = AsignacionMesa.objects.select_related(
        'mesa'
    )
    invitados = (
        Invitado.objects
        .filter(grupo__evento=evento)
        .select_related('grupo')
        .prefetch_related(
            Prefetch(
                'asignaciones_mesa',
                queryset=asignaciones,
                to_attr='asignaciones_export',
            )
        )
        .order_by(
            'grupo__nombre_grupo',
            'grupo_id',
            'orden',
            'id',
        )
    ) if evento else Invitado.objects.none()

    wb = Workbook()
    ws = wb.active
    ws.title = 'Confirmaciones'

    ws.append([
        'Evento',
        'Invitación / grupo',
        'Tipo de invitación',
        'Persona',
        'Adulto o niño',
        'Buffet asignado',
        'Asistencia',
        'Mesa',
        'Acompañante extra',
        'Código',
    ])

    for invitado in invitados:
        asignacion = (
            invitado.asignaciones_export[0]
            if invitado.asignaciones_export
            else None
        )

        ws.append([
            str(evento),
            invitado.grupo.nombre_grupo,
            invitado.grupo.get_tipo_display(),
            invitado.nombre,
            invitado.get_tipo_persona_display(),
            invitado.menu_buffet_efectivo.title(),
            estado_asistencia(invitado.asistira),
            asignacion.mesa.nombre if asignacion else 'Sin mesa',
            'Sí' if invitado.es_acompanante_extra else 'No',
            str(invitado.grupo.codigo),
        ])

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    response = HttpResponse(
        content_type=(
            'application/vnd.openxmlformats-officedocument.'
            'spreadsheetml.sheet'
        )
    )
    response['Content-Disposition'] = (
        'attachment; filename=confirmaciones_evento.xlsx'
    )
    wb.save(response)
    return response


@login_required(login_url=LOGIN_DASHBOARD_URL)
def exportar_resumen_evento(request):
    evento, _ = obtener_evento_dashboard(request)
    if not evento:
        raise PermissionDenied('No hay un evento autorizado para exportar.')
    exigir_permiso_evento(request, evento, Actions.EVENT_OPERATIONS)
    metricas = construir_metricas_dashboard(evento)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Resumen del evento'

    ws.append(['Evento', metricas['evento']])
    ws.append([])
    ws.append(['Indicador', 'Valor'])
    for clave, valor in metricas['resumen'].items():
        ws.append([clave.replace('_', ' ').title(), valor])

    ws.append([])
    ws.append(['Proveedores'])
    ws.append(['Servicio', 'Proveedor', 'Estado', 'Costo proveedor'])
    for servicio in ServicioEvento.objects.filter(evento=evento).select_related('proveedor'):
        ws.append([
            servicio.nombre_servicio,
            servicio.proveedor.nombre_comercial if servicio.proveedor else '',
            servicio.get_estado_display(),
            servicio.costo_proveedor,
        ])

    ws.append([])
    ws.append(['Gastos'])
    ws.append(['Concepto', 'Categoria', 'Estado', 'Estimado', 'Real', 'Pagado', 'Saldo', 'Vence'])
    for gasto in GastoEvento.objects.filter(evento=evento).select_related('categoria'):
        ws.append([
            gasto.concepto,
            gasto.categoria.nombre,
            gasto.get_estado_display(),
            gasto.monto_estimado,
            gasto.monto_real,
            gasto.total_pagado,
            gasto.saldo_pendiente,
            gasto.fecha_limite.isoformat() if gasto.fecha_limite else '',
        ])

    ws.append([])
    ws.append(['Tareas'])
    ws.append(['Titulo', 'Categoria', 'Prioridad', 'Estado', 'Avance', 'Vence'])
    for tarea in TareaEvento.objects.filter(evento=evento):
        ws.append([
            tarea.titulo,
            tarea.get_categoria_display(),
            tarea.get_prioridad_display(),
            tarea.get_estado_display(),
            tarea.porcentaje_avance,
            tarea.fecha_limite.isoformat() if tarea.fecha_limite else '',
        ])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=resumen_evento.xlsx'
    wb.save(response)
    return response


def estado_asistencia(valor):
    if valor is True:
        return 'Sí'
    if valor is False:
        return 'No'
    return 'Pendiente'


@login_required(login_url=LOGIN_DASHBOARD_URL)
@require_POST
def marcar_envio_invitacion(request, grupo_id):
    grupo = get_object_or_404(Grupoinvitacion.objects.select_related('evento'), id=grupo_id)
    exigir_permiso_evento(request, grupo.evento, Actions.EVENT_GUESTS)
    grupo.estado_envio = 'INVITACION_PREPARADA'
    grupo.fecha_ultimo_envio = timezone.now()
    grupo.save()
    return redirect(f'/dashboard/?evento={grupo.evento_id}#invitados' if grupo.evento_id else 'dashboard')


@login_required(login_url=LOGIN_DASHBOARD_URL)
@require_POST
def marcar_recordatorio(request, grupo_id):
    grupo = get_object_or_404(Grupoinvitacion.objects.select_related('evento'), id=grupo_id)
    exigir_permiso_evento(request, grupo.evento, Actions.EVENT_GUESTS)
    grupo.estado_envio = 'RECORDATORIO_PREPARADO'
    grupo.fecha_ultimo_recordatorio = timezone.now()
    grupo.save()
    return redirect(f'/dashboard/?evento={grupo.evento_id}#invitados' if grupo.evento_id else 'dashboard')
